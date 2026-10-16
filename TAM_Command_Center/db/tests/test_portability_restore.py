"""T-175 (Run 6): restore-readiness checks against the existing schema.

Run: python -m unittest discover -s TAM_Command_Center/db/tests -p "test_portability*.py" -v
Proves: audit vocabulary covers restore events with no schema change; explicit-id
inserts (restored ids, timestamps, row_version) work with every trigger/unique guard;
parents-before-children works with foreign_keys=ON; the fixtures behave as labelled.
"""
from __future__ import annotations

import csv
import sqlite3
import sys
import unittest
import warnings
from pathlib import Path

DB_DIR = Path(__file__).resolve().parents[1]
FX = DB_DIR.parent / "tests" / "fixtures" / "portability"
sys.path.insert(0, str(DB_DIR))
sys.path.insert(0, str(FX))
import migrate  # noqa: E402
import build_fixtures as bf  # noqa: E402

warnings.filterwarnings("ignore", message=r"migrate\(channel=None\)")
ORDER = ["Accounts", "Contacts", "Tasks", "Notes", "Appliances", "Metrics",
         "Operations Quarters", "Operations Records"]
SKIP_COLS = {"account_name", "appliance_name"}  # display-only parent names
INT_COLS = {"is_primary", "archived", "enabled", "row_version", "stale_after_seconds",
            "stale_suppressed", "schedule_year"}


def fresh():
    c = sqlite3.connect(":memory:", isolation_level=None)
    c.execute("PRAGMA foreign_keys = ON")
    migrate.migrate(c, channel=None)
    return c


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f, delimiter=";" if "semicolon" in path.name else ","))


def insert(c, sheet, r):
    table, _ = bf.SHEETS[sheet]
    cols = [k for k in r if k not in SKIP_COLS]
    vals = [(int(r[k]) if k in INT_COLS and r[k] != "" else (None if r[k] == "" else r[k])) for k in cols]
    c.execute(f"INSERT INTO {table} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", vals)


def load(c, sheets=ORDER, folder="valid"):
    for s in sheets:
        for r in read_csv(FX / folder / f"{s.replace(' ', '_')}.csv"):
            insert(c, s, r)


class AuditVocabulary(unittest.TestCase):
    def test_restore_events_fit_existing_vocabulary(self):
        c = fresh()
        # Table-level restore event (D-130): one per restored table, no per-row events.
        for t in ("clients", "contacts", "tasks", "notes", "bams_objects", "bams_metrics",
                  "quarterly_upgrade_windows", "server_upgrades"):
            c.execute("INSERT INTO audit_events(category,object_type,object_id,object_label,action,new_value,source)"
                      " VALUES ('data','import',?,?,'restored','{\"added\":3,\"identical\":1}','Bulk Import')",
                      (t, t))
        # Whole-database restore events (D-140): post-restore (System) and pre-restore backup reference.
        c.execute("INSERT INTO audit_events(category,object_type,object_id,action,new_value,source)"
                  " VALUES ('data','backup','tam_backup_x_pre-restore.db','restored','{\"safety_backup\":\"x\"}','System')")
        c.execute("INSERT INTO audit_events(category,object_type,object_id,action,source)"
                  " VALUES ('system','schema','restore','restored','System')")
        self.assertEqual(c.execute("SELECT count(*) FROM audit_events").fetchone()[0], 10)
        with self.assertRaises(sqlite3.IntegrityError):  # an invented type is still refused
            c.execute("INSERT INTO audit_events(category,object_type,action,source)"
                      " VALUES ('data','restore','restored','System')")
        with self.assertRaises(sqlite3.DatabaseError):  # append-only still holds
            c.execute("UPDATE audit_events SET action='x' WHERE id=1")
        with self.assertRaises(sqlite3.DatabaseError):
            c.execute("DELETE FROM audit_events WHERE id=1")


class ExplicitIdInserts(unittest.TestCase):
    def test_full_valid_set_parents_first(self):
        c = fresh()
        load(c)
        counts = {t: c.execute(f"SELECT count(*) FROM {t}").fetchone()[0] for t, _ in bf.SHEETS.values()}
        self.assertEqual(counts, {"clients": 2, "contacts": 3, "tasks": 3, "notes": 1, "bams_objects": 1,
                                  "bams_metrics": 1, "quarterly_upgrade_windows": 1, "server_upgrades": 1})
        self.assertEqual(c.execute("PRAGMA foreign_key_check").fetchall(), [])
        self.assertEqual(c.execute("PRAGMA integrity_check").fetchone()[0], "ok")

    def test_triggers_do_not_rewrite_restored_values(self):
        c = fresh()
        load(c)
        r = c.execute("SELECT row_version, created_at, updated_at FROM clients WHERE id='fx-acc-1'").fetchone()
        self.assertEqual(r, (1, bf.TS, bf.TS))  # insert fires no bump/guard; timestamps kept as given
        c.execute("INSERT INTO clients(id,name,row_version,created_at,updated_at) VALUES ('x','Seven',7,?,?)",
                  (bf.TS, bf.TS))  # a high row_version is storable; planner resets New rows to 1 (D-127)
        # id immutability and version guards still apply after restore
        with self.assertRaises(sqlite3.DatabaseError):
            c.execute("UPDATE clients SET id='other' WHERE id='fx-acc-1'")
        c.execute("UPDATE clients SET status='Inactive' WHERE id='fx-acc-1'")
        self.assertEqual(c.execute("SELECT row_version FROM clients WHERE id='fx-acc-1'").fetchone()[0], 2)

    def test_children_before_parents_fail(self):
        c = fresh()
        with self.assertRaises(sqlite3.IntegrityError):
            load(c, ["Contacts"])
        c = fresh()
        load(c, ["Accounts"])
        with self.assertRaises(sqlite3.IntegrityError):  # release row must exist before an upgrade record
            load(c, ["Operations Records"])
        c = fresh()
        load(c, ["Accounts", "Appliances"])
        with self.assertRaises(sqlite3.IntegrityError):  # metric needs its appliance (passes: ordered) / orphan check
            c.execute("INSERT INTO bams_metrics(id,object_id,metric_name,ticker) VALUES ('m','nope','n','T1')")

    def test_orphan_and_invalid_fixtures_rejected(self):
        for name, sheet, prep in [
            ("missing_parent_Contacts.csv", "Contacts", []),
            ("missing_parent_Tasks.csv", "Tasks", []),
            ("bad_values_Tasks.csv", "Tasks", []),
            ("two_primaries_Contacts.csv", "Contacts", ["Accounts"]),
            ("dup_name_Accounts.csv", "Accounts", ["Accounts"]),
            ("dup_ticker_Metrics.csv", "Metrics", ["Accounts", "Appliances", "Metrics"]),
        ]:
            with self.subTest(name):
                c = fresh()
                load(c, prep)
                c.execute("SAVEPOINT s")
                with self.assertRaises(sqlite3.IntegrityError):
                    for r in read_csv(FX / "invalid" / name):
                        insert(c, sheet, r)
                c.execute("ROLLBACK TO s")
                self.assertEqual(c.execute("PRAGMA integrity_check").fetchone()[0], "ok")

    def test_reinsert_same_id_and_unique_guards(self):
        c = fresh()
        load(c)
        for sheet in ("Accounts", "Contacts", "Operations Quarters", "Operations Records"):
            with self.assertRaises(sqlite3.IntegrityError, msg=sheet):  # planner must classify Identical/Differs first
                insert(c, sheet, read_csv(FX / "valid" / f"{sheet.replace(' ', '_')}.csv")[0])
        # same upgrade identity under a new id is refused (D-127 'upgrade identity')
        r = read_csv(FX / "valid" / "Operations_Records.csv")[0]
        r["id"] = "fx-ops-dup"
        with self.assertRaises(sqlite3.IntegrityError):
            insert(c, "Operations Records", r)

    def test_atomic_rollback(self):
        c = fresh()
        c.execute("BEGIN")
        load(c, ["Accounts", "Contacts"])
        try:
            insert(c, "Contacts", read_csv(FX / "invalid" / "two_primaries_Contacts.csv")[0])
            insert(c, "Contacts", read_csv(FX / "invalid" / "two_primaries_Contacts.csv")[1])
        except sqlite3.IntegrityError:
            c.execute("ROLLBACK")
        self.assertEqual(c.execute("SELECT count(*) FROM clients").fetchone()[0], 0)


class Fixtures(unittest.TestCase):
    def test_headers_are_unambiguous_signatures(self):
        sigs = {s: frozenset(cols) for s, (_, cols) in bf.SHEETS.items()}
        for a in sigs:
            for b in sigs:
                if a != b:
                    self.assertFalse(sigs[a] <= sigs[b], f"{a} header set is a subset of {b}: ambiguous")
        for s in ORDER:
            hdr = set(next(csv.reader(open(FX / "valid" / f"{s.replace(' ', '_')}.csv", encoding="utf-8"))))
            self.assertEqual(hdr, sigs[s])
        hdr = set(next(csv.reader(open(FX / "invalid" / "Unrecognized.csv"))))
        self.assertFalse(any(hdr >= sig for sig in sigs.values()))

    def test_variants_load_same_rows(self):
        base = read_csv(FX / "valid" / "Accounts.csv")
        for v in ("semicolon", "bom", "shuffled"):
            self.assertEqual(read_csv(FX / f"variants_Accounts_{v}.csv"), base if v != "shuffled"
                             else [dict(r) for r in base])

    def test_xlsx_fixtures(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        for name, about_schema, channel, sheets in [
            ("valid_export_dev.xlsx", "1", "DEV", True), ("valid_export_prod.xlsx", "1", "PROD", True),
            ("newer_schema.xlsx", "999", "PROD", True), ("not_an_export.xlsx", None, None, False)]:
            wb = openpyxl.load_workbook(FX / "xlsx" / name)
            if sheets:
                about = {r[0]: r[1] for r in wb["About"].iter_rows(min_row=2, values_only=True)}
                self.assertEqual((about["schema_version"], about["channel"]), (about_schema, channel))
                self.assertEqual(wb.sheetnames, ["About"] + ORDER)
            else:
                self.assertNotIn("About", wb.sheetnames)
        # the valid workbook restores cleanly
        c = fresh()
        wb = openpyxl.load_workbook(FX / "xlsx" / "valid_export_prod.xlsx")
        for s in ORDER:
            rows = list(wb[s].iter_rows(values_only=True))
            for vals in rows[1:]:
                insert(c, s, {k: ("" if v is None else str(v)) for k, v in zip(rows[0], vals)})
        self.assertEqual(c.execute("PRAGMA foreign_key_check").fetchall(), [])


if __name__ == "__main__":
    unittest.main()
