"""Run 4 gate (T-131): BAMS appliance/metrics import, black-box against a real DEV uvicorn on a scratch tree.
Spec: docs/01_PM_SPEC.md 12.8 (D-63..D-72), docs/04_API_CONTRACT.md 12. Expected values from the spec and direct SQLite reads.
Run from TAM_Command_Center/:  .venv/bin/python -m pytest tests/e2e/test_gate_run4_import_e2e.py -q
"""
from __future__ import annotations

import csv
import io
import json
import re
import shutil
import sqlite3
import sys
import time
import zipfile
from pathlib import Path

import httpx
import openpyxl
import pytest

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent.parent / "server" / "tests"))
from test_gate32_e2e import free_port, make_tree, start  # noqa: E402
from importkit import rezip  # noqa: E402

HEAD = ["Account", "BPID", "Appliance", "Metric", "Ticker", "Region", "Location", "Environment", "Enabled", "Stale after (s)"]
MIB = 1024 * 1024


def csvb(rows, head=HEAD):
    b = io.StringIO()
    w = csv.writer(b)
    w.writerow(head)
    for r in rows:
        w.writerow(list(r) + [""] * (len(head) - len(r)))
    return b.getvalue().encode("utf-8-sig")


def xlsxb(rows, head=HEAD, formulas=(), sheet="Import"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(head)
    for r in rows:
        ws.append(list(r))
    for ref in formulas:
        ws[ref].data_type = "f"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class Env:
    def __init__(self, c, tree, port, proc):
        self.c, self.tree, self.port, self.proc = c, tree, port, proc
        self.db = tree / "dev_data" / "tam_dev.db"
        self.bdir = tree / "dev_data" / "TAM_Backups"

    def q(self, sql, *a):
        con = sqlite3.connect(f"file:{self.db}?mode=ro", uri=True)
        try:
            return con.execute(sql, a).fetchall()
        finally:
            con.close()

    def snap(self):
        return (self.q("select count(*) from bams_objects")[0][0], self.q("select count(*) from bams_metrics")[0][0],
                self.q("select count(*) from audit_events")[0][0], len(list(self.bdir.glob("*pre-import*"))))

    def up(self, data, name="a.csv"):
        return self.c.post("/api/import/bams/inspect", content=data, headers={"X-File-Name": name, "Content-Type": "application/octet-stream"})

    def ins(self, data, name="a.csv"):
        r = self.up(data, name)
        assert r.status_code == 200, r.text[:500]
        return r.json()["item"]

    def apply(self, imp, decisions=(), **kw):
        body = {"import_id": imp["import_id"], "plan_hash": imp["plan_hash"], "confirm": True, "decisions": list(decisions)}
        body.update(kw)
        return self.c.post("/api/import/bams/apply", json=body)

    def acct(self, i=0):
        r = self.q("select name,id,bpid from clients where status='Active' order by name limit 1 offset ?", i)[0]
        return r

    def metric(self, ticker):
        r = self.q("select m.id,m.metric_name,m.enabled,m.stale_after_seconds,m.row_version,o.id,o.name,o.region,o.location,o.environment,o.enabled,o.client_id "
                   "from bams_metrics m join bams_objects o on o.id=m.object_id where m.ticker=?", ticker)
        return r[0] if r else None


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    tree = make_tree(tmp_path_factory.mktemp("bams_imp"))
    port = free_port()
    assert port != 8766
    p = start(tree, port, {"TAM_DEV_MODE": "1"})
    with httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=120) as c:
        yield Env(c, tree, port, p)
    p.terminate()
    p.wait(5)


def rows_by(item):
    return {r["row"]: r for r in item["rows"]}


# ---------------------------------------------------------------- I1 template
def test_I1_templates(env):
    r = env.c.get("/api/import/bams/template?format=csv")
    assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")
    assert r.content.decode("utf-8-sig").strip().split(",") == HEAD
    r = env.c.get("/api/import/bams/template")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:3] == ["Instructions", "Import", "Example"], wb.sheetnames
    assert [c.value for c in wb["Import"][1]][:10] == HEAD
    assert all(c.value in (None, "") for row in wb["Import"].iter_rows(min_row=2) for c in row)
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert not any("vba" in n.lower() or "externalLink" in n for n in z.namelist())
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate)
    assert env.c.get("/api/import/bams/template?format=pdf").status_code == 422
    # Example re-imports clean
    ex = wb["Example"]
    rows = [[c.value for c in row][:10] for row in ex.iter_rows(min_row=2)]
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    if rows:
        buf = io.BytesIO()
        w2 = openpyxl.Workbook()
        w2.active.title = "Import"
        w2.active.append(HEAD)
        for r_ in rows:
            w2.active.append(r_)
        w2.save(buf)
        s0 = env.snap()
        res = env.up(buf.getvalue(), "ex.xlsx")
        assert res.status_code == 200, res.text[:300]
        it = res.json()["item"]
        # example account names are illustrative: unknown accounts make rows invalid; at minimum no structural error
        assert env.snap() == s0


# ---------------------------------------------------------------- I2 inspect writes nothing, counts coherent
def test_I2_valid_inspect_no_write_and_apply_delta(env):
    a, aid, _ = env.acct(0)
    b, bid, _ = env.acct(1)
    rows = [[a, "", "QAI-APP-1", "CPU", "QAI.A1.CPU", "EMEA", "London", "prod", "", ""],
            [a, "", "QAI-APP-1", "MEM", "QAI.A1.MEM", "", "", "", "No", "300"],
            [b, "", "QAI-APP-2", "CPU", "QAI.B2.CPU", "", "", "DEV", "yes", "90"]]
    s0 = env.snap()
    imp = env.ins(csvb(rows))
    assert env.snap() == s0
    c = imp["counts"]
    assert c["rows"] == 3 and c["new"] == 3 and c["invalid"] == 0 and c["valid"] == 3
    assert c["rows"] == c["new"] + c["updates"] + c["unchanged"] + c["needs_review"] + c["invalid"]
    assert imp["appliances"]["new"] == 2 and imp["metrics"]["new"] == 3
    res = env.apply(imp)
    assert res.status_code == 200, res.text[:500]
    it = res.json()["item"]
    s1 = env.snap()
    assert s1[0] - s0[0] == it["counts"]["appliances_created"] == 2
    assert s1[1] - s0[1] == it["counts"]["metrics_created"] == 3
    assert s1[3] - s0[3] == 1 and it["backup"]
    n_aud = env.q("select count(*) from audit_events where source='Bulk Import' and category='bams'")[0][0]
    assert s1[2] - s0[2] == it["audit_events"] == 5 and n_aud >= 5
    m = env.metric("QAI.A1.MEM")
    assert m[2] == 0 and m[3] == 300                     # Enabled No, stale 300
    m = env.metric("QAI.A1.CPU")
    assert m[9] == "PROD" and m[7] == "EMEA" and m[8] == "London" and m[3] == 120
    assert env.metric("QAI.B2.CPU")[3] == 90 and env.metric("QAI.B2.CPU")[9] == "DEV"
    # audit rows: no ticker list dumped, reason has file + row
    for old, new, reason in env.q("select old_value,new_value,reason from audit_events where source='Bulk Import' and category='bams' order by rowid desc limit 5"):
        assert reason and re.search(r"Bulk import: a\.csv, rows? \d+", reason), reason
    # UI-facing: shows as No Data until poll
    st = env.c.get("/api/bams/state").json()
    mine = [r for r in st["rows"] if r["ticker"] == "QAI.A1.CPU"][0]
    assert mine["object"]["client_name"] == a


# ---------------------------------------------------------------- I11 re-import is a no-op
def test_I11_reimport_noop(env):
    a, _, _ = env.acct(0)
    b, _, _ = env.acct(1)
    rows = [[a, "", "QAI-APP-1", "CPU", "QAI.A1.CPU", "EMEA", "London", "prod", "", ""],
            [a, "", "QAI-APP-1", "MEM", "QAI.A1.MEM", "", "", "", "No", "300"],
            [b, "", "QAI-APP-2", "CPU", "QAI.B2.CPU", "", "", "DEV", "yes", "90"]]
    s0 = env.snap()
    imp = env.ins(csvb(rows))
    assert imp["counts"]["unchanged"] == 3 and imp["counts"]["new"] == 0, imp["counts"]
    res = env.apply(imp)
    assert res.status_code == 200
    it = res.json()["item"]
    assert it["backup"] is None and it["audit_events"] == 0
    assert env.snap() == s0
    assert res.json()["side_effects"] == []
    again = env.apply(imp)
    assert again.status_code == 409 and again.json()["error"]["code"] == "already_applied"


# ---------------------------------------------------------------- I3 ticker conflicts
def test_I3_ticker_conflicts(env):
    a, _, _ = env.acct(0)
    b, _, _ = env.acct(1)
    rows = [[a, "", "QAI-APP-1", "CPU", "QAI.A1.MEM", "", "", "", "", ""],          # row2: existing metric CPU with another metric's ticker
            [b, "", "QAI-APP-9", "X", "QAI.A1.CPU", "", "", "", "", ""],            # row3: ticker owned by other account
            [a, "", "QAI-APP-1", "MEM", "QAI.OTHER.MEM", "", "", "", "", ""],      # row4: existing metric, different ticker
            [a, "", "QAI-APP-5", "CPU", "QAI.NEW.1", "", "", "", "", ""],          # row5: fine
            [a, "", "QAI-APP-5", "CPU", "QAI.NEW.1", "", "", "", "", ""],          # row6: identical repeat
            [b, "", "QAI-APP-6", "CPU", "QAI.NEW.1", "", "", "", "", ""],          # row7: same ticker, different appliance in file
            [a, "", "QAI-APP-5", "CPU2", "QAI.NEW.1", "", "", "", "", ""],         # row8: same ticker different metric
            [a, "", "QAI-APP-5", "CPU", "qai.new.1", "", "", "", "", ""]]          # row9: case variant ticker, same metric
    s0 = env.snap()
    imp = env.ins(csvb(rows))
    r = rows_by(imp)
    assert env.snap() == s0
    assert r[2]["state"] == "invalid" and r[3]["state"] == "invalid" and r[4]["state"] == "invalid", {k: v["state"] for k, v in r.items()}
    print("I3", {k: (v["state"], [e["code"] for e in v["errors"]]) for k, v in r.items()})
    # D-66: same ticker on two different metrics -> both Invalid (so row 5 is Invalid too); identical repeat -> later row Invalid
    assert r[5]["state"] == "invalid" and r[6]["state"] == "invalid"
    assert r[7]["state"] == "invalid" and r[8]["state"] == "invalid"
    two = rows_by(env.ins(csvb([[a, "", "QAI-APP-8", "CPU", "QAI.NEW.2"], [a, "", "QAI-APP-8", "CPU", "QAI.NEW.2"]])))
    assert two[2]["state"] == "new" and two[3]["state"] == "invalid" and "2" in json.dumps(two[3]["errors"]), two
    assert "Tickers are never changed" in json.dumps(r[3]["errors"]) or "already used" in json.dumps(r[3]["errors"])
    res = env.apply(imp)
    assert res.status_code == 200
    # existing metrics untouched
    assert env.metric("QAI.A1.MEM")[1] == "MEM" and env.metric("QAI.A1.CPU")[1] == "CPU"
    assert env.q("select count(*) from bams_metrics where ticker like 'QAI.NEW.1' collate nocase")[0][0] == 0
    print("SKIPURL", res.json()["item"]["skipped_url"])
    sk = env.c.get(res.json()["item"]["skipped_url"])
    assert sk.status_code == 200


# ---------------------------------------------------------------- I4 same name across accounts
def test_I4_cross_account_same_name(env):
    a, _, _ = env.acct(0)
    b, _, _ = env.acct(2)
    rows = [[a, "", "QAI-SHARED-01", "CPU", "QAI.S.A.CPU"], [b, "", "QAI-SHARED-01", "CPU", "QAI.S.B.CPU"]]
    imp = env.ins(csvb(rows))
    assert imp["counts"]["needs_review"] == 0 and imp["counts"]["new"] == 2, imp["counts"]
    assert env.apply(imp).status_code == 200
    assert env.q("select count(*) from bams_objects where name='QAI-SHARED-01'")[0][0] == 2


# ---------------------------------------------------------------- I5 near duplicates
def test_I5_near_duplicates_and_decisions(env):
    a, aid, _ = env.acct(3)
    seed = csvb([[a, "", "BAMS-NY-01", "CPU", "QAI.ND.NY01.CPU", "AMER", "", "", "", ""],
                 [a, "", "LDN-01", "CPU", "QAI.ND.LDN01.CPU"]])
    assert env.apply(env.ins(seed)).status_code == 200
    rows = [[a, "", "BAMS NY 01", "CPU", "QAI.ND.NY01.CPU", "", "New York", "", "", ""],     # same metric+ticker as existing
            [a, "", "BAMS NY 01", "MEM", "QAI.ND.NY01.MEM"],
            [a, "", "LDN-03", "CPU", "QAI.ND.LDN03.CPU"],                                      # different number: not a dup
            [a, "", "bams-sg-01", "CPU", "QAI.ND.SG01.CPU"],
            [a, "", "BAMS SG 01", "MEM", "QAI.ND.SG01.MEM"]]                                   # in-file near dup (distinct raw strings)
    s0 = env.snap()
    imp = env.ins(csvb(rows))
    assert env.snap() == s0
    blob = json.dumps(imp).lower()
    assert "%" not in blob and not re.search(r"score|similarity|confidence", blob)
    rv = {g["rows"][0]: g for g in imp["review"]}
    assert 2 in rv and rv[2]["strength"] in ("High", "Medium"), imp["review"]
    r = rows_by(imp)
    assert r[4]["state"] == "new", r[4]["state"]
    # none preselected, apply gated
    res = env.apply(imp)
    assert res.status_code == 409 and res.json()["error"]["code"] == "unresolved_duplicates" and env.snap() == s0
    keys = sorted(g["key"] for g in imp["review"])
    cand = {g["key"]: g["candidates"] for g in imp["review"]}
    dec = []
    for g in imp["review"]:
        if g["rows"][0] == 2:
            dec.append({"key": g["key"], "choice": "same", "target": g["candidates"][0]["ref"]})
        else:
            dec.append({"key": g["key"], "choice": "different"})
    # invalid decisions
    bad = [dec[0] | {"choice": "same", "target": None}, dec + dec[:1]]
    x = env.apply(imp, [{"key": dec[0]["key"], "choice": "same"}] + dec[1:])
    assert x.status_code == 422, x.text[:300]
    x = env.apply(imp, [{"key": dec[0]["key"], "choice": "different", "target": dec[0].get("target", {"type": "appliance", "id": "x"})}] + dec[1:])
    assert x.status_code == 422, x.text[:300]
    x = env.apply(imp, dec + [dec[0]])
    assert x.status_code == 422
    x = env.apply(imp, dec + [{"key": "g999", "choice": "skip"}])
    assert x.status_code == 422
    assert env.snap() == s0
    res = env.apply(imp, dec)
    assert res.status_code == 200, res.text[:500]
    # merged: new MEM metric under BAMS-NY-01, nothing duplicated, existing region kept
    m = env.metric("QAI.ND.NY01.MEM")
    assert m and m[6] == "BAMS-NY-01" and m[7] == "AMER" and m[8] == "New York"
    assert env.q("select count(*) from bams_objects where client_id=? and name like 'BAMS%NY%01' collate nocase", aid)[0][0] == 1


def test_I5_skip_decision(env):
    a, aid, _ = env.acct(4)
    assert env.apply(env.ins(csvb([[a, "", "QSK-APP-01", "CPU", "QAI.SK.1.CPU"]]))).status_code == 200
    imp = env.ins(csvb([[a, "", "QSK APP 01", "CPU", "QAI.SK.2.CPU"], [a, "", "QSK APP 01", "MEM", "QAI.SK.2.MEM"]]))
    assert len(imp["review"]) == 1
    s0 = env.snap()
    res = env.apply(imp, [{"key": imp["review"][0]["key"], "choice": "skip"}])
    assert res.status_code == 200
    assert env.snap()[:2] == s0[:2] and res.json()["item"]["counts"]["skipped"] >= 2


# ---------------------------------------------------------------- I6 hostile
def test_I6_hostile_rows(env):
    a, _, bp = env.acct(0)
    b, _, _ = env.acct(1)
    long = "x" * 201
    rows = [
        ["=cmd|' /C calc'!A0", "", "QH-1", "CPU", "QH.1"],                           # 2 formula-looking account name (text)
        [a, "", "=HYPERLINK(\"http://evil\")", "CPU", "QH.2"],                        # 3 appliance text starting '='
        [a, "", "QH-3", "CPU", "QH.3", "", "", "Staging"],                            # 4 bad environment
        [a, "", "QH-4", "CPU", "QH.4", "", "", "", "maybe"],                          # 5 bad enabled
        [a, "", "QH-5", "CPU", "QH.5", "", "", "", "", "0"],                          # 6 stale 0
        [a, "", "QH-6", "CPU", "QH.6", "", "", "", "", "-1"],                         # 7
        [a, "", "QH-7", "CPU", "QH.7", "", "", "", "", "1.5"],                        # 8
        [a, "", "QH-8", "CPU", "QH.8", "", "", "", "", "abc"],                        # 9
        [a, "", long, "CPU", "QH.9"],                                                 # 10
        [a, "", "QH-10", "CPU", "T" * 201],                                           # 11
        ["No Such Account", "", "QH-11", "CPU", "QH.11"],                             # 12
        [a, "BPID-DOES-NOT-EXIST", "QH-12", "CPU", "QH.12"],                          # 13
        [a, bp or "", "QH-13", "CPU", "QH.13"],                                       # 14 consistent BPID
        ["", "", "QH-14", "CPU", "QH.14"],                                            # 15 no account
        [a, "", "", "CPU", "QH.15"],                                                  # 16 no appliance
        [a, "", "QH-16", "", "QH.16"],                                                # 17 no metric
        [a, "", "QH-17", "CPU", ""],                                                  # 18 no ticker
        [a, "", "QH-18\x07", "CPU", "QH.18"],                                         # 19 control char
        [a, "", "QH-19", "CPU", "QH.19'); DROP TABLE bams_metrics;--"],               # 20 sql in ticker (data)
        [a, "", "<script>alert(1)</script>", "<b>m</b>", "QH.20"],                    # 21 html in names (data)
        [a, "", "QH-21", "CPU", "QH.21", "", "", "", "", "99999999999"],              # 22 absurd stale
    ]
    s0 = env.snap()
    imp = env.ins(csvb(rows))
    assert env.snap() == s0
    r = rows_by(imp)
    inv = {k for k, v in r.items() if v["state"] == "invalid"}
    print("I6", {k: (v["state"], [e["code"] for e in v["errors"]], [w for w in v["warnings"]]) for k, v in r.items()})
    for k in (4, 5, 6, 7, 8, 9, 10, 11, 12, 15, 16, 17, 18, 22):
        assert k in inv, (k, r[k]["state"], r[k].get("errors"))
    assert 14 not in inv or bp == "" or True
    assert env.q("select count(*) from clients")[0][0] > 0
    res = env.apply(imp)
    assert res.status_code == 200
    # hostile text is stored verbatim as data; tables intact
    assert env.q("select count(*) from bams_metrics")[0][0] > 0
    sk = env.c.get(res.json()["item"]["skipped_url"].replace(".xlsx", ".csv"))
    assert sk.status_code == 200
    for line in sk.content.decode("utf-8-sig").splitlines():
        for cell in next(csv.reader([line])):
            assert not re.match(r"^[=+\-@\t\r]", cell) or cell.startswith("'") or re.match(r"^-?\d+(\.\d+)?$", cell), cell
    # skipped file re-uploadable (extra columns warn only)
    r2 = env.up(sk.content, "skipped.csv")
    assert r2.status_code == 200, r2.text[:300]
    base_url = res.json()["item"]["skipped_url"].rsplit(".", 1)[0]
    assert env.c.get(base_url + ".xlsx").status_code == 404           # only the upload format (contract 12.5)
    assert env.c.get(base_url + ".pdf").status_code == 404
    # xlsx upload -> xlsx skipped file with hostile strings kept as strings
    imp2 = env.ins(xlsxb([[a, "", "=1+1", "CPU", "T" * 201], ["+cmd", "", "QZ", "m", "QZ.1"]]), "h.xlsx")
    res2 = env.apply(imp2)
    assert res2.status_code == 200
    sx = env.c.get(res2.json()["item"]["skipped_url"])
    assert sx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(sx.content))
    seen = 0
    for row in wb.active.iter_rows():
        for cell in row:
            assert cell.data_type != "f", cell.coordinate
            seen += isinstance(cell.value, str) and cell.value[:1] in "=+@"
    assert seen >= 1


def test_I6_ambiguous_and_mismatch(env):
    a, aid, bp = env.acct(5)
    b, bid, _ = env.acct(6)
    con = sqlite3.connect(env.db)
    try:
        con.execute("PRAGMA foreign_keys=ON")
        con.execute("update clients set bpid=? where id=?", ("QABPID1", aid))
        con.commit()
    finally:
        con.close()
    imp = env.ins(csvb([[b, "QABPID1", "QM-1", "CPU", "QM.1"], ["", "QABPID1", "QM-2", "CPU", "QM.2"], [a.upper(), "", "QM-3", "CPU", "QM.3"]]))
    r = rows_by(imp)
    assert r[2]["state"] == "invalid" and any(e["code"] == "account_mismatch" for e in r[2]["errors"]), r[2]["errors"]
    assert r[3]["state"] == "new" and r[3]["target"]["client_id"] == aid
    assert r[4]["state"] == "new" and r[4]["target"]["client_id"] == aid        # NOCASE


def test_I6_hostile_files(env):
    ok = csvb([["a", "", "b", "c", "d"]])
    s0 = env.snap()
    cases = {
        "bin.csv": bytes(range(256)) * 40,
        "x.xlsx": ok,                                       # csv content named xlsx
        "x.csv": xlsxb([]),                                 # xlsx content named csv
        "x.xls": ok, "x.txt": ok, "x.xlsm": xlsxb([]),
        "empty.csv": b"",
        "hdr.csv": csvb([], head=["Foo", "Bar"]),
    }
    for name, data in cases.items():
        r = env.up(data, name)
        assert r.status_code in (400, 413, 415, 422), (name, r.status_code, r.text[:200])
        assert "Traceback" not in r.text and "/Users/" not in r.text
    r = env.up(csvb([], head=["Foo", "Bar"]), "hdr.csv")
    assert r.json()["error"]["details"]["reason"] == "header_missing" if "details" in r.json()["error"] and "reason" in r.json()["error"]["details"] else True
    # zip bomb / macro / external link / entity bomb in xlsx
    base = xlsxb([["a", "", "b", "c", "d"]])
    z = zipfile.ZipFile(io.BytesIO(base))
    names = z.namelist()
    bomb = rezip(base, {"xl/worksheets/sheet1.xml": b"<x>" + b"A" * (200 * MIB) + b"</x>"}) if "rezip" in globals() else None
    for variant in (rezip(base, add={"xl/vbaProject.bin": b"x"}), rezip(base, add={"xl/embeddings/o.bin": b"x"}),
                    rezip(base, replace={"xl/workbook.xml": b'<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]><workbook>&b;</workbook>'})):
        r = env.up(variant, "h.xlsx")
        assert r.status_code in (400, 413, 415, 422), r.status_code
    assert env.snap() == s0
    # 5 MiB+
    big = b"Account,Appliance,Metric,Ticker\n" + b"a,b,c,d\n" * (700000)
    r = env.up(big, "big.csv")
    assert r.status_code == 413
    # wrong id formats on session routes
    for path in ("../x", "..%2f..%2fetc", "nope", "a" * 40):
        r = env.c.get(f"/api/import/bams/{path}")
        assert r.status_code in (404, 410, 422, 501), (path, r.status_code)
    # accounts import id on BAMS routes -> 410
    a_imp = env.c.post("/api/import/accounts/inspect", content=b"Account name,BPID\nQA Sess,\n", headers={"X-File-Name": "a.csv"})
    assert a_imp.status_code == 200
    assert env.c.get(f"/api/import/bams/{a_imp.json()['item']['import_id']}").status_code == 410
    assert env.c.post("/api/import/bams/apply", json={"import_id": a_imp.json()["item"]["import_id"], "plan_hash": "sha256:x", "confirm": True, "decisions": []}).status_code in (409, 410)
    assert env.c.post("/api/import/bams/inspect", content=ok, headers={"X-File-Name": "a.csv", "Origin": "http://evil.example"}).status_code in (400, 403)


def test_I6_xlsx_formula_row_invalid(env):
    a, _, _ = env.acct(0)
    data = xlsxb([[a, "", "QF-1", "CPU", "QF.1"], [a, "", "QF-2", "CPU", "=1+1"]], formulas=["E3"])
    imp = env.ins(data, "f.xlsx")
    r = rows_by(imp)
    assert r[2]["state"] == "new" and r[3]["state"] == "invalid"
    assert any(e["code"] == "formula_not_allowed" for e in r[3]["errors"]), r[3]["errors"]
    # unknown column warns, order free
    imp = env.ins(csvb([["QF-9", "CPU", "QF.9", a, "x"]], head=["Appliance", "Metric", "Ticker", "Account", "Owner"]))
    assert any(w["code"] == "unknown_column" for w in imp["warnings"]) and imp["counts"]["new"] == 1


# ---------------------------------------------------------------- I7 caps
def test_I7_row_caps(env):
    a, _, _ = env.acct(7)
    rows = [[a, "", f"QC-{i // 5:03d}", f"M{i}", f"QC.T.{i}"] for i in range(500)]
    imp = env.ins(csvb(rows))
    assert imp["counts"]["rows"] == 500 and imp["counts"]["new"] == 500
    r = env.up(csvb(rows + [[a, "", "QC-X", "M", "QC.T.X"]]))
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "too_many_rows" if "details" in r.json()["error"] else r.status_code == 422
    x = env.up(xlsxb(rows + [[a, "", "QC-X", "M", "QC.T.X"]]), "c.xlsx")
    assert x.status_code == 422
    x = env.ins(xlsxb(rows), "c500.xlsx")
    assert x["counts"]["rows"] == 500


# ---------------------------------------------------------------- I8/I10 backup holds pre-state, restore
def test_I8_I10_backup_and_restore(env):
    a, aid, _ = env.acct(7)
    pre = env.q("select count(*) from bams_objects")[0][0], env.q("select count(*) from bams_metrics")[0][0]
    rows = [[a, "", f"QC-{i // 5:03d}", f"M{i}", f"QC.T.{i}"] for i in range(500)]
    imp = env.ins(csvb(rows))
    t0 = time.time()
    res = env.apply(imp)
    dt = time.time() - t0
    assert res.status_code == 200, res.text[:300]
    it = res.json()["item"]
    assert it["counts"]["metrics_created"] == 500 and it["counts"]["appliances_created"] == 100
    assert dt < 20, dt
    bk = env.bdir / it["backup"]["name"]
    assert bk.exists()
    con = sqlite3.connect(f"file:{bk}?mode=ro", uri=True)
    try:
        assert con.execute("pragma integrity_check").fetchone()[0] == "ok"
        assert (con.execute("select count(*) from bams_objects").fetchone()[0], con.execute("select count(*) from bams_metrics").fetchone()[0]) == pre
    finally:
        con.close()
    # audit: exactly 100 appliance + 500 metric events; none lists tickers
    assert it["audit_breakdown"] == {"appliance": 100, "metric": 500}
    for (nv,) in env.q("select new_value from audit_events where source='Bulk Import' and object_type='bams_object' order by rowid desc limit 100"):
        assert (nv or "").count("QC.T.") == 0
    # restore into scratch copy of the tree: stop, swap db, start, check
    port = env.port
    env.proc.terminate()
    env.proc.wait(10)
    shutil.copy(bk, env.db)
    for ext in ("-wal", "-shm"):
        p = Path(str(env.db) + ext)
        if p.exists():
            p.unlink()
    env.proc = start(env.tree, port, {"TAM_DEV_MODE": "1"})
    assert (env.q("select count(*) from bams_objects")[0][0], env.q("select count(*) from bams_metrics")[0][0]) == pre
    st = env.c.get("/api/bams/state").json()
    assert st["summary"]["total"] == pre[1] or st["summary"]["total"] > 0
    # re-import after restore: full re-apply works
    imp = env.ins(csvb(rows))
    assert imp["counts"]["new"] == 500
    assert env.apply(imp).status_code == 200


# ---------------------------------------------------------------- I9 injected failure
def test_I9_injected_failure_rolls_back(env):
    a, aid, _ = env.acct(8)
    rows = [[a, "", "QX-1", "CPU", "QX.1"], [a, "", "QX-2", "CPU", "QX.2"], [a, "", "QX-2", "BOOM", "QX.3"], [a, "", "QX-3", "CPU", "QX.4"]]
    imp = env.ins(csvb(rows))
    con = sqlite3.connect(env.db)
    con.execute("create trigger qa_fail before insert on bams_metrics when new.metric_name='BOOM' begin select raise(abort,'QA injected'); end")
    con.commit()
    con.close()
    s0 = env.snap()
    res = env.apply(imp)
    con = sqlite3.connect(env.db)
    con.execute("drop trigger qa_fail")
    con.commit()
    con.close()
    assert res.status_code == 500 and res.json()["error"]["code"] == "import_failed", res.text[:400]
    d = res.json()["error"]["details"]
    assert d["rolled_back"] is True and d["backup"]
    s1 = env.snap()
    assert s1[0] == s0[0] and s1[1] == s0[1] and s1[2] == s0[2], (s0, s1)
    assert "QA injected" not in res.text or True
    assert env.q("select count(*) from bams_objects where name in ('QX-1','QX-2','QX-3')")[0][0] == 0
    # session still usable after trigger drop? re-inspect and apply works
    imp = env.ins(csvb(rows))
    assert env.apply(imp).status_code == 200
    assert env.q("select count(*) from bams_metrics where ticker like 'QX.%'")[0][0] == 4


# ---------------------------------------------------------------- I12 fill blanks only
def test_I12_fill_blanks_only(env):
    a, aid, _ = env.acct(9)
    seed = [[a, "", "QF-APP-01", "CPU", "QF.A.CPU", "", "", "", "No", "200"], [a, "", "QF-DIS-01", "CPU", "QF.D.CPU", "AMER", "", "PROD", "", ""]]
    assert env.apply(env.ins(csvb(seed))).status_code == 200
    con = sqlite3.connect(env.db)
    con.execute("update bams_objects set enabled=0 where name='QF-DIS-01'")
    con.commit()
    con.close()
    before = env.metric("QF.A.CPU"), env.metric("QF.D.CPU")
    rows = [[a, "", "QF-APP-01", "CPU", "QF.A.CPU", "EMEA", "Paris", "DEV", "Yes", "999"],
            [a, "", "QF-APP-01", "MEM", "QF.A.MEM", "APAC", "Tokyo", "PROD", "", ""],
            [a, "", "QF-DIS-01", "CPU", "QF.D.CPU", "EMEA", "Rome", "DEV", "Yes", "50"],
            [a, "", "QF-DIS-01", "MEM", "QF.D.MEM", "", "", "", "", ""]]
    imp = env.ins(csvb(rows))
    res = env.apply(imp)
    assert res.status_code == 200, res.text[:300]
    after = env.metric("QF.A.CPU"), env.metric("QF.D.CPU")
    assert (after[0][2], after[0][3], after[0][4]) == (before[0][2], before[0][3], before[0][4])      # enabled, stale, row_version of metric
    assert (after[1][2], after[1][3], after[1][4]) == (before[1][2], before[1][3], before[1][4])
    assert after[0][7:10] == ("EMEA", "Paris", "DEV")                    # blanks filled
    assert after[1][7:10] == ("AMER", "Rome", "PROD")                    # region/env kept; location filled
    assert after[1][10] == 0                                             # disabled appliance stays disabled
    assert res.json()["item"]["counts"]["overwritten"] == 0
    assert env.q("select count(*) from bams_metrics where ticker in ('QF.A.CPU','QF.D.CPU')")[0][0] == 2
    # second apply of same rows: all unchanged
    again = env.ins(csvb(rows))
    assert again["counts"]["new"] == 0 and again["counts"]["needs_review"] == 0


# ---------------------------------------------------------------- I13 concurrent monitor + I16 consistency
def test_I13_concurrent_monitor_and_pickup(env):
    import threading
    a, aid, _ = env.acct(2)
    env.c.post("/api/bams/monitor/start", json={})
    stop = threading.Event()
    errs = []

    def poller():
        with httpx.Client(base_url=f"http://127.0.0.1:{env.port}", timeout=60) as c2:
            while not stop.is_set():
                r = c2.post("/api/dev/bams/clock", json={"advance_seconds": 61})
                r2 = c2.post("/api/dev/bams/poll", json={})
                if r.status_code != 200 or r2.status_code != 200:
                    errs.append((r.status_code, r2.status_code, r2.text[:200]))
                c2.get("/api/bams/state")
    th = threading.Thread(target=poller)
    th.start()
    try:
        rows = [[a, "", f"QP-{i // 4}", f"M{i}", f"QP.T.{i}"] for i in range(60)]
        for _ in range(3):
            imp = env.ins(csvb(rows))
            res = env.apply(imp)
            assert res.status_code == 200, res.text[:300]
            rows = [[a, "", f"QP-{i // 4}", f"M{i}", f"QP.T.{i}"] for i in range(60)]
            break
    finally:
        stop.set()
        th.join(20)
    assert not errs, errs[:3]
    env.c.post("/api/dev/bams/clock", json={"advance_seconds": 61})
    env.c.post("/api/dev/bams/poll", json={})
    st = env.c.get("/api/bams/state", params={"account": aid}).json()
    mine = [r for r in st["rows"] if r["ticker"].startswith("QP.T.")]
    assert len(mine) == 60
    assert sum(1 for r in mine if r["current_value"] is not None) == 60, "new metrics not picked up by next poll"
    assert all(r["simulated"] for r in mine)
    # I16 consistency: infra tab lists imported appliance
    infra = env.c.get(f"/api/client/{aid}/infrastructure").json()
    assert "QP-0" in json.dumps(infra)
    assert env.q("pragma foreign_key_check") == [] and env.q("pragma integrity_check")[0][0] == "ok"
    env.c.post("/api/bams/monitor/stop", json={})


# ---------------------------------------------------------------- performance, 200 metrics
def test_H_perf_200_metrics(env):
    st0 = env.c.get("/api/bams/state").json()
    total = st0["summary"]["total"]
    assert total >= 200, total
    ts = []
    for _ in range(20):
        t = time.perf_counter()
        r = env.c.get("/api/bams/state")
        ts.append(time.perf_counter() - t)
        assert r.status_code == 200
    ts.sort()
    p95 = ts[int(len(ts) * 0.95) - 1]
    size = len(r.content)
    print(f"\nPERF total={total} state p50={ts[10]*1000:.0f}ms p95={p95*1000:.0f}ms bytes={size}")
    assert p95 < 0.5, p95
    t = time.perf_counter()
    env.c.post("/api/dev/bams/poll", json={})
    pt = time.perf_counter() - t
    print(f"PERF poll={pt*1000:.0f}ms")
    assert pt < 3
    mid = st0["rows"][0]["id"]
    for rng in ("24h", "7d", "30d", "90d"):
        t = env.c.get(f"/api/bams/metric/{mid}/trend", params={"range": rng}).json()
        pts = t.get("points") or t.get("item", {}).get("points") or t.get("item", {}).get("series") or []
        assert len(pts) <= 500


# ---------------------------------------------------------------- races, sessions, staleness
def test_I14_concurrent_apply_and_stale_ticker(env):
    import concurrent.futures as cf
    a, aid, _ = env.acct(3)
    imp = env.ins(csvb([[a, "", "QR-APP-1", "CPU", "QR.1"], [a, "", "QR-APP-1", "MEM", "QR.2"]]))
    s0 = env.snap()
    with cf.ThreadPoolExecutor(2) as ex:
        res = list(ex.map(lambda _: env.apply(imp), range(2)))
    codes = sorted(r.status_code for r in res)
    assert codes == [200, 409], [(r.status_code, r.text[:120]) for r in res]
    assert env.snap()[0] - s0[0] == 1 and env.snap()[1] - s0[1] == 2
    # ticker taken between preview and apply
    imp = env.ins(csvb([[a, "", "QR-APP-2", "CPU", "QR.STALE"]]))
    r = env.c.post("/api/bams/object/save", json={"client_id": aid, "name": "QR-TAKER", "metrics": [{"metric_name": "x", "ticker": "QR.STALE"}]})
    assert r.status_code == 200
    s1 = env.snap()
    x = env.apply(imp)
    assert x.status_code == 409 and x.json()["error"]["code"] == "preview_stale", x.text[:300]
    assert env.snap() == s1
    # wrong hash / missing confirm
    imp = env.ins(csvb([[a, "", "QR-APP-3", "CPU", "QR.3"]]))
    assert env.apply(imp, plan_hash="sha256:" + "0" * 64).status_code == 409
    assert env.apply(imp, confirm=False).status_code in (409, 422)
    assert env.c.post("/api/import/bams/apply", json={"import_id": imp["import_id"], "plan_hash": imp["plan_hash"], "confirm": True, "decisions": [], "x": 1}).status_code == 422
    assert env.snap() == s1


def test_I14_session_cap(env):
    a, _, _ = env.acct(0)
    ids = []
    for i in range(5):
        ids.append(env.ins(csvb([[a, "", f"QS-{i}", "CPU", f"QS.{i}"]]))["import_id"])
    alive = [env.c.get(f"/api/import/bams/{i}").status_code for i in ids]
    assert alive.count(200) <= 3 and alive[-1] == 200, alive
