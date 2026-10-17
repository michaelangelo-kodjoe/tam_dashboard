"""Restore from file: sessions, Inspect, session views, Apply, skipped download (D-126..D-135).
Routes call these functions only. Existing rows are never changed or deleted."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import logging
import sqlite3
import threading
import time
import uuid
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional
from urllib.parse import unquote

from ..audit import write_audit
from ..dbkit import migrate, transaction
from ..errors import (AlreadyApplied, Conflict, ConfirmationRequired, ImportExpired, ImportFailed,
                      NotFound, PayloadTooLarge, PreviewStale, ValidationFailed)
from ..importer import reader as R
from ..importer.reader import _bad, detect_format, display_file_name
from . import planner as P
from . import readers as RD
from . import tables as T

log = logging.getLogger("tam.restore")
HOOKS: Dict[str, Callable[..., None]] = {}               # test seams (failure injection), never set in production
SESSION_TTL = 10 * 60
RESULT_TTL = 60 * 60
MAX_SESSIONS = 3
CSV_NOTICE = "CSV files cannot tell which database they came from; check the source."


def _hook(name: str, *a: Any) -> None:
    fn = HOOKS.get(name)
    if fn:
        fn(*a)


@dataclass
class Staged:
    file_id: str
    name: str
    format: str
    size: int
    table: Optional[str]
    table_source: str                        # sheets | headers | filename | chosen | none
    headers: List[str]
    headers_cf: List[str]
    rows: List[RD.RawRow]
    encoding: Optional[str] = None
    delimiter: Optional[str] = None
    unknown: List[str] = field(default_factory=list)

    @property
    def missing(self) -> List[str]:
        if self.format == "xlsx" or self.table is None:
            return []
        return RD.missing_columns(T.BY_KEY[self.table], self.headers_cf)


@dataclass
class Session:
    id: str
    kind: str                                # xlsx | csv
    now: float
    files: List[Staged] = field(default_factory=list)
    about: Optional[Dict[str, Any]] = None
    warnings: List[Dict[str, str]] = field(default_factory=list)
    touched: float = 0.0
    applied: Optional[Dict[str, Any]] = None
    applied_at: Optional[float] = None
    skipped: List[P.Result] = field(default_factory=list)
    expires: str = ""


class Store:
    def __init__(self, clock: Callable[[], float] = time.monotonic) -> None:
        self.now = clock
        self._lock = threading.Lock()
        self._s: Dict[str, Session] = {}
        self.apply_lock = threading.Lock()

    def _sweep(self) -> None:
        t = self.now()
        for k, s in list(self._s.items()):
            if t - (s.applied_at if s.applied else s.touched) > (RESULT_TTL if s.applied else SESSION_TTL):
                del self._s[k]

    def add(self, s: Session) -> None:
        with self._lock:
            self._sweep()
            open_ = sorted((x for x in self._s.values() if not x.applied), key=lambda x: x.touched)
            while len(open_) >= MAX_SESSIONS:
                del self._s[open_.pop(0).id]
            self._s[s.id] = s

    def get(self, sid: str) -> Optional[Session]:
        with self._lock:
            self._sweep()
            s = self._s.get(sid)
            if s is not None and not s.applied:
                s.touched = self.now()
            return s

    def drop(self, sid: str) -> None:
        with self._lock:
            self._s.pop(sid, None)


def store(rt: Any) -> Store:
    if getattr(rt, "restore_sessions", None) is None:
        rt.restore_sessions = Store()
    return rt.restore_sessions


def _iso(ts: Optional[_dt.datetime] = None) -> str:
    t = ts or _dt.datetime.now(_dt.timezone.utc)
    return t.strftime("%Y-%m-%dT%H:%M:%S.") + f"{t.microsecond // 1000:03d}Z"


def _session(rt: Any, sid: str) -> Session:
    s = store(rt).get(sid)
    if s is None:
        raise ImportExpired("This restore session has expired or is unknown. Upload the file again.")
    return s


# ----------------------------------------------------------------------------------------------- inspect
def _check_about(rt: Any, about: Dict[str, str], conn: sqlite3.Connection) -> None:
    if about.get("app") != T.APP_NAME:
        raise _bad("not_an_export", "This workbook is not an export from this app.")
    try:
        ver = int(about.get("schema_version", ""))
    except ValueError:
        raise _bad("not_an_export", "This workbook is not an export from this app (no schema version).") from None
    ch = about.get("channel", "").upper()
    if ch not in ("DEV", "PROD"):
        raise _bad("not_an_export", "This workbook is not an export from this app (no channel).")
    cur = migrate.current_version(conn)
    if ver > cur:
        raise _bad("schema_newer", "This export was made by a newer version of the app. Update the app, then try again.",
                   file_schema=ver, app_schema=cur)
    if ver < cur:
        raise _bad("schema_older", "This export was made by an older version of the app and cannot be restored here. "
                                   "Restore it in the version that made it.", file_schema=ver, app_schema=cur)
    if ch == "DEV" and rt.config.mode == "PROD":
        raise _bad("channel_mismatch", "Sample data cannot be restored into the production database.")


def inspect(rt: Any, data: bytes, raw_name: Optional[str], session_id: Optional[str] = None,
            table: Optional[str] = None) -> Dict[str, Any]:
    name = display_file_name(unquote(raw_name or ""))
    if len(data) > RD.MAX_FILE_BYTES:
        raise PayloadTooLarge(f"The file is larger than {RD.MAX_FILE_BYTES // (1024 * 1024)} MB.",
                              details={"max_bytes": RD.MAX_FILE_BYTES, "reason": "too_large"})
    if not data:
        raise _bad("corrupt", "The file is empty.")
    if table is not None and table not in T.BY_KEY:
        raise ValidationFailed("Request validation failed.", fields={"table": "unknown table"})
    fmt = detect_format(data, name)
    st = store(rt)
    s: Optional[Session] = None
    if session_id:
        s = _session(rt, session_id)
        if s.applied:
            raise AlreadyApplied("This restore was already applied.", details={"result": s.applied})
        if s.kind != fmt:
            raise _bad("mixed_formats", "You cannot mix an Excel file and CSV files in one restore. Remove the other "
                                        "files first, or start again.")
        if fmt == "xlsx":
            raise _bad("xlsx_single", "Only one Excel file can be restored at a time. Remove the current file first.")
        if len(s.files) >= RD.MAX_CSV_FILES:
            raise _bad("too_many_files", f"At most {RD.MAX_CSV_FILES} CSV files (one per table) can be added.")
    new_files: List[Staged] = []
    about = None
    warnings: List[Dict[str, str]] = []
    with rt.db() as conn:
        if fmt == "xlsx":
            wb = RD.read_workbook(data)
            _check_about(rt, wb.about, conn)
            about = {"app": wb.about.get("app"), "version": wb.about.get("version"),
                     "schema_version": int(wb.about["schema_version"]), "channel": wb.about["channel"].upper(),
                     "exported_at": wb.about.get("exported_at")}
            warnings = wb.warnings
            wb_id = uuid.uuid4().hex[:12]                       # all sheets of one workbook share one file_id (QA-7-2)
            for k, sh in wb.sheets.items():
                new_files.append(Staged(wb_id, name, "xlsx", len(data), k, "sheets", sh.headers,
                                        [h.strip().casefold() for h in sh.headers], sh.rows, unknown=sh.unknown))
        else:
            cf = RD.read_csv_file(data)
            key, source = table, "chosen"
            if key is None:
                key = RD.detect_table(cf.headers_cf, name)
                source = "headers" if key else "none"
                if key and RD.missing_columns(T.BY_KEY[key], cf.headers_cf) == [] and not _headers_win(cf, key):
                    source = "filename"
            t = T.BY_KEY.get(key) if key else None
            unknown = [h for h, c in zip(cf.headers, cf.headers_cf) if c and (t is None or c not in t.by_header())] if t else []
            new_files.append(Staged(uuid.uuid4().hex[:12], name, "csv", len(data), key, source, cf.headers, cf.headers_cf,
                                    cf.rows, cf.encoding, cf.delimiter, unknown))
            if key and s is not None and any(f.table == key for f in s.files):
                raise _bad("table_already_added", f"{T.BY_KEY[key].sheet} is already added; remove it first.", table=key)
        base = (s.files if s else [])
        if sum(len(f.rows) for f in base + new_files) > RD.MAX_TOTAL_ROWS:
            raise _bad("too_many_rows", f"More than {RD.MAX_TOTAL_ROWS} rows in total.", limit=RD.MAX_TOTAL_ROWS)
        if s is None:
            s = Session(uuid.uuid4().hex, fmt, st.now(), touched=st.now(), about=about, warnings=warnings)
            s.files = new_files
            st.add(s)
        else:
            s.files.extend(new_files)
        if about and s.about is None:
            s.about = about
        log.info("restore inspect: file=%s format=%s size=%d rows=%d", name, fmt, len(data), sum(len(f.rows) for f in new_files))
        return view(rt, s, conn)


def _headers_win(cf: RD.CsvFile, key: str) -> bool:
    return RD.detect_table(cf.headers_cf, "") == key


# ----------------------------------------------------------------------------------------------- views
def _staged_map(s: Session) -> Dict[str, Any]:
    out = {}
    for f in s.files:
        if f.table and not f.missing:
            out[f.table] = (f.file_id, f.rows)
    return out


def plan_for(s: Session, conn: sqlite3.Connection) -> P.Plan:
    return P.build_plan(_staged_map(s), conn, csv_parent_hint=(s.kind == "csv"))


def view(rt: Any, s: Session, conn: sqlite3.Connection, *, table: Optional[str] = None, outcome: Optional[str] = None,
         limit: int = 200, offset: int = 0) -> Dict[str, Any]:
    plan = plan_for(s, conn)
    tot = plan.totals
    exp = _dt.datetime.now(_dt.timezone.utc) + _dt.timedelta(seconds=SESSION_TTL)
    blocking: List[Dict[str, Any]] = []
    files = []
    seen_wb: Dict[str, Dict[str, Any]] = {}
    for f in s.files:
        if f.format == "xlsx":                       # one entry for the workbook, its sheets listed inside
            e = seen_wb.get(f.file_id)
            sheet = {"table": f.table, "label": T.BY_KEY[f.table].sheet, "rows": len(f.rows)}
            if e is None:
                e = {"file_id": f.file_id, "name": f.name, "format": "xlsx", "size_bytes": f.size, "table": None,
                     "table_label": None, "table_source": "sheets", "state": "ready", "missing_columns": [], "rows": 0,
                     "encoding": None, "delimiter": None, "ignored_columns": [], "sheets": []}
                seen_wb[f.file_id] = e
                files.append(e)
            e["sheets"].append(sheet)
            e["rows"] += len(f.rows)
            e["ignored_columns"] += f.unknown
            continue
        state = "ready"
        if f.table is None:
            state = "unrecognized"
            blocking.append({"code": "unrecognized_file", "file_id": f.file_id,
                             "text": f'"{f.name}" could not be matched to a table. Choose the table for this file.'})
        elif f.missing:
            state = "missing_columns"
            blocking.append({"code": "missing_columns", "file_id": f.file_id,
                             "text": f'"{f.name}" does not fit {T.BY_KEY[f.table].sheet}: missing column(s) '
                                     + ", ".join(f.missing) + "."})
        files.append({"file_id": f.file_id, "name": f.name, "format": f.format, "size_bytes": f.size, "table": f.table,
                      "table_label": T.BY_KEY[f.table].sheet if f.table else None, "table_source": f.table_source,
                      "state": state, "missing_columns": f.missing, "rows": len(f.rows), "encoding": f.encoding,
                      "delimiter": f.delimiter, "ignored_columns": f.unknown})
    tables = []
    for t in T.TABLES:
        c = plan.counts[t.key]
        staged = any(f.table == t.key for f in s.files)
        tables.append({"key": t.key, "label": t.sheet, "staged": staged, **c, "rows": sum(c.values())})
    wanted = [r for r in plan.results if r.outcome in ("differs", "invalid")
              and (table is None or r.table == table) and (outcome is None or r.outcome == outcome)]
    issues = [_issue(r) for r in wanted[offset:offset + limit]]
    can_apply = bool(s.files) and not blocking and tot["new"] > 0
    return {"restore_id": s.id, "plan_hash": plan.plan_hash, "expires_at": _iso(exp), "kind": s.kind, "source": s.about,
            "csv_notice": CSV_NOTICE if s.kind == "csv" else None,
            "limits": {"max_file_bytes": RD.MAX_FILE_BYTES, "max_sheet_rows": RD.MAX_SHEET_ROWS,
                       "max_total_rows": RD.MAX_TOTAL_ROWS, "max_csv_files": RD.MAX_CSV_FILES},
            "files": files, "tables": tables, "totals": tot, "blocking": blocking, "warnings": s.warnings,
            "can_apply": can_apply, "nothing_to_restore": bool(s.files) and not blocking and tot["new"] == 0,
            "apply_label": f"Add {tot['new']} missing rows" if tot["new"] != 1 else "Add 1 missing row",
            "issues": issues, "issues_total": len(wanted)}


def _issue(r: P.Result) -> Dict[str, Any]:
    t = T.BY_KEY[r.table]
    p = r.problems[0] if r.problems else {"column": "", "code": r.outcome, "text": ""}
    return {"table": r.table, "table_label": t.sheet, "row": r.row, "outcome": r.outcome, "column": p["column"] or None,
            "code": p["code"], "text": f"{t.sheet} row {r.row}: {p['text']}" if p["text"] else f"{t.sheet} row {r.row}",
            "problems": len(r.problems), "file_id": r.file_id}


def get_plan(rt: Any, sid: str, table: Optional[str] = None, outcome: Optional[str] = None, limit: int = 200,
             offset: int = 0) -> Dict[str, Any]:
    s = _session(rt, sid)
    if s.applied:
        raise AlreadyApplied("This restore was already applied.", details={"result": s.applied})
    with rt.db() as conn:
        return view(rt, s, conn, table=table, outcome=outcome, limit=limit, offset=offset)


def remove(rt: Any, sid: str, file_id: str) -> Dict[str, Any]:
    s = _session(rt, sid)
    if s.applied:
        raise AlreadyApplied("This restore was already applied.", details={"result": s.applied})
    if not any(f.file_id == file_id for f in s.files):
        raise NotFound("That file is not part of this restore.")
    s.files = [f for f in s.files if f.file_id != file_id]           # removes every sheet of a workbook
    if s.kind == "xlsx" and not any(f.format == "xlsx" for f in s.files):
        s.about = None
    with rt.db() as conn:
        return view(rt, s, conn)


def set_table(rt: Any, sid: str, file_id: str, table: str) -> Dict[str, Any]:
    s = _session(rt, sid)
    if s.applied:
        raise AlreadyApplied("This restore was already applied.", details={"result": s.applied})
    f = next((x for x in s.files if x.file_id == file_id), None)
    if f is None:
        raise NotFound("That file is not part of this restore.")
    if s.kind != "csv":
        raise ValidationFailed("Request validation failed.", fields={"file_id": "the table of an Excel sheet cannot be changed"})
    if table not in T.BY_KEY:
        raise ValidationFailed("Request validation failed.", fields={"table": "unknown table"})
    if any(x.table == table and x.file_id != file_id for x in s.files):
        raise _bad("table_already_added", f"{T.BY_KEY[table].sheet} is already added; remove it first.", table=table)
    f.table, f.table_source = table, "chosen"
    t = T.BY_KEY[table]
    f.unknown = [h for h, c in zip(f.headers, f.headers_cf) if c and c not in t.by_header()]
    with rt.db() as conn:
        return view(rt, s, conn)


# ----------------------------------------------------------------------------------------------- apply
def _insert(conn: sqlite3.Connection, t: T.Table, vals: Dict[str, Any]) -> None:
    cols = [c for c, v in vals.items() if v is not None]
    sql = f"INSERT INTO {t.db} ({', '.join(cols)}) VALUES ({', '.join(':' + c for c in cols)})"  # nosec B608 - fixed identifiers
    conn.execute(sql, {c: vals[c] for c in cols})


def apply(rt: Any, restore_id: str, plan_hash: str, confirm: bool) -> Dict[str, Any]:
    if confirm is not True:
        raise ConfirmationRequired('Restoring needs confirmation. Resend with "confirm": true.')
    st = store(rt)
    s = _session(rt, restore_id)
    if not st.apply_lock.acquire(blocking=False):
        raise Conflict("Another restore is being applied. Try again in a moment.", details={"reason": "apply_in_progress"})
    try:
        if s.applied:
            raise AlreadyApplied("This restore was already applied.", details={"result": s.applied})
        with rt.db() as conn:
            plan = plan_for(s, conn)
            if plan.plan_hash != plan_hash:
                raise PreviewStale("The database changed since the preview. Preview again.", details={"why": "plan_hash"})
            blockers = [f for f in s.files if f.table is None or f.missing]
            if blockers or not s.files:
                raise Conflict("Choose the table for every file before restoring.", details={"reason": "unrecognized_file"})
            if plan.totals["new"] == 0:
                raise Conflict("There is nothing to restore: every row is already in the database or was skipped.",
                               details={"reason": "nothing_to_restore"})
            backup = rt.backups.create(conn, "pre-import")
            _hook("after_backup", rt, plan)
            n_written = 0
            audit_n = 0
            try:
                with transaction(conn):
                    again = plan_for(s, conn)
                    if again.plan_hash != plan.plan_hash:
                        raise PreviewStale("The database changed while restoring. Nothing was written.",
                                           details={"why": "plan_hash_in_transaction"})
                    for t in T.TABLES:
                        rows = plan.new_rows(t.key)
                        for r in rows:
                            _insert(conn, t, r.values or {})
                            n_written += 1
                            _hook("after_write", n_written, t.key)
                        if t.key in plan.counts and (rows or any(plan.counts[t.key].values())):
                            c = plan.counts[t.key]
                            names = ", ".join(sorted({f.name for f in s.files if f.table == t.key}))[:500]
                            write_audit(conn, category="data", object_type="import", object_id=t.db,
                                        label=f"Restore: {t.sheet}", action="restored",
                                        new={"added": c["new"], "identical": c["identical"], "skipped_differs": c["differs"],
                                             "skipped_invalid": c["invalid"]},
                                        reason=f"Restore from file: {names}"[:2000], source="Bulk Import")
                            audit_n += 1
            except (PreviewStale, Conflict):
                raise
            except Exception as exc:  # noqa: BLE001 - rolled back by the transaction; report plainly
                log.exception("restore failed and was rolled back")
                raise ImportFailed("The restore could not be completed and nothing was written.",
                                   details={"rolled_back": True, "backup": backup["name"],
                                            "reason": "A database error stopped the restore; no changes were kept."}) from exc
        skipped = [r for r in plan.results if r.outcome in ("differs", "invalid")]
        s.skipped = skipped
        result = {
            "restore_id": s.id, "applied_at": _iso(), "kind": s.kind, "files": list(dict.fromkeys(f.name for f in s.files)),
            "counts": {"added": plan.totals["new"], "identical": plan.totals["identical"], "differs": plan.totals["differs"],
                       "invalid": plan.totals["invalid"], "skipped": len(skipped)},
            "tables": [{"key": t.key, "label": t.sheet, "added": plan.counts[t.key]["new"],
                        "identical": plan.counts[t.key]["identical"], "differs": plan.counts[t.key]["differs"],
                        "invalid": plan.counts[t.key]["invalid"]} for t in T.TABLES if any(f.table == t.key for f in s.files)],
            "backup": {"name": backup["name"], "reason": "pre-import", "size_bytes": backup["size_bytes"]},
            "audit_events": audit_n, "skipped_url_xlsx": f"/api/restore/{s.id}/skipped.xlsx" if skipped else None,
            "skipped_url_csv": f"/api/restore/{s.id}/skipped.csv" if skipped else None,
            "skipped_preview": [_issue(r) for r in skipped[:200]], "skipped_total": len(skipped),
        }
        s.applied, s.applied_at = result, st.now()
        log.info("restore applied: added=%d skipped=%d backup=%s", result["counts"]["added"], len(skipped), backup["name"])
        return {"item": result, "side_effects": [{"type": "backup_created", "file": backup["name"], "reason": "pre-import"}]}
    finally:
        st.apply_lock.release()


# ----------------------------------------------------------------------------------------------- skipped
def skipped_file(rt: Any, sid: str, fmt: str) -> Dict[str, Any]:
    from .export import guard
    s = _session(rt, sid)
    if not s.applied:
        raise Conflict("This restore has not been applied yet.")
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M")
    extra = ["source_row", "outcome", "reason"]
    if fmt == "csv":
        out = io.StringIO(newline="")
        w = csv.writer(out, lineterminator="\r\n")
        w.writerow(["table", "source_row", "outcome", "reason", "values"])
        for r in s.skipped:
            detail = "; ".join(f"{k}={v}" for k, v in r.raw.items() if v)[:30000]
            w.writerow([guard(x) for x in (T.BY_KEY[r.table].sheet, str(r.row), r.outcome,
                                           r.problems[0]["text"] if r.problems else "", detail)])
        return {"data": b"\xef\xbb\xbf" + out.getvalue().encode("utf-8"), "media": "text/csv; charset=utf-8",
                "filename": f"tam_restore_skipped_{stamp}.csv"}
    openpyxl = R._load_openpyxl()
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for t in T.TABLES:
        rows = [r for r in s.skipped if r.table == t.key]
        if not rows:
            continue
        ws = wb.create_sheet(t.sheet)
        heads = [c.header for c in t.cols] + extra
        ws.append(heads)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            line = [guard(r.raw.get(c.header.casefold(), "")) for c in t.cols] + \
                   [str(r.row), r.outcome, guard(r.problems[0]["text"] if r.problems else "")]
            ws.append(line)
        for rowc in ws.iter_rows(min_row=2):
            for cell in rowc:
                cell.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return {"data": buf.getvalue(), "media": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "filename": f"tam_restore_skipped_{stamp}.xlsx"}
