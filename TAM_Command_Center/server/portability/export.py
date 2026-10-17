"""Excel export and database copy (D-123..D-125). Read-only; no audit event, one log line."""
from __future__ import annotations

import datetime as _dt
import io
import logging
import sqlite3
import uuid
from pathlib import Path
from typing import Any, Dict, List

from ..dbkit import migrate
from ..errors import BackupFailed, PayloadTooLarge
from ..importer.reader import _load_openpyxl
from . import tables as T

log = logging.getLogger("tam.export")
MAX_EXPORT_ROWS = 50_000
_GUARD = ("=", "+", "-", "@", "\t", "\r")
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def guard(s: str) -> str:
    """D-124: text a spreadsheet could evaluate gets a leading apostrophe (the reader removes exactly one)."""
    rest = s.lstrip("'")
    if rest and rest[0] in _GUARD:
        return "'" + s
    return s


def unguard(s: str) -> str:
    if s[:1] == "'":
        rest = s.lstrip("'")
        if rest and rest[0] in _GUARD:
            return s[1:]
    return s


# --- display columns (account / appliance names) are looked up here, never typed
_ACCOUNT_NAMES = "SELECT id, name FROM clients"
_APPLIANCE_NAMES = "SELECT o.id, o.name, c.name AS account FROM bams_objects o JOIN clients c ON c.id = o.client_id"


def _cell_out(kind: str, v: Any) -> Any:
    if v is None:
        return None
    if kind == "bool":
        return "Yes" if v else "No"
    if kind == "int":
        return int(v)
    return guard(str(v))


def collect(conn: sqlite3.Connection) -> Dict[str, List[List[Any]]]:
    acc = {r[0]: r[1] for r in conn.execute(_ACCOUNT_NAMES)}
    app = {r[0]: (r[1], r[2]) for r in conn.execute(_APPLIANCE_NAMES)}
    out: Dict[str, List[List[Any]]] = {}
    for t in T.TABLES:
        cols = [c.db for c in t.data_cols]
        sql = f"SELECT {', '.join(cols)}, row_version FROM {t.db} ORDER BY {t.order_by}"  # nosec B608 - fixed identifiers
        rows: List[List[Any]] = []
        for r in conn.execute(sql):
            d = dict(zip(cols + ["row_version"], r))
            line: List[Any] = []
            for c in t.cols:
                if c.db:
                    line.append(_cell_out(c.kind, d[c.db]))
                elif c.header == "row_version":
                    line.append(d["row_version"])
                elif c.header == "account_name":
                    cid = d.get("client_id")
                    if cid:
                        line.append(guard(acc.get(cid, "")))
                    elif t.key == "metrics":
                        line.append(guard(app.get(d["object_id"], ("", ""))[1]))
                    else:
                        line.append(None)
                elif c.header == "appliance_name":
                    line.append(guard(app.get(d["object_id"], ("", ""))[0]))
                else:
                    line.append(None)
            rows.append(line)
        out[t.sheet] = rows
    return out


def build_xlsx(conn: sqlite3.Connection, *, version: str, channel: str, now: _dt.datetime) -> bytes:
    conn.execute("BEGIN")                                 # one read snapshot for all sheets
    try:
        data = collect(conn)
        schema = migrate.current_version(conn)
    finally:
        conn.execute("COMMIT")
    total = sum(len(v) for v in data.values())
    if total > MAX_EXPORT_ROWS:
        raise PayloadTooLarge(f"This database has {total} rows; the Excel export is limited to {MAX_EXPORT_ROWS}. "
                              "Use the database copy instead.", details={"reason": "too_many_rows", "found": total,
                                                                         "limit": MAX_EXPORT_ROWS})
    openpyxl = _load_openpyxl()
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    about = wb.active
    about.title = T.ABOUT_SHEET
    inst = now.strftime("%Y-%m-%dT%H:%M:%S.") + f"{now.microsecond // 1000:03d}Z"
    about.append(["key", "value"])
    for k, v in (("app", T.APP_NAME), ("version", version), ("schema_version", schema), ("channel", channel.upper()),
                 ("exported_at", inst)):
        about.append([k, v])
    for t in T.TABLES:
        about.append([f"rows: {t.sheet}", len(data[t.sheet])])
    for c in about[1]:
        c.font = Font(bold=True)
    about.column_dimensions["A"].width = 30
    about.column_dimensions["B"].width = 40
    for t in T.TABLES:
        ws = wb.create_sheet(t.sheet)
        ws.append([c.header for c in t.cols])
        for c in ws[1]:
            c.font = Font(bold=True)
        for line in data[t.sheet]:
            ws.append(line)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
        for i in range(len(t.cols)):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = 22
    wb.properties.creator = T.APP_NAME
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_filename(now_local: _dt.datetime, mode: str) -> str:
    return f"tam_export_{now_local.strftime('%Y%m%d-%H%M')}_{mode}.xlsx"


def database_copy(rt: Any) -> Dict[str, Any]:
    """Online backup to a temp file in the backup folder, integrity check, return its path for streaming."""
    tmp = rt.config.backup_dir / f"tam_download_{uuid.uuid4().hex}.tmp"
    with rt.db() as conn:
        try:
            version = migrate.current_version(conn)
            migrate.online_backup(conn, tmp)             # sqlite backup API + integrity_check on the copy
        except Exception as exc:
            log.error("database copy failed: %s", exc)
            _rm(tmp)
            raise BackupFailed("The database copy could not be created. Nothing was downloaded.") from exc
    stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M")
    log.info("database copy prepared: %d bytes", tmp.stat().st_size)
    return {"path": tmp, "filename": f"tam_backup_{stamp}_v{version}.db"}


def _rm(p: Path) -> None:
    try:
        p.unlink()
    except OSError:
        pass
