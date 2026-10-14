"""Builders for import tests: well-formed workbooks/CSVs and hostile variants, all generated in code."""
from __future__ import annotations

import io
import zipfile
from typing import Any, Dict, Iterable, List, Optional, Sequence

import openpyxl

HEAD = ["Account name", "BPID", "Products", "Status", "Next contact date", "Notes"]


def xlsx(rows: Iterable[Sequence[Any]], headers: Optional[Sequence[str]] = HEAD, sheet: str = "Accounts",
         formulas: Sequence[str] = (), padding: int = 0) -> bytes:
    """rows of up to 6 values. Strings are stored as strings (even if they start with '='); cells listed in
    `formulas` (e.g. "A3") are written as real formulas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    r0 = 1
    if headers is not None:
        for c, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c, value=h)
        r0 = 2
    n = 0
    for n, row in enumerate(rows):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r0 + n, column=c, value=v)
            if isinstance(v, str):
                cell.data_type = "s"
    for ref in formulas:
        ws[ref].data_type = "f"
    if padding:
        ws.cell(row=r0 + n + padding, column=1).number_format = "@"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def csv_bytes(text: str, enc: str = "utf-8") -> bytes:
    return text.encode(enc)


def rezip(data: bytes, add: Optional[Dict[str, bytes]] = None, replace: Optional[Dict[str, bytes]] = None,
          remove: Sequence[str] = ()) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            if i.filename in remove:
                continue
            zout.writestr(i.filename, (replace or {}).get(i.filename, zin.read(i.filename)))
        for k, v in (add or {}).items():
            zout.writestr(k, v)
    return out.getvalue()


def sheet_part(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return next(n for n in z.namelist() if n.startswith("xl/worksheets/sheet"))


def upload(api, data: bytes, name: str = "accounts.xlsx"):
    return api.c.post("/api/import/accounts/inspect", content=data,
                      headers={"X-File-Name": name, "Content-Type": "application/octet-stream"})
