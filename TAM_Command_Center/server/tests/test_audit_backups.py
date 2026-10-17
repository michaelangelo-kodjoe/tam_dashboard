"""Audit log API, notes-only Activity, weekly backups and retention (contract section 18)."""
from __future__ import annotations

import csv
import datetime as dt
import io

import openpyxl
import pytest

from server import constants as C
from server.backups import kind_of

UTC = dt.timezone.utc


# ------------------------------------------------------------------------------------------ activity
def test_activity_is_notes_only_and_audit_has_the_system_rows(build):
    api = build()
    c = api.mk_client("Acme Capital")
    t = api.mk_task(c["id"], "Send report")
    api.ok("/api/task/save", {"id": t["id"], "status": "Complete"})
    api.ok("/api/contact/save", {"client_id": c["id"], "name": "Pat"})
    n = api.ok("/api/note/save", {"client_id": c["id"], "note_type": "General Note", "subject": "Hello"})["item"]
    items = api.get(f"/api/client/{c['id']}/activity").json()["items"]
    assert [i["id"] for i in items] == [n["id"]] and all(i["kind"] == "note" for i in items)
    sysv = api.get(f"/api/client/{c['id']}/activity", filter="system").json()
    assert sysv["items"] == [] and sysv["has_more"] is False
    a = api.get("/api/audit", client_id=c["id"], object_type="task").json()
    assert {i["action"] for i in a["items"]} >= {"created", "completed"}
    assert any(i["summary"] == "Task completed: Send report" and i["client_name"] == "Acme Capital" for i in a["items"])
    today = api.get("/api/today").json()["recent_changes"]
    assert any(r["object_type"] == "task" for r in today)            # Today "Recently changed" is unchanged


# ------------------------------------------------------------------------------------------ audit log
def seed_audit(api):
    c1, c2 = api.mk_client("Alpha Co")["id"], api.mk_client("Beta Co")["id"]
    rows = [("2026-10-01T10:00:00.000Z", "crm", c1, "task", "completed", "GUI", "Task A", "done early"),
            ("2026-10-02T10:00:00.000Z", "crm", c1, "contact", "created", "GUI", "Pat", None),
            ("2026-10-03T10:00:00.000Z", "bams", c2, "bams_object", "created", "Bulk Import", "BPOD-1", "file x"),
            ("2026-10-04T10:00:00.000Z", "settings", None, "setting", "updated", "GUI", "log_level", None),
            ("2026-10-05T10:00:00.000Z", "data", None, "import", "restored", "Bulk Import", "Restore: Accounts", None)]
    for ts, cat, cid, ot, act, src, label, reason in rows:
        api.exec("INSERT INTO audit_events (timestamp, category, client_id, object_type, object_id, object_label, action,"
                 " reason, source) VALUES (?,?,?,?,?,?,?,?,?)", (ts, cat, cid, ot, "o1", label, act, reason, src))
    return c1, c2


def test_filters_paging_order_and_summaries(build):
    api = build()
    c1, c2 = seed_audit(api)
    j = api.get("/api/audit", limit=3).json()
    ts = [i["timestamp"] for i in j["items"]]
    assert ts == sorted(ts, reverse=True) and j["total"] >= 7 and j["has_more"] and j["limit"] == 3
    page2 = api.get("/api/audit", limit=3, offset=3).json()["items"]
    assert not ({i["id"] for i in j["items"]} & {i["id"] for i in page2})
    f = lambda **k: api.get("/api/audit", **k).json()  # noqa: E731
    assert [i["object_label"] for i in f(category="settings")["items"]] == ["log_level"]
    assert f(source="Bulk Import")["total"] == 2
    assert {i["object_label"] for i in f(client_id=c1, category="crm", object_type="task")["items"]} == {"Task A"}
    assert f(action="restored")["items"][0]["summary"] == "Restore: Accounts (restored)"
    r = f(**{"from": "2026-10-02", "to": "2026-10-03"})
    assert [i["object_label"] for i in r["items"] if i["timestamp"].startswith("2026-10")] == ["BPOD-1", "Pat"]
    assert f(**{"to": "2026-10-01"})["total"] >= 1 and all(i["timestamp"] < "2026-10-02" for i in f(**{"to": "2026-10-01"})["items"])
    assert [i["object_label"] for i in f(q="EARLY")["items"]] == ["Task A"]                 # reason, case-insensitive
    assert {i["object_label"] for i in f(q="beta")["items"]} >= {"BPOD-1"}                    # account name
    assert f(q="100%")["total"] == 0                                                         # LIKE wildcards are escaped
    first = f(category="bams")["items"][0]
    assert first["client_name"] == "Beta Co" and first["summary"].startswith("BAMS object created")
    assert set(first) >= {"id", "timestamp", "category", "client_id", "object_type", "object_id", "object_label", "action",
                          "source", "reason", "summary", "old_value", "new_value"}


def test_validation_caps_readonly_and_facets(build):
    api = build()
    seed_audit(api)
    for bad in ({"from": "yesterday"}, {"to": "2026-13-01"}, {"category": "nope"}, {"source": "Bot"}, {"limit": 201},
                {"limit": 0}, {"offset": -1}, {"offset": 100001}, {"q": "x" * 201}):
        assert api.get("/api/audit", **bad).status_code == 422, bad
    assert api.c.post("/api/audit", json={}).status_code == 405 and api.c.delete("/api/audit").status_code in (404, 405)
    n = api.sql("SELECT count(*) n FROM audit_events")[0]["n"]
    api.get("/api/audit")
    api.get("/api/audit/export", format="csv")
    assert api.sql("SELECT count(*) n FROM audit_events")[0]["n"] == n                       # reading writes nothing
    fac = api.get("/api/audit/facets").json()
    assert "crm" in fac["categories"] and "task" in fac["object_types"] and "restored" in fac["actions"] and "GUI" in fac["sources"]


def test_deleted_account_keeps_the_row_with_null_name(build):
    api = build()
    c = api.mk_client("Gone Co")["id"]
    api.exec("INSERT INTO audit_events (category, client_id, object_type, object_id, object_label, action, source)"
             " VALUES ('crm', ?, 'task', 't1', 'Orphan task', 'created', 'GUI')", (c,))
    api.exec("DELETE FROM clients WHERE id = ?", (c,))
    row = api.get("/api/audit", q="Orphan").json()["items"][0]
    assert row["client_name"] is None and row["summary"] == "Task created: Orphan task"


def test_export_csv_and_xlsx_are_guarded_and_filtered(build):
    api = build()
    api.exec("INSERT INTO audit_events (category, object_type, object_id, object_label, action, source, reason)"
             " VALUES ('settings', 'setting', 'k', '=cmd|calc', 'updated', 'GUI', '+SUM(1)')")
    r = api.get("/api/audit/export", format="csv", category="settings")
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store" and ".csv" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["id", "timestamp", "category"] and len(rows) == 2
    assert "'=cmd|calc" in rows[1] and "'+SUM(1)" in rows[1]
    x = api.get("/api/audit/export", category="settings")
    ws = openpyxl.load_workbook(io.BytesIO(x.content))["Audit log"]
    assert ws.max_row == 2 and all(c.data_type != "f" for row in ws.iter_rows() for c in row)
    assert api.get("/api/audit/export", format="pdf").status_code == 422


def test_export_cap(build, monkeypatch):
    from server import auditlog
    api = build()
    seed_audit(api)
    monkeypatch.setattr(auditlog, "EXPORT_MAX", 2)
    r = api.get("/api/audit/export", format="csv")
    assert r.status_code == 413 and r.json()["error"]["details"]["reason"] == "too_many_rows"


# ------------------------------------------------------------------------------------------ backups
@pytest.fixture
def bk(build):
    api = build()
    api.mk_client("Data Co")                                         # the weekly backup needs data
    clock = {"t": dt.datetime(2026, 10, 1, 8, 0, tzinfo=UTC)}
    api.rt.backups.clock = lambda: clock["t"]
    api.clock = clock
    api.after = lambda **kw: clock.update(t=clock["t"] + dt.timedelta(**kw))
    return api


def names(api, kind=None):
    return [b["name"] for b in api.rt.backups.list(500) if kind is None or b["kind"] == kind]


def make(api, reason):
    with api.rt.db() as conn:
        return api.rt.backups.create(conn, reason)["name"]


def test_kinds_and_list_order(bk):
    assert [kind_of(r) for r in ("weekly", "manual", "pre-import", "pre-restore", "pre-delete")] == \
        ["weekly", "manual", "safety", "safety", "safety"]
    make(bk, "manual"); bk.after(minutes=1); make(bk, "pre-delete"); bk.after(minutes=1); bk.rt.backups.run_scheduled(bk.rt.db)
    items = bk.get("/api/data/backups").json()["items"]
    assert [i["kind"] for i in items] == ["weekly", "safety", "manual"]          # newest first, grouped by the UI
    assert [i["name"] for i in items] == sorted((i["name"] for i in items), reverse=True)
    assert bk.ok("/api/data/backup")["item"]["kind"] == "manual"


def test_weekly_cadence_with_injected_clock(bk):
    svc = bk.rt.backups
    assert svc.weekly_due() and svc.run_scheduled(bk.rt.db)["reason"] == "weekly"
    for hours in (1, 24, 24 * 5):                                    # within 7 days: nothing new
        bk.after(hours=hours)
        assert svc.run_scheduled(bk.rt.db) is None
    assert len(names(bk, "weekly")) == 1
    bk.after(days=2)                                                 # 7+ days since the first
    assert svc.run_scheduled(bk.rt.db)["kind"] == "weekly" and len(names(bk, "weekly")) == 2
    assert svc.run_scheduled(bk.rt.db) is None                        # never twice in one week


def test_weekly_skipped_for_an_empty_database_and_failure_is_swallowed(build, monkeypatch):
    api = build()
    assert api.rt.backups.run_scheduled(api.rt.db) is None and not api.rt.backups.list()
    api.mk_client("X")
    monkeypatch.setattr(api.rt.backups, "create", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("disk full")))
    assert api.rt.backups.run_scheduled(api.rt.db) is None                    # logged, not raised


def test_weekly_keeps_only_the_newest_eight(bk):
    for _ in range(11):
        bk.rt.backups.run_scheduled(bk.rt.db)
        bk.after(days=7)
    w = names(bk, "weekly")
    assert len(w) == C.WEEKLY_KEEP == 8 and w == sorted(w, reverse=True)


def test_safety_backups_keep_five_per_reason_and_thirty_days(bk):
    for _ in range(8):
        make(bk, "pre-import"); bk.after(minutes=1)
    for _ in range(3):
        make(bk, "pre-delete"); bk.after(minutes=1)
    assert len([n for n in names(bk) if "_pre-import_" in n]) == 5 and len([n for n in names(bk) if "_pre-delete_" in n]) == 3
    bk.after(days=31)                                                 # everything older than 30 days goes...
    newest = make(bk, "pre-import")
    left = names(bk)
    assert newest in left and not [n for n in left if "_pre-delete_" in n]
    assert len([n for n in left if "_pre-import_" in n]) == 1        # ...only today's pre-import survives


def test_pre_restore_and_manual_are_never_pruned_and_newest_is_kept(bk):
    keep = [make(bk, "pre-restore"), make(bk, "manual")]
    bk.after(days=400)
    for _ in range(12):
        make(bk, "pre-delete"); bk.after(minutes=1)
    assert set(keep) <= set(names(bk))
    bk.after(days=100)
    bk.rt.backups.prune()
    left = names(bk)
    assert set(keep) <= set(left) and len([n for n in left if "_pre-delete_" in n]) == 1       # newest of any kind kept
    assert left[0].count("_pre-delete_") == 1                                                   # and it is the newest overall


def test_newest_backup_survives_even_when_it_is_old_and_alone(bk):
    only = make(bk, "pre-import")
    bk.after(days=90)
    assert bk.rt.backups.prune() == [] and names(bk) == [only]


def test_restore_flow_and_listing_still_work_with_new_retention(bk):
    snap = bk.get("/api/export/database").content
    bk.mk_client("Later Co")
    item = bk.c.post("/api/restore/database/inspect", content=snap, headers={"X-File-Name": "s.db"}).json()["item"]
    bk.ok("/api/restore/database/apply", {"restore_id": item["restore_id"], "confirm": "RESTORE"})
    got = bk.get("/api/data/backups").json()["items"]
    assert got[0]["reason"] == "pre-restore" and got[0]["kind"] == "safety"
    back = bk.post("/api/restore/database/from-backup", {"name": got[0]["name"]})
    assert back.status_code == 200
