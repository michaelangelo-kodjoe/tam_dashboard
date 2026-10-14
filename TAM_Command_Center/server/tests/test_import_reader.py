"""T-82/T-83: template + hardened reader. Every refusal reason has a generated hostile file."""
from __future__ import annotations

import io
import re
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

from server import constants as C
from server.errors import InvalidFile, PayloadTooLarge
from server.importer import reader, template
from server.tests import importkit as K

ROOT = Path(__file__).resolve().parents[2]


def refused(data, name="a.xlsx"):
    with pytest.raises((InvalidFile, PayloadTooLarge)) as ei:
        reader.read_upload(data, name)
    return ei.value


def reason(data, name="a.xlsx"):
    return refused(data, name).details["reason"]


GOOD = K.xlsx([["Acme Ltd", "00123", "P", "active", "2026-11-05", "n"]])


# ------------------------------------------------------------------ template (T-82)
def test_status_constant_pinned_to_schema_check():
    sql = (ROOT / "db" / "schema.sql").read_text(encoding="utf-8")
    m = re.search(r"CREATE TABLE IF NOT EXISTS clients \((.*?)\n\);", sql, re.S)
    chk = re.search(r"status\s+TEXT NOT NULL DEFAULT 'Active' CHECK \(status IN \(([^)]*)\)\)", m.group(1))
    assert tuple(re.findall(r"'([^']+)'", chk.group(1))) == C.CLIENT_STATUSES == ("Active", "Inactive")
    from server.schemas import CLIENT_STATUSES
    assert CLIENT_STATUSES is C.CLIENT_STATUSES


def test_template_workbook_content():
    wb = openpyxl.load_workbook(io.BytesIO(template.build_template_xlsx()))
    assert wb.sheetnames == ["Instructions", "Accounts", "Lists"]
    assert wb["Lists"].sheet_state == "hidden" and wb["Accounts"].sheet_state == "visible"
    ws = wb["Accounts"]
    assert [c.value for c in ws[1]] == ["Account name", "BPID", "Products", "Status", "Next contact date", "Notes"]
    assert [wb["Lists"].cell(row=i, column=1).value for i in (1, 2)] == list(C.CLIENT_STATUSES)
    dvs = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    assert dvs["D2:D501"].type == "list" and dvs["D2:D501"].formula1 == "=Lists!$A$1:$A$2"
    assert dvs["E2:E501"].type == "date" and "YYYY-MM-DD" in dvs["E2:E501"].prompt
    assert dvs["A2:A501"].type == "textLength"
    assert ws["B2"].number_format == "@" and ws["B501"].number_format == "@" and ws["E2"].number_format == "yyyy-mm-dd"
    for sheet in wb:      # no data rows, no formulas anywhere
        for row in sheet.iter_rows():
            for c in row:
                assert c.data_type != "f"
    assert ws.max_row >= 501 and all(ws.cell(row=r, column=1).value is None for r in range(2, 502))


def test_template_has_no_macro_or_links_and_reimport_reports_no_rows():
    data = template.build_template_xlsx()
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        assert not any("vba" in n.lower() or "external" in n.lower() or "embed" in n.lower() for n in names)
        assert b"TargetMode" not in b"".join(z.read(n) for n in names if n.endswith(".rels"))
    assert reason(data) == "no_rows"
    assert reason(template.build_template_csv(), "t.csv") == "no_rows"


def test_template_csv_bytes():
    b = template.build_template_csv()
    assert b.startswith(b"\xef\xbb\xbf") and b.decode("utf-8-sig") == "Account name,BPID,Products,Status,Next contact date,Notes\r\n"


def test_pip_audit_pin_present():
    req = (ROOT / "server" / "requirements.txt").read_text()
    assert re.search(r"^openpyxl==3\.1\.5$", req, re.M) and re.search(r"^et_xmlfile==2\.0\.0$", req, re.M)


def test_app_starts_and_other_routes_work_without_openpyxl(tmp_path):
    code = r'''
import sys, builtins
real = builtins.__import__
def blocked(name, *a, **k):
    if name.split(".")[0] in ("openpyxl", "et_xmlfile"):
        raise ImportError("blocked for test")
    return real(name, *a, **k)
builtins.__import__ = blocked
sys.path.insert(0, %r)
from fastapi.testclient import TestClient
from server.tests.conftest import _build
api, app = _build(__import__("pathlib").Path(%r))
assert api.get("/api/health").json()["ok"] is True
r = api.c.get("/api/import/accounts/template")
assert r.status_code == 500 and r.json()["error"]["code"] == "internal", r.text
r = api.c.get("/api/import/accounts/template", params={"format": "csv"})
assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")
r = api.c.post("/api/import/accounts/inspect", content=b"Account name\r\nAcme\r\n", headers={"X-File-Name": "a.csv"})
assert r.status_code == 200 and r.json()["item"]["file"]["format"] == "csv", r.text
app.state.runtime.close()
print("OK")
''' % (str(ROOT), str(tmp_path))
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=str(ROOT))
    assert out.returncode == 0 and "OK" in out.stdout, out.stdout + out.stderr


# ------------------------------------------------------------------ happy path + normalisation
def test_reads_xlsx_rows_in_any_column_order_and_normalises():
    data = K.xlsx([["  Acme\t\nCorp  ", "12", "x", "INACTIVE", "2026-01-31", " line1\r\nline2 "]],
                  headers=["Notes", "Next contact date", "Status", "Products", "BPID", "Account name*"])
    # headers reordered: reuse values in header order
    data = K.xlsx([[" line1\r\nline2 ", "2026-01-31", "INACTIVE", "x", 12, "  Acme\t\nCorp  "]],
                  headers=["Notes", "Next contact date", "Status", "Products", "BPID", "Account name*"])
    p = reader.read_upload(data, "a.xlsx")
    v = p.rows[0].values
    assert v == {"name": "Acme Corp", "bpid": "12", "products": "x", "status": "Inactive",
                 "next_contact": "2026-01-31", "notes": "line1\nline2"}
    assert p.format == "xlsx" and p.sheet == "Accounts" and p.rows_found == 1 and p.rows[0].row == 2


def test_numbers_dates_blank_rows_and_row_numbers_match_excel():
    import datetime as dt
    data = K.xlsx([["A1", 123456.0, None, None, dt.datetime(2026, 11, 5, 13, 30), None],
                   [None] * 6, ["  "], ["B2", "007", None, None, None, None]])
    p = reader.read_upload(data, "a.xlsx")
    assert [r.row for r in p.rows] == [2, 5]
    assert p.rows[0].values["bpid"] == "123456" and p.rows[0].values["next_contact"] == "2026-11-05"
    assert p.rows[1].values["bpid"] == "007" and p.rows_found == 2


@pytest.mark.parametrize("cell,code", [
    ("x" * 201, "name_too_long"),
])
def test_name_limits_200_ok_201_invalid(cell, code):
    ok = reader.read_upload(K.xlsx([["n" * 200]]), "a.xlsx").rows[0]
    assert ok.valid and len(ok.values["name"]) == 200
    bad = reader.read_upload(K.xlsx([[cell]]), "a.xlsx").rows[0]
    assert not bad.valid and bad.errors[0]["code"] == code and bad.values["name"] is None


def test_row_errors_name_row_field_and_fix():
    rows = [["", "x", "", "", "", ""], ["N", "b" * 33], ["N2", None, "p" * 501], ["N3", None, None, "Dormant"],
            ["N4", None, None, None, "11/05/2026"], ["N5", None, None, None, "2026-02-30"], ["N6", None, None, None, 45000],
            ["N7", 1.5], ["N8", None, None, None, None, "n" * 20001], ["N9", True]]
    p = reader.read_upload(K.xlsx(rows), "a.xlsx")
    codes = [r.errors[0]["code"] for r in p.rows]
    assert codes == ["name_required", "bpid_too_long", "products_too_long", "status_not_allowed", "date_invalid",
                     "date_invalid", "date_invalid", "bpid_invalid", "notes_too_long", "value_type"]
    t = p.rows[3].errors[0]["text"]
    assert t.startswith("Row 5: Status:") and "Active, Inactive" in t and "skipped" in t
    assert "YYYY-MM-DD" in p.rows[4].errors[0]["text"]


def test_header_rules():
    assert reason(K.xlsx([["a"]], headers=["Name only"])) == "header_missing"
    assert reason(K.xlsx([["a"]], headers=["Account name", "bpid", "BPID"])) == "header_duplicate"
    assert reason(K.xlsx([["a"]], sheet="Other")) == "sheet_missing"
    assert refused(K.xlsx([["a"]], sheet="Other")).details["sheets"] == ["Other"]
    p = reader.read_upload(K.xlsx([["a", None, None, None, None, None, "x"]], headers=HEAD7 if False else K.HEAD + ["Region"]),
                           "a.xlsx")
    assert p.warnings == [{"code": "unknown_column", "text": 'Column G "Region" is ignored.'}]
    assert reader.read_upload(K.xlsx([["a"]], sheet="accounts "), "a.xlsx").rows_found == 1     # case-insensitive


# ------------------------------------------------------------------ formulas
def test_formula_rows_are_invalid_and_never_evaluated():
    data = K.xlsx([["Good"], ["=1+1"], ["Real formula"], ["=SUM(1,2)", None, "=\"abc\"", None, None, "=A2"]],
                  formulas=["A3", "A4", "A5", "C5"])
    p = reader.read_upload(data, "a.xlsx")
    by = {r.row: r for r in p.rows}
    assert by[2].valid
    assert by[3].values["name"] is None and by[3].errors[0]["code"] == "formula_not_allowed"   # =1+1 as a formula
    assert by[4].errors[0]["code"] == "formula_not_allowed"
    assert by[5].errors[0]["code"] == "formula_not_allowed" and len(by[5].errors) >= 2
    # text that merely starts with = + - @ is data
    t = K.xlsx([["=1+1"], ["+cmd"], ["-5"], ["@SUM(A1)"]])
    vals = [r.values["name"] for r in reader.read_upload(t, "a.xlsx").rows]
    assert vals == ["=1+1", "+cmd", "-5", "@SUM(A1)"]


def test_cached_formula_value_is_never_used():
    """A formula whose cached value would be a valid name stays invalid: data_only is never used."""
    base = K.xlsx([["x"]])
    part = K.sheet_part(base)
    xml = zipfile.ZipFile(io.BytesIO(base)).read(part).decode()
    patched = re.sub(r'<c r="A2"[^>]*>.*?</c>', '<c r="A2" t="str"><f>1+1</f><v>Cached Valid Name</v></c>', xml, flags=re.S)
    assert "Cached Valid Name" in patched
    data = K.rezip(base, replace={part: patched.encode()})
    r = reader.read_upload(data, "a.xlsx").rows[0]
    assert not r.valid and r.errors[0]["code"] == "formula_not_allowed" and r.values["name"] is None


def test_formulas_outside_data_columns_are_ignored():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Accounts"
    for i, h in enumerate(K.HEAD, start=1):
        ws.cell(row=1, column=i, value=h)
    ws["A2"] = "Ok"; ws["J2"] = "=1+1"
    ws2 = wb.create_sheet("Instructions"); ws2["A1"] = "=NOW()"
    b = io.BytesIO(); wb.save(b)
    p = reader.read_upload(b.getvalue(), "a.xlsx")
    assert p.rows[0].valid and p.rows_found == 1


# ------------------------------------------------------------------ limits
def test_500_rows_ok_501_fails_whole_file():
    ok = reader.read_upload(K.xlsx([[f"Acct {i}"] for i in range(500)]), "a.xlsx")
    assert ok.rows_found == 500 and len(ok.rows) == 500
    e = refused(K.xlsx([[f"Acct {i}"] for i in range(501)]))
    assert e.details == {"reason": "too_many_rows", "found": 501, "limit": 500}
    assert "501 account rows" in e.message and "limit is 500" in e.message
    e = refused(K.xlsx([[f"Acct {i}"] for i in range(812)]))
    assert "812 account rows" in e.message


def test_blank_rows_do_not_count_and_padding_never_hangs():
    rows = [["A"]] + [[None]] * 600 + [["B"]]
    assert reader.read_upload(K.xlsx(rows), "a.xlsx").rows_found == 2
    t = time.perf_counter()
    e = refused(K.xlsx([["a"]] * 10, padding=20000))
    assert e.details["reason"] == "too_many_rows" and e.details.get("physical") is True
    assert time.perf_counter() - t < 5
    t = time.perf_counter()
    assert reader.read_upload(K.xlsx([["a"]] * 10, padding=9000), "a.xlsx").rows_found == 10
    assert time.perf_counter() - t < 5


def test_data_beyond_scan_limit_is_refused_even_if_dimension_lies():
    base = K.xlsx([["a"]])
    part = K.sheet_part(base)
    xml = zipfile.ZipFile(io.BytesIO(base)).read(part).decode()
    xml = xml.replace("</sheetData>", f'<row r="{C.IMPORT_SCAN_ROWS + 1}"><c r="A{C.IMPORT_SCAN_ROWS + 1}" t="inlineStr"><is><t>late</t></is></c></row></sheetData>')
    xml = re.sub(r"<dimension[^>]*/>", '<dimension ref="A1:F5"/>', xml)
    assert reason(K.rezip(base, replace={part: xml.encode()})) == "too_many_rows"


def test_huge_cell_is_an_invalid_row_not_a_crash():
    p = reader.read_upload(K.xlsx([["n" * 100000, "x"], ["ok"]]), "a.xlsx")
    assert not p.rows[0].valid and p.rows[0].errors[0]["code"] == "name_too_long" and p.rows[1].valid
    assert len(p.rows[0].raw["name"]) <= reader.RAW_KEEP


def test_control_characters_never_survive():
    # xlsx (XML) cannot carry C0 controls, but can carry bidi/zero-width/C1 characters
    p = reader.read_upload(K.xlsx([["Ac\u200bme\u202e Corp\x85", "1\u200f2", None, None, None, "n\u2028ote"]]), "a.xlsx")
    assert p.rows[0].values["name"] == "Acme Corp"
    assert p.rows[0].values["bpid"] == "12" and p.rows[0].values["notes"] == "note"
    c = reader.read_upload("Account name,BPID,Notes\r\nA\x07c\x1bme Corp,1\x0b2,n\x7fote\r\n".encode(), "a.csv")
    assert c.rows[0].values == {"name": "Acme Corp", "bpid": "12", "products": None, "status": None,
                                "next_contact": None, "notes": "note"}


# ------------------------------------------------------------------ hostile files: one per refusal reason
def test_refusal_not_xlsx_family():
    assert reason(GOOD, "a.xls") == "unsupported_type"
    assert reason(GOOD, "a.txt") == "unsupported_type"
    assert reason(GOOD, "noext") == "unsupported_type"
    assert reason(b"hello, not a zip", "a.xlsx") == "not_xlsx"
    assert reason(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600, "a.xlsx") == "not_xlsx"           # .xls renamed
    assert reason(b"Account name\n", "a.csv") == "no_rows"
    assert reason(GOOD, "a.csv") == "not_csv"                                                          # zip named .csv
    assert reason(b"\x00\x01\x02binary", "a.csv") == "not_csv"
    docx = K.rezip(GOOD, remove=("xl/workbook.xml",), add={"word/document.xml": b"<w/>"})
    assert reason(docx) == "not_xlsx"


def test_refusal_encrypted_corrupt_empty():
    enc = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100 + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 100
    assert reason(enc) == "encrypted"
    assert reason(GOOD[:len(GOOD) // 2]) == "corrupt"
    assert reason(b"PK\x03\x04" + b"\x00" * 50) == "corrupt"
    assert reason(b"") == "corrupt"
    # a ZIP whose general-purpose flag marks encryption
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        zi = zipfile.ZipInfo("xl/workbook.xml"); zi.flag_bits |= 0x1
        z.writestr(zi, b"x")
        z.writestr("[Content_Types].xml", b"x")
    raw = bytearray(buf.getvalue())
    for m in re.finditer(rb"PK\x01\x02", bytes(raw)):
        raw[m.start() + 8] |= 1                       # general-purpose flag bit 0 in the central directory
    assert reason(bytes(raw)) == "encrypted"


def test_refusal_too_large_is_413():
    e = refused(b"PK\x03\x04" + b"0" * (C.MAX_IMPORT_BYTES + 1))
    assert isinstance(e, PayloadTooLarge) and e.status == 413 and e.details["reason"] == "too_large"


def test_refusal_zip_bomb_100_to_1_a_few_kb():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", b"<Types/>")
        z.writestr("xl/workbook.xml", b"<workbook/>")
        z.writestr("xl/bomb.xml", b"\x00" * (5 * 1024 * 1024))
    data = buf.getvalue()
    assert len(data) < 20000
    assert reason(data) == "zip_unsafe"


def test_refusal_zip_entry_total_and_name_limits():
    def zipped(entries):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
            for n, b in entries:
                z.writestr(n, b)
        return buf.getvalue()
    many = [("[Content_Types].xml", b"x"), ("xl/workbook.xml", b"x")] + [(f"xl/p{i}.xml", b"x") for i in range(201)]
    assert reason(zipped(many)) == "zip_unsafe"
    for bad in ("../evil.xml", "/abs.xml", "xl/../../x.xml", "C:/x.xml", "xl\\evil.xml"):
        assert reason(zipped([("[Content_Types].xml", b"x"), ("xl/workbook.xml", b"x"), (bad, b"x")])) == "zip_unsafe", bad
    dup = io.BytesIO()
    with zipfile.ZipFile(dup, "w") as z:
        z.writestr("[Content_Types].xml", b"x"); z.writestr("xl/workbook.xml", b"x")
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            z.writestr("xl/workbook.xml", b"y")
    assert reason(dup.getvalue()) == "zip_unsafe"


def test_refusal_declared_size_lie_cannot_bypass(monkeypatch):
    """Over-limit entry sizes are refused from the directory without reading anything."""
    monkeypatch.setattr(reader, "MAX_ENTRY", 100)
    assert reason(GOOD) == "zip_unsafe"


def test_refusal_macros():
    assert reason(K.rezip(GOOD, add={"xl/vbaProject.bin": b"\xd0\xcf"})) == "macros"
    ct = zipfile.ZipFile(io.BytesIO(GOOD)).read("[Content_Types].xml").replace(
        b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml")
    assert reason(K.rezip(GOOD, replace={"[Content_Types].xml": ct})) == "macros"


def test_refusal_external_links_and_objects():
    assert reason(K.rezip(GOOD, add={"xl/externalLinks/externalLink1.xml": b"<externalLink/>"})) == "external_links"
    rels = b'<Relationships><Relationship Id="r1" Type="x" Target="http://evil.example/x" TargetMode="External"/></Relationships>'
    assert reason(K.rezip(GOOD, add={"xl/worksheets/_rels/sheet1.xml.rels": rels})) == "external_links"
    assert reason(K.rezip(GOOD, add={"xl/embeddings/oleObject1.bin": b"x"})) == "embedded_objects"
    assert reason(K.rezip(GOOD, add={"xl/activeX/activeX1.xml": b"<x/>"})) == "embedded_objects"
    assert reason(K.rezip(GOOD, add={"xl/foo/blob.bin": b"x"})) == "embedded_objects"
    # ordinary Excel printer settings are tolerated (documented deviation)
    ok = K.rezip(GOOD, add={"xl/printerSettings/printerSettings1.bin": b"x"})
    assert reader.read_upload(ok, "a.xlsx").rows_found == 1


@pytest.mark.parametrize("payload", [
    b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;">]><x>&lol2;</x>',
    b"<!doctype html><x/>",
    b'<!ENTITY xxe SYSTEM "file:///etc/passwd">',
    "<?xml version='1.0' encoding='UTF-16'?><!DOCTYPE x [<!ENTITY a 'b'>]><x/>".encode("utf-16-le"),
    "<!DOCTYPE x><x/>".encode("utf-16-be"),
])
def test_refusal_xml_entities_and_doctype(payload):
    for part in ("xl/workbook.xml", "xl/styles.xml", "docProps/core.xml"):
        assert reason(K.rezip(GOOD, replace={part: payload})) == "xml_unsafe", part
    assert reason(K.rezip(GOOD, add={"xl/theme/extra.xml": payload})) == "xml_unsafe"


def test_hostile_files_never_reach_openpyxl(monkeypatch):
    called = []
    monkeypatch.setattr(reader, "_load_openpyxl", lambda: called.append(1))
    for data in (K.rezip(GOOD, add={"xl/vbaProject.bin": b"x"}), K.rezip(GOOD, replace={"xl/workbook.xml": b"<!DOCTYPE a>"})):
        refused(data)
    assert called == []


# ------------------------------------------------------------------ csv
def test_csv_delimiters_encodings_and_row_numbers():
    base = "Account name;BPID;Products;Status;Next contact date;Notes\r\nAcme;00123;P;inactive;2026-11-05;line\r\n\r\nBeta;;;;;\r\n"
    p = reader.read_upload(base.encode("utf-8"), "a.csv")
    assert p.format == "csv" and p.delimiter == "semicolon" and p.encoding == "UTF-8" and p.sheet == "CSV"
    assert [r.row for r in p.rows] == [2, 4] and p.rows[0].values["bpid"] == "00123" and p.rows[0].values["status"] == "Inactive"
    assert p.notes == ["Read as CSV: encoding UTF-8, delimiter semicolon."]
    t = base.replace(";", "\t")
    assert reader.read_upload(t.encode("utf-8-sig"), "a.csv").delimiter == "tab"
    assert reader.read_upload(t.encode("utf-8-sig"), "a.csv").encoding == "UTF-8 (with BOM)"
    assert reader.read_upload(t.encode("utf-16"), "a.csv").encoding == "UTF-16"
    cp = "Account name,Notes\r\nSoci\xe9t\xe9 G\xe9n\xe9rale,caf\xe9\r\n".encode("cp1252")
    q = reader.read_upload(cp, "a.csv")
    assert q.encoding == "Windows-1252" and q.rows[0].values["name"] == "Société Générale" and q.delimiter == "comma"
    assert reader.read_upload(base.replace(";", "|").encode(), "a.csv").delimiter == "pipe"


def test_csv_quoted_fields_formula_text_and_quote_prefix_roundtrip():
    text = 'Account name,Notes\r\n"Acme, Inc.","line1\r\nline2"\r\n=1+1,@SUM(A1)\r\n\'=cmd,x\r\n'
    p = reader.read_upload(text.encode(), "a.csv")
    assert [r.values["name"] for r in p.rows] == ["Acme, Inc.", "=1+1", "=cmd"]      # text, never evaluated
    assert p.rows[0].values["notes"] == "line1\nline2" and p.rows[1].values["notes"] == "@SUM(A1)"
    assert all(r.valid for r in p.rows)


def test_csv_errors_limits_and_headers():
    assert reason(b"Name\r\nx\r\n", "a.csv") == "header_missing"
    assert reason(b"Account name,BPID,bpid\r\nx\r\n", "a.csv") == "header_duplicate"
    big = "Account name\r\n" + "".join(f"A{i}\r\n" for i in range(501))
    assert refused(big.encode(), "a.csv").details["found"] == 501
    assert reader.read_upload(("Account name\r\n" + "".join(f"A{i}\r\n" for i in range(500))).encode(), "a.csv").rows_found == 500
    w = reader.read_upload(b"Account name,Region\r\nA,EU\r\n", "a.csv").warnings
    assert w[0]["code"] == "unknown_column"
    p = reader.read_upload(b"Account name,Next contact date\r\nA,5/11/2026\r\n", "a.csv")
    assert p.rows[0].errors[0]["code"] == "date_invalid"


def test_skipped_files_are_text_only_and_roundtrip():
    rows = [["=SUM(A1)", "+1", "@x", "Active", "2026-01-01", "-5"], ["Plain", "", "", "", "", ""]]
    x = template.build_skipped_xlsx([r + ["7", "reason", ""] for r in rows])
    ws = openpyxl.load_workbook(io.BytesIO(x))["Accounts"]
    assert ws["A2"].value == "=SUM(A1)" and ws["A2"].data_type == "s" and ws["C2"].data_type == "s"
    assert [c.value for c in ws[1]][6:] == ["Source row", "Reason", "Possible duplicate of"]
    p = reader.read_upload(x, "s.xlsx")
    assert p.rows[0].values["name"] == "=SUM(A1)" and p.rows[0].valid
    assert p.warnings and {w["code"] for w in p.warnings} == {"unknown_column"} and len(p.warnings) == 3
    c = template.build_skipped_csv([r + ["7", "reason", ""] for r in rows])
    q = reader.read_upload(c, "s.csv")
    assert q.rows[0].values["name"] == "=SUM(A1)" and q.rows[0].values["bpid"] == "+1" and q.rows[0].values["notes"] == "-5"
    assert b"'=SUM(A1)" in c
