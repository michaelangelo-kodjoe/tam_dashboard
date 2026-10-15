"""Run 4 (T-123..T-125): BAMS appliance + metrics import."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import openpyxl
import pytest

from server import dupes
from server.importer import bams_reader, bams_service
from server.importer.bams_plan import find_similar_appliances
from server.tests.importkit import rezip, xlsx
from server.tests.test_bams_monitor import T0

HEAD = bams_reader.HEADERS
FIX = Path(__file__).parent / "fixtures" / "dupes_bams_golden.json"
ROWS_FIX = Path(__file__).resolve().parents[2] / "db" / "fixtures" / "bams_import_rows.json"


def R(account="Acme Capital", bpid="", appliance="BPOD-01", metric="CPU %", ticker="T.CPU", region="", location="",
      env="", enabled="", stale=""):
    return [account, bpid, appliance, metric, ticker, region, location, env, enabled, stale]


def up(api, rows, name="a.xlsx", headers=HEAD, **kw):
    if name.endswith(".csv"):
        buf = io.StringIO(newline="")
        w = csv.writer(buf)
        w.writerow(headers)
        w.writerows(rows)
        data = buf.getvalue().encode("utf-8")
    else:
        data = xlsx(rows, headers=headers, sheet="Import", **kw)
    return api.c.post("/api/import/bams/inspect", content=data, headers={"X-File-Name": name})


def inspect_ok(api, rows, **kw):
    r = up(api, rows, **kw)
    assert r.status_code == 200, r.text
    return r.json()["item"]


def apply(api, item, decisions=None, status=200, confirm=True, **extra):
    body = {"import_id": item["import_id"], "plan_hash": item["plan_hash"], "confirm": confirm,
            "decisions": decisions or [], **extra}
    r = api.post("/api/import/bams/apply", body)
    assert r.status_code == status, r.text
    return r.json()


def counts(api):
    return api.sql("SELECT (SELECT count(*) FROM bams_objects) o, (SELECT count(*) FROM bams_metrics) m,"
                   " (SELECT count(*) FROM audit_events) a")[0]


def backups(api):
    return list(api.rt.config.backup_dir.glob("*pre-import*"))


@pytest.fixture
def ia(build):
    api = build(app_kw={"utc_base": lambda: T0})
    api.cid = api.mk_client("Acme Capital")["id"]
    api.cid2 = api.mk_client("Beta Fund", bpid="777")["id"]
    return api


def mk_obj(api, name, metrics, cid=None, **kw):
    return api.ok("/api/bams/object/save", {"client_id": cid or api.cid, "name": name, "metrics": [
        {"metric_name": m, "ticker": t} for m, t in metrics], **kw})["item"]["object"]


# ------------------------------------------------------------------ template and reader
def test_template_files(api):
    r = api.get("/api/import/bams/template")
    assert r.status_code == 200 and "tam_bams_import_template.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Instructions", "Import", "Example", "Lists"] and wb["Lists"].sheet_state == "hidden"
    assert [c.value for c in wb["Import"][1]] == HEAD
    assert wb["Import"].max_row == 501 or wb["Import"].cell(row=2, column=1).value is None
    assert all(c.data_type != "f" for ws in wb for row in ws.iter_rows() for c in row)
    assert "Unit" not in HEAD
    c = api.get("/api/import/bams/template", format="csv")
    assert c.content.startswith(b"\xef\xbb\xbf") and c.content.decode("utf-8-sig").strip().split(",") == HEAD
    assert api.get("/api/import/bams/template", format="pdf").status_code == 422


def test_example_sheet_reads_back(ia):
    wb = openpyxl.load_workbook(io.BytesIO(ia.get("/api/import/bams/template").content))
    rows = [[c.value or "" for c in r] for r in wb["Example"].iter_rows(min_row=2)]
    item = inspect_ok(ia, rows)
    assert item["counts"]["invalid"] == 3 and {e["code"] for r in item["rows"] for e in r["errors"]} == {"account_not_found"}


def test_headers_formulas_unknown_columns_limits(ia):
    r = up(ia, [["x"]], headers=["Appliance", "Metric"])
    j = r.json()
    assert r.status_code == 422 and j["error"]["code"] == "invalid_file" and j["error"]["details"]["reason"] == "header_missing"
    assert j["error"]["details"]["missing"] == ["Account or BPID", "Ticker"] and j["error"]["details"]["found"] == ["Appliance", "Metric"]
    item = inspect_ok(ia, [R() + ["x"], R(appliance="B2", ticker="T2") + ["y"]], headers=HEAD + ["Owner"], formulas=["C3"])
    assert item["warnings"] == [{"code": "unknown_column", "text": 'Column K "Owner" is ignored.'}]
    assert [r["state"] for r in item["rows"]] == ["new", "invalid"] and item["rows"][1]["errors"][0]["code"] == "formula_not_allowed"
    assert inspect_ok(ia, [R(appliance=f"A{i}", ticker=f"T{i}") for i in range(500)])["counts"]["rows"] == 500
    r = up(ia, [R(appliance=f"A{i}", ticker=f"T{i}") for i in range(501)])
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "too_many_rows"
    assert up(ia, [R()], name="a.csv").status_code == 200
    bad = rezip(xlsx([R()], headers=HEAD, sheet="Import"), add={"xl/vbaProject.bin": b"x"})
    r = ia.c.post("/api/import/bams/inspect", content=bad, headers={"X-File-Name": "m.xlsx"})
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "macros"
    assert ia.c.get("/api/import/bams/" + "0" * 32).status_code == 410
    assert ia.c.get("/api/import/bams/inspect").status_code == 405


def test_account_import_session_is_separate(ia):
    item = inspect_ok(ia, [R()])
    assert ia.c.get(f"/api/import/accounts/{item['import_id']}").status_code == 410
    sk = ia.c.get(f"/api/import/bams/{item['import_id']}/skipped.xlsx")
    assert sk.status_code == 409


def test_row_validation_and_account_resolution(ia):
    mk_obj(ia, "X", [("a", "TX")], cid=ia.cid2)
    rows = [R(account="acme  CAPITAL", ticker="T1"), R(account="", bpid="777", appliance="B1", ticker="T2"),
            R(account="Nope", ticker="T3"), R(account="Acme Capital", bpid="777", ticker="T4"),
            R(env="staging", ticker="T5"), R(stale="0", ticker="T6"), R(enabled="maybe", ticker="T7"),
            R(account="", bpid="", ticker="T8"), R(env="dev", enabled="No", stale="300", ticker="T9", appliance="E9")]
    item = inspect_ok(ia, rows)
    codes = [[e["code"] for e in r["errors"]] for r in item["rows"]]
    assert codes == [[], [], ["account_not_found"], ["account_mismatch"], ["environment_invalid"], ["stale_after_invalid"],
                     ["enabled_invalid"], ["account_required"], []]
    last = item["rows"][-1]
    assert last["values"]["environment"] == "DEV" and last["values"]["enabled"] is False and last["values"]["stale_after_seconds"] == 300
    assert item["rows"][1]["target"]["client_name"] == "Beta Fund"
    assert item["counts"]["rows"] == 9 and item["counts"]["invalid"] == 6
    c = item["counts"]
    assert c["rows"] == c["new"] + c["updates"] + c["unchanged"] + c["needs_review"] + c["invalid"] and c["valid"] == c["rows"] - c["invalid"]


def test_in_file_duplicates_and_ticker_rules(ia):
    mk_obj(ia, "EXIST", [("CPU", "TE1"), ("OTH", "TE2")])
    rows = [R(ticker="D1"), R(ticker="D1"),                                   # repeats row 2
            R(appliance="B2", metric="M2", ticker="S1"), R(appliance="B3", metric="M3", ticker="s1"),   # both invalid
            R(appliance="B4", metric="M4", ticker="Q1"), R(appliance="B4", metric="M4", ticker="Q2"),   # differs in file
            R(appliance="EXIST", metric="CPU", ticker="OTHER"),                # ticker_differs vs DB
            R(appliance="EXIST", metric="NEW", ticker="te2")]                  # ticker owned by another metric
    item = inspect_ok(ia, rows)
    got = [(r["state"], [e["code"] for e in r["errors"]]) for r in item["rows"]]
    assert got == [("new", []), ("invalid", ["repeats_row"]), ("invalid", ["ticker_repeated_in_file"]),
                   ("invalid", ["ticker_repeated_in_file"]), ("new", []), ("invalid", ["ticker_differs"]),
                   ("invalid", ["ticker_differs"]), ("invalid", ["ticker_conflict"])]
    assert "repeats row 2" in item["rows"][1]["errors"][0]["text"]
    assert "already used by OTH on EXIST (Acme Capital)" in item["rows"][7]["errors"][0]["text"]


def test_same_name_other_account_is_independent(ia):
    mk_obj(ia, "SHARED-01", [("CPU", "TA")])
    item = inspect_ok(ia, [R(account="Beta Fund", appliance="shared-01", metric="CPU", ticker="TB")])
    assert item["rows"][0]["state"] == "new" and item["review"] == []


# ------------------------------------------------------------------ fixture file against the seeded DB
def test_fixture_rows_against_seed(seeded):
    d = json.loads(ROWS_FIX.read_text())
    cols = d["columns"]
    item = inspect_ok(seeded, [[r[c] for c in cols] for r in d["rows"]])
    expect = {r["_expect"]: [] for r in d["rows"]}
    for row, src in zip(item["rows"], d["rows"]):
        expect[src["_expect"]].append(row)
    st = lambda k: expect[k][0]["state"]  # noqa: E731
    assert st("new_appliance") in ("new", "invalid")        # row 2 shares its ticker with row 14: both Invalid (D-66)
    assert expect["new_appliance"][1]["state"] == "new"
    assert st("new_appliance_first_account") == "new" and st("exact_existing_noop") == "unchanged"
    assert st("new_metric_on_existing") == "new"
    assert [r["state"] for r in expect["near_duplicate_of_existing"]] == ["needs_review", "needs_review"]
    assert st("near_duplicate_in_file") == "new" and expect["near_duplicate_in_file"][1]["state"] == "needs_review"
    assert st("same_name_other_account_not_flagged") == "new" and st("inactive_account_accepted") == "new"
    for k, code in (("invalid_ticker_conflict", "ticker_conflict"), ("invalid_ticker_repeated_in_file", "ticker_repeated_in_file"),
                    ("invalid_unknown_account", "account_not_found"), ("invalid_bpid_name_mismatch", "account_mismatch"),
                    ("invalid_environment", "environment_invalid"), ("invalid_stale_after", "stale_after_invalid"),
                    ("invalid_over_limit", "appliance_too_long")):
        assert [e["code"] for e in expect[k][0]["errors"]] == [code], k
    g = [x for x in item["review"] if x["file"]["appliance"] == "scs bpipe nyc 01"][0]
    assert g["candidates"][0]["ref"]["type"] == "group" and g["strength"] == "High"
    assert item["appliances"]["needs_review"] == len(item["review"]) == 3
    before = counts_of(seeded)
    assert counts_of(seeded) == before                                   # inspect writes nothing
    assert seeded.get(f"/api/import/bams/{item['import_id']}").json()["item"]["plan_hash"] == item["plan_hash"]


def counts_of(api):
    return api.sql("SELECT (SELECT count(*) FROM bams_objects) o, (SELECT count(*) FROM bams_metrics) m,"
                   " (SELECT count(*) FROM audit_events) a")[0]


# ------------------------------------------------------------------ golden and parity (T-124)
def test_appliance_matcher_golden_and_engine_parity():
    d = json.loads(FIX.read_text())
    assert len(d["cases"]) >= 8
    for case in d["cases"]:
        mt = dupes.Matcher()
        for p in case["pool"]:
            mt.add(p["ref"], p["name"], None, data={"kind": "appliance", "obj": {"name": p["name"]}})
        ex = mt.exact(case["new"]["name"])
        assert (ex.ref if ex else None) == case["expected"]["exact"], case["id"]
        got = [] if ex else find_similar_appliances(mt, case["new"]["name"])
        assert [{"ref": c["entry"].ref, "strength": c["strength"], "reasons": [r["code"] for r in c["reasons"]],
                 "differences": [x["code"] for x in c["differences"]]} for c in got] == case["expected"]["candidates"], case["id"]
        raw = {c["entry"].ref["id"]: c for c in mt.find(case["new"]["name"], None, limit=500)["candidates"]}
        for c in got:                       # parity: identical strength word and reason sentences as the Run 3 engine
            e = raw[c["entry"].ref["id"]]
            assert c["strength"] == e["strength"] and c["reasons"] == e["reasons"]


# ------------------------------------------------------------------ Apply (T-125)
def test_apply_new_and_merge_rules_audit_backup(ia):
    disabled = mk_obj(ia, "OLD-01", [("CPU", "TO1")], region="Americas", enabled=False)
    ia.exec("UPDATE bams_metrics SET stale_after_seconds = 999, enabled = 0")
    rows = [R(appliance="OLD-01", metric="CPU", ticker="TO1", region="Europe", location="Paris", env="DEV", enabled="Yes", stale="5"),
            R(appliance="old-01", metric="MEM", ticker="TO2", enabled="No", stale="45"),
            R(appliance="NEW-01", metric="CPU", ticker="TN1", region="EMEA", location="Rome", env="prod"),
            R(appliance="NEW-01", metric="MEM", ticker="TN2", enabled="no", stale="300"),
            R(account="Nope", ticker="TX")]
    item = inspect_ok(ia, rows)
    assert item["appliances"] == {"new": 1, "updated": 1, "unchanged": 0, "needs_review": 0}
    assert item["counts"]["updates"] == 1 and item["metrics"] == {"new": 3, "updated": 1, "unchanged": 0}
    r = apply(ia, item, confirm=False, status=409)
    assert r["error"]["code"] == "confirmation_required"
    res = apply(ia, item)["item"]
    c = res["counts"]
    assert c["appliances_created"] == 1 and c["appliances_updated"] == 1 and c["metrics_created"] == 3
    assert c["metrics_updated"] == 1 and c["skipped"] == 1 and c["invalid"] == 1 and c["overwritten"] == 0
    assert res["fields_filled"] == 2 and res["backup"]["reason"] == "pre-import" and len(backups(ia)) == 1
    assert res["touched_accounts"] == [{"id": ia.cid, "name": "Acme Capital"}]
    o = ia.sql("SELECT * FROM bams_objects WHERE name = 'OLD-01'")[0]
    assert (o["region"], o["location"], o["environment"], o["enabled"]) == ("Americas", "Paris", "DEV", 0)   # fills blanks only
    m = {r["metric_name"]: r for r in ia.sql("SELECT * FROM bams_metrics WHERE object_id = ?", (o["id"],))}
    assert m["CPU"]["stale_after_seconds"] == 999 and m["CPU"]["enabled"] == 0 and m["CPU"]["ticker"] == "TO1"
    assert m["MEM"]["enabled"] == 0 and m["MEM"]["stale_after_seconds"] == 45
    nm = {r["metric_name"]: r for r in ia.sql("SELECT m.* FROM bams_metrics m JOIN bams_objects o ON o.id = m.object_id WHERE o.name = 'NEW-01'")}
    assert nm["MEM"]["enabled"] == 0 and nm["MEM"]["stale_after_seconds"] == 300 and nm["CPU"]["enabled"] == 1
    new = ia.sql("SELECT * FROM bams_objects WHERE name = 'NEW-01'")[0]
    assert (new["enabled"], new["environment"], new["region"]) == (1, "PROD", "EMEA")
    aud = ia.audit(source="Bulk Import")
    assert len(aud) == res["audit_events"] == 5 and all(a["category"] == "bams" for a in aud)
    assert res["audit_breakdown"] == {"appliance": 2, "metric": 3}
    assert all(a["reason"].startswith("Bulk import: a.xlsx") for a in aud)
    assert all("TN1" not in (a["new_value"] or "") or a["object_type"] == "bams_metric" for a in aud)
    sk = ia.c.get(res["skipped_url"])
    assert sk.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(sk.content))["Import"]
    assert [c.value for c in ws[1]] == HEAD + ["Source row", "Reason", "Possible duplicate of"] and ws["K2"].value == "6"
    again = apply(ia, item, status=409)
    assert again["error"]["code"] == "already_applied"
    stale = apply(ia, {**item, "import_id": "0" * 32}, status=410)
    assert stale["error"]["code"] == "import_expired"


def test_idempotent_reimport_writes_nothing(ia):
    rows = [R(appliance="A1", ticker="TA", region="EU"), R(appliance="A1", metric="MEM", ticker="TB")]
    apply(ia, inspect_ok(ia, rows))
    n_backups, before = len(backups(ia)), counts(ia)
    item = inspect_ok(ia, rows)
    assert item["counts"]["unchanged"] == 2 and item["appliances"]["unchanged"] == 1
    res = apply(ia, item)["item"]
    assert res["backup"] is None and res["audit_events"] == 0 and res["counts"]["metrics_unchanged"] == 2
    assert res["counts"]["appliances_created"] == 0 and res["counts"]["appliances_updated"] == 0
    assert counts(ia) == before and len(backups(ia)) == n_backups


def flagged_file(ia):
    mk_obj(ia, "BPOD-NY-01", [("CPU", "EXIST.CPU")], region="Americas")
    return [R(appliance="BPOD NY 01", metric="CPU", ticker="NEW.CPU", location="New York"),
            R(appliance="BPOD NY 01", metric="Disk", ticker="NEW.DISK", region="Europe")]


def test_decisions_gate_and_validation(ia):
    item = inspect_ok(ia, flagged_file(ia))
    assert len(item["review"]) == 1 and item["counts"]["needs_review"] == 2
    g = item["review"][0]
    assert g["key"] == "g2" and g["candidates"][0]["ref"]["type"] == "appliance" and g["strength"] == "High"
    cand = g["candidates"][0]
    assert cand["fills"] == {"location": "New York"} and cand["kept"]["region"]["differs"] is True
    assert cand["metrics_conflicts"] == [{"metric": "CPU", "ticker": "NEW.CPU", "owner": "BPOD-NY-01", "existing_ticker": "EXIST.CPU"}]
    assert cand["metrics_added"] == ["Disk"]
    b = counts(ia)
    e = apply(ia, item, status=409)["error"]
    assert e["code"] == "unresolved_duplicates" and e["details"]["keys"] == ["g2"] and e["details"]["rows"] == [2, 3]
    assert counts(ia) == b and backups(ia) == []
    tgt = {"type": "appliance", "id": cand["ref"]["id"]}
    for dec in ([{"key": "g9", "choice": "skip"}], [{"key": "g2", "choice": "skip"}, {"key": "g2", "choice": "skip"}],
                [{"key": "g2", "choice": "same"}], [{"key": "g2", "choice": "different", "target": tgt}],
                [{"key": "g2", "choice": "same", "target": {"type": "appliance", "id": "nope"}}], [{"key": "bad", "choice": "skip"}]):
        assert apply(ia, item, dec, status=422)["error"]["code"] == "validation_error"
    assert counts(ia) == b


def test_decision_same_different_skip(ia):
    item = inspect_ok(ia, flagged_file(ia))
    cand = item["review"][0]["candidates"][0]
    res = apply(ia, item, [{"key": "g2", "choice": "same", "target": {"type": "appliance", "id": cand["ref"]["id"]}}])["item"]
    # CPU hits a different existing ticker -> invalid, not blocking; Disk is added; blank location filled
    assert res["counts"]["metrics_created"] == 1 and res["counts"]["invalid"] == 1 and res["counts"]["appliances_updated"] == 1
    o = ia.sql("SELECT * FROM bams_objects")[0]
    assert (o["location"], o["region"]) == ("New York", "Americas")
    assert {r["metric_name"]: r["ticker"] for r in ia.sql("SELECT * FROM bams_metrics")} == {"CPU": "EXIST.CPU", "Disk": "NEW.DISK"}
    assert res["decisions"][0]["label"] == "BPOD NY 01 merged into BPOD-NY-01" and res["breakdown"]["merged_marked_same"] == 1
    # different
    item2 = inspect_ok(ia, [R(appliance="BPOD NY 01", metric="Net", ticker="NEW.NET")])
    r2 = apply(ia, item2, [{"key": "g2", "choice": "different"}])["item"]
    assert r2["counts"]["appliances_created"] == 1 and r2["breakdown"]["created_marked_different"] == 1
    assert ia.sql("SELECT count(*) AS n FROM bams_objects WHERE name = 'BPOD NY 01'")[0]["n"] == 1
    # skip
    item3 = inspect_ok(ia, [R(appliance="BPOD-NY-1", metric="Q", ticker="NEW.Q")])
    key = item3["review"][0]["key"]
    n = counts(ia)
    r3 = apply(ia, item3, [{"key": key, "choice": "skip"}])["item"]
    assert r3["counts"]["skipped"] == 1 and r3["breakdown"]["skipped_by_user"] == 1 and counts(ia) == n and r3["backup"] is None
    assert "Skipped by you (possible duplicate of" in r3["skipped"][0]["reason"]
    sk = ia.c.get(r3["skipped_url"].replace(".xlsx", ".csv"))
    assert sk.status_code == 404


def test_same_onto_group_target_and_cross_group_skip(ia):
    rows = [R(appliance="scs-bpipe-nyc-01", metric="CPU", ticker="G1"), R(appliance="scs bpipe nyc 01", metric="MEM", ticker="G2")]
    item = inspect_ok(ia, rows)
    assert item["rows"][0]["state"] == "new" and item["review"][0]["candidates"][0]["ref"] == {"type": "group", "key": "g2"}
    res = apply(ia, item, [{"key": "g3", "choice": "same", "target": {"type": "group", "key": "g2"}}])["item"]
    assert res["counts"]["appliances_created"] == 1 and res["counts"]["metrics_created"] == 2
    assert ia.sql("SELECT count(*) AS n FROM bams_objects")[0]["n"] == 1
    item = inspect_ok(ia, [R(appliance="yy-unit-b", metric="A", ticker="Y1"), R(appliance="yy unit b", metric="B", ticker="Y2")])
    assert item["review"][0]["key"] == "g3"
    res = apply(ia, item, [{"key": "g3", "choice": "same", "target": {"type": "group", "key": "g2"}}])["item"]
    assert res["counts"]["metrics_created"] == 2 and res["counts"]["appliances_created"] == 1
    item = inspect_ok(ia, [R(appliance="ww-unit-c", metric="A", ticker="W1"), R(appliance="ww unit c", metric="B", ticker="W2")])
    e = apply(ia, item, [{"key": "g3", "choice": "skip"}])["item"]
    assert e["counts"]["skipped"] == 1 and e["counts"]["appliances_created"] == 1


def test_atomic_failure_injection_rolls_back_everything(ia):
    item = inspect_ok(ia, [R(appliance=f"U{i}", ticker=f"TT{i}") for i in range(4)])
    before = counts(ia)

    def boom(n, dest):
        if n == 3:
            raise RuntimeError("injected")
    bams_service.HOOKS["after_write"] = boom
    try:
        e = apply(ia, item, status=500)["error"]
    finally:
        bams_service.HOOKS.clear()
    assert e["code"] == "import_failed" and e["details"]["rolled_back"] is True and e["details"]["row"] >= 2
    assert e["details"]["backup"].endswith(".db") and "injected" not in json.dumps(e)
    assert counts(ia) == before
    ok = apply(ia, item)["item"]                       # the session survives a failed apply
    assert ok["counts"]["appliances_created"] == 4


def test_preview_stale_when_ticker_taken_between_preview_and_apply(ia):
    item = inspect_ok(ia, [R(appliance="N1", ticker="RACE")])
    mk_obj(ia, "OTHER", [("x", "RACE")])
    e = apply(ia, item, status=409)["error"]
    assert e["code"] == "preview_stale" and e["details"]["why"] == "ticker_conflict"
    assert ia.sql("SELECT count(*) AS n FROM bams_objects WHERE name = 'N1'")[0]["n"] == 0
    fresh = inspect_ok(ia, [R(appliance="N1", ticker="RACE")])
    assert fresh["rows"][0]["state"] == "invalid"


def test_preview_stale_when_data_changes_in_transaction(ia):
    item = inspect_ok(ia, [R(appliance="N2", ticker="T2")])
    hits = []

    def sneaky(rt, plan):
        hits.append(1)
        with rt.db() as conn:
            conn.execute("INSERT INTO bams_objects (id, client_id, name) VALUES ('sneak', ?, 'N2')", (ia.cid,))
    bams_service.HOOKS["after_backup"] = sneaky
    try:
        e = apply(ia, item, status=409)["error"]
    finally:
        bams_service.HOOKS.clear()
    assert hits and e["code"] == "preview_stale" and ia.sql("SELECT count(*) AS n FROM bams_metrics")[0]["n"] == 0


def test_imported_metric_polled_by_running_monitor_without_restart(ia):
    ia.ok("/api/bams/monitor/start", {})
    ia.ok("/api/dev/bams/poll", {})
    res = apply(ia, inspect_ok(ia, [R(appliance="P1", metric="CPU", ticker="POLL.ME")]))["item"]
    row = [r for r in ia.get("/api/bams/state").json()["rows"] if r["ticker"] == "POLL.ME"][0]
    assert row["freshness"] == "No Data" and res["touched_accounts"]
    ia.ok("/api/dev/bams/clock", {"advance_seconds": 60})
    ia.ok("/api/dev/bams/poll", {})
    row = [r for r in ia.get("/api/bams/state").json()["rows"] if r["ticker"] == "POLL.ME"][0]
    assert row["freshness"] == "Live" and row["current_value"] is not None
    assert len(ia.get("/api/bams/state", account=ia.cid).json()["rows"]) == 1


def test_csv_skipped_file_is_formula_safe(ia):
    rows = [R(appliance="=cmd()", ticker="=T1", account="Nope")]
    item = inspect_ok(ia, rows, name="x.csv")
    res = apply(ia, item)["item"]
    sk = ia.c.get(res["skipped_url"])
    assert sk.status_code == 200 and sk.headers["content-disposition"].startswith('attachment; filename="tam_bams_skipped_')
    text = sk.content.decode("utf-8-sig")
    assert "'=cmd()" in text and "'=T1" in text
    again = inspect_ok(ia, [r for r in csv.reader(io.StringIO(text))][1:], headers=next(csv.reader(io.StringIO(text))), name="y.csv")
    assert any(w["code"] == "unknown_column" for w in again["warnings"])


def test_prod_import_allowed_but_no_numbers(build):
    api = build(prod=True)
    cid = api.mk_client("P")["id"]
    res = apply(api, inspect_ok(api, [R(account="P", appliance="PA", ticker="PT")]))["item"]
    assert res["counts"]["metrics_created"] == 1
    s = api.get("/api/bams/state").json()
    assert s["rows"][0]["freshness"] == "No Data" and s["rows"][0]["current_value"] is None and s["banner"]
