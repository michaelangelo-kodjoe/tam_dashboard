"""T-84/T-85/T-86: Inspect/Preview, Apply, skipped download, check-similar and client/save enforcement."""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3

import openpyxl
import pytest

from server import constants as C
from server.importer import service as svc
from server.tests import importkit as K
from server.tests.conftest import add_route  # noqa: F401  (kept for parity with other suites)


@pytest.fixture(autouse=True)
def _clean_hooks():
    svc.HOOKS.clear()
    yield
    svc.HOOKS.clear()


def insp(api, rows, name="accounts.xlsx", **kw):
    r = K.upload(api, K.xlsx(rows, **kw), name)
    assert r.status_code == 200, r.text
    return r.json()["item"]


def apply_(api, item, decisions=(), status=200, confirm=True, plan_hash=None):
    r = api.c.post("/api/import/accounts/apply", json={"import_id": item["import_id"], "confirm": confirm,
                                                     "plan_hash": plan_hash or item["plan_hash"], "decisions": list(decisions)})
    assert r.status_code == status, r.text
    return r.json()


def err(api, item, decisions, status, code, **kw):
    j = apply_(api, item, decisions, status=status, **kw)
    assert j["ok"] is False and j["error"]["code"] == code, j
    return j["error"]


def snapshot(api):
    """Everything an import may touch, to prove 'nothing written'."""
    return {"clients": api.sql("SELECT * FROM clients ORDER BY id"), "audit": api.sql("SELECT * FROM audit_events ORDER BY id"),
            "backups": sorted(p.name for p in api.rt.config.backup_dir.iterdir())}


def by_row(item):
    return {r["row"]: r for r in item["rows"]}


def same(row, c):
    return {"row": row, "choice": "same", "target": c["ref"]}


def acct(api, name):
    return api.sql("SELECT * FROM clients WHERE name = ?", (name,))[0]


# ======================================================================= template route
def test_template_routes_and_headers(api):
    r = api.c.get("/api/import/accounts/template")
    assert r.status_code == 200 and r.headers["content-type"] == template_media()
    assert 'filename="tam_accounts_import_template.xlsx"' in r.headers["content-disposition"]
    assert r.headers["cache-control"] == "no-store" and r.headers["x-content-type-options"] == "nosniff"
    assert "default-src 'self'" in r.headers["content-security-policy"]
    assert openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames == ["Instructions", "Accounts", "Lists"]
    c = api.c.get("/api/import/accounts/template", params={"format": "csv"})
    assert c.status_code == 200 and c.headers["content-type"].startswith("text/csv")
    assert 'tam_accounts_import_template.csv' in c.headers["content-disposition"]
    b = api.c.get("/api/import/accounts/template", params={"format": "xls"})
    assert b.status_code == 422 and b.json()["error"]["code"] == "validation_error"


def template_media():
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_route_methods_and_stubs(api):
    assert api.c.get("/api/import/accounts/inspect").status_code == 405
    assert api.c.get("/api/import/accounts/apply").status_code == 405
    assert api.c.get("/api/import/accounts/" + "0" * 32).status_code == 410
    assert api.c.post("/api/import/other").status_code == 501 and api.c.get("/api/import").status_code == 501
    assert api.c.get("/api/import/accounts/" + "0" * 32 + "/skipped.xlsx").status_code == 410
    assert api.c.get("/api/import/accounts/" + "0" * 32 + "/skipped.txt").status_code == 404
    spec = api.c.get("/openapi.json").json()["paths"]
    for p in ("/api/import/accounts/template", "/api/import/accounts/inspect", "/api/import/accounts/apply",
              "/api/import/accounts/{import_id}", "/api/client/check-similar"):
        assert p in spec


def test_cors_preflight_allows_file_name_header(build):
    a = build(extra_env={"TAM_CORS_ORIGINS": "http://127.0.0.1:5173"}) if False else None
    from server.app import create_app  # noqa: F401
    # default DEV config has the Vite origins allowed
    api = build()
    r = api.c.options("/api/import/accounts/inspect", headers={
        "Origin": "http://127.0.0.1:5173", "Access-Control-Request-Method": "POST",
        "Access-Control-Request-Headers": "x-file-name,content-type"})
    assert r.status_code == 200 and "x-file-name" in r.headers["access-control-allow-headers"].lower()


# ======================================================================= body caps
def test_five_mib_cap_applies_to_the_upload_route_only(api):
    big = b"PK\x03\x04" + b"0" * (3 * 1024 * 1024)
    r = K.upload(api, big)                                 # 3 MiB: over the global cap, under the import cap
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_file"
    r = api.c.post("/api/client/save", content=b"x" * (3 * 1024 * 1024), headers={"Content-Type": "application/json"})
    assert r.status_code == 413 and r.json()["error"]["code"] == "payload_too_large"
    r = K.upload(api, b"PK\x03\x04" + b"0" * (C.MAX_IMPORT_BYTES + 10))
    assert r.status_code == 413 and r.json()["error"]["details"] == {"max_bytes": C.MAX_IMPORT_BYTES, "reason": "too_large"}
    assert r.headers["x-content-type-options"] == "nosniff"
    # chunked body (no Content-Length) is capped too
    def gen():
        for _ in range(6):
            yield b"0" * (1024 * 1024)
    r = api.c.post("/api/import/accounts/inspect", content=gen(), headers={"X-File-Name": "a.xlsx"})
    assert r.status_code == 413


# ======================================================================= Inspect / Preview
def test_inspect_counts_states_and_no_write(api):
    api.mk_client("Web Mgmt Ltd", bpid="104233", products="Terminal")
    api.mk_client("Zeta Inc", notes="keep me")
    api.mk_client("Dormant Co", status="Inactive")
    before = snapshot(api)
    rows = [["Web Ltd", "104233", "B-PIPE", "Active", "2026-11-05", "n"],      # 2 needs_review
            ["  ZETA   inc ", None, "P", None, None, "other note"],          # 3 update (exact)
            ["Brand New Corp", "9", None, None, None, None],                 # 4 new
            ["Bad Status", None, None, "Dormant"],                           # 5 invalid
            ["Dormant Co", None, "X", "Active"],                             # 6 update on Inactive
            ["brand   new CORP", None, "P2", None, None, "late"],            # 7 in-file exact twin of row 4
            ["Zeta Inc"]]                                                    # 8 twin of exact update row 3
    item = insp(api, rows)
    assert item["counts"] == {"rows": 7, "valid": 6, "new": 2, "updates": 3, "needs_review": 1, "invalid": 1, "unchanged": 1}
    c = item["counts"]
    assert c["valid"] == c["new"] + c["updates"] + c["needs_review"] and c["valid"] + c["invalid"] == c["rows"]
    r = by_row(item)
    assert r[2]["state"] == "needs_review" and r[3]["state"] == "update" and r[4]["state"] == "new"
    assert r[5]["state"] == "invalid" and r[5]["raw"]["status"] == "Dormant"
    assert r[3]["match"]["fills"] == {"products": "P"} and r[3]["match"]["kept"]["notes"] == {
        "existing": "keep me", "file": "other note", "differs": True}
    assert r[6]["match"]["notes"] == ["Existing account is Inactive and stays Inactive"]
    assert r[7]["merged_into"] == {"row": 4} and r[7]["state"] == "new" and r[7]["candidates"] == []
    assert r[8]["merged_into"] == {"row": 3} and r[8]["state"] == "update" and r[8]["unchanged"] is True
    assert item["file"] == {"name": "accounts.xlsx", "size_bytes": item["file"]["size_bytes"], "format": "xlsx",
                            "sheet": "Accounts", "rows_found": 7, "encoding": None, "delimiter": None, "notes": []}
    assert item["plan_hash"].startswith("sha256:") and item["limits"]["max_rows"] == 500
    assert re.fullmatch(r"[0-9a-f]{32}", item["import_id"]) and item["expires_at"].endswith("Z")
    cand = r[2]["candidates"][0]
    assert cand["strength"] == "High" and [x["code"] for x in cand["reasons"]] == ["contained_name", "same_bpid"]
    assert cand["fills"] == {"next_contact": "2026-11-05", "notes": "n"} and cand["kept"]["products"]["differs"] is True
    assert snapshot(api) == before                       # Inspect wrote nothing at all
    assert not re.search(r"score|percent|similarity", json.dumps(item), re.I)


def test_inspect_invalid_rows_take_no_part_in_matching(api):
    api.mk_client("Orbit Systems")
    item = insp(api, [["Orbit Systems Ltd", None, None, "Bogus"], ["Orbit Systems Group"]])
    assert by_row(item)[2]["state"] == "invalid"
    assert by_row(item)[3]["state"] == "update" or by_row(item)[3]["state"] == "needs_review"
    # the invalid row 2 is not a candidate for row 3
    assert all(c["ref"]["type"] == "account" for c in by_row(item)[3]["candidates"])


def test_in_file_similar_rows_flag_the_later_row(api):
    item = insp(api, [["Pinecrest Advisory"], ["Pinecrest Advisory Group"], ["Other Co", "77"], ["Third Co", "77"]])
    r = by_row(item)
    assert r[2]["state"] == "new" and r[3]["state"] == "needs_review"
    assert r[3]["candidates"][0]["ref"] == {"type": "row", "row": 2} and "row_version" not in r[3]["candidates"][0]
    assert r[5]["state"] == "needs_review" and r[5]["candidates"][0]["reasons"][0]["code"] == "same_bpid"
    assert r[4]["state"] == "new"


def test_new_rows_that_collide_after_normalisation_are_flagged_not_merged(api):
    item = insp(api, [["ACME, Inc."], ["Acme Inc"]])
    assert by_row(item)[3]["state"] == "needs_review" and by_row(item)[3]["candidates"][0]["strength"] == "High"


def test_bpid_edge_rules(api):
    api.mk_client("Alpha One", bpid="55")
    api.mk_client("Beta Two", bpid="66")
    item = insp(api, [["Alpha One", "66", None], ["Quartz Zed", "55"]])
    r = by_row(item)
    w = {x["code"] for x in r[2]["warnings"]}
    assert r[2]["state"] == "update" and w == {"different_bpid", "bpid_used_elsewhere"}      # warning, not a flag
    assert r[2]["match"]["kept"]["bpid"]["existing"] == "55"
    assert r[3]["state"] == "needs_review" and r[3]["candidates"][0]["reasons"][0]["code"] == "same_bpid"


def test_inspect_refusals_leave_no_session_and_no_write(api):
    before = snapshot(api)
    st = svc.store(api.rt)
    for data, name, status in [(b"junk", "a.xlsx", 422), (K.xlsx([["a"]]), "a.txt", 422),
                               (K.rezip(K.xlsx([["a"]]), add={"xl/vbaProject.bin": b"x"}), "a.xlsx", 422),
                               (K.xlsx([["x"]] * 501), "a.xlsx", 422)]:
        r = K.upload(api, data, name)
        assert r.status_code == status and r.json()["error"]["code"] == "invalid_file"
        assert r.json()["error"]["details"]["reason"]
    assert st.count() == 0 and snapshot(api) == before


def test_session_limits_and_expiry(api):
    ids = [insp(api, [[f"Co {i}"]])["import_id"] for i in range(4)]
    assert api.c.get(f"/api/import/accounts/{ids[0]}").status_code == 410          # oldest discarded (max 3)
    assert all(api.c.get(f"/api/import/accounts/{i}").status_code == 200 for i in ids[1:])
    st = svc.store(api.rt)
    t = [1000.0]
    st.now = lambda: t[0]
    keep = insp(api, [["Co X"]])["import_id"]
    t[0] += C.IMPORT_SESSION_TTL_SECONDS - 5
    assert api.c.get(f"/api/import/accounts/{keep}").status_code == 200            # touching refreshes the idle timer
    t[0] += C.IMPORT_SESSION_TTL_SECONDS + 5
    r = api.c.get(f"/api/import/accounts/{keep}")
    assert r.status_code == 410 and r.json()["error"]["code"] == "import_expired"


def test_restart_loses_sessions(build):
    a = build()
    item = insp(a, [["Co"]])
    a.rt.import_sessions = None                                                  # what a restart does
    assert a.c.get(f"/api/import/accounts/{item['import_id']}").status_code == 410
    err(a, item, [], 410, "import_expired")


def test_get_plan_refreshes_hash_after_a_change(api):
    item = insp(api, [["Gamma Delta"]])
    again = api.c.get(f"/api/import/accounts/{item['import_id']}").json()["item"]
    assert again["plan_hash"] == item["plan_hash"]
    api.mk_client("Gamma Delta Ltd", duplicates_confirmed=True)
    fresh = api.c.get(f"/api/import/accounts/{item['import_id']}").json()["item"]
    assert fresh["plan_hash"] != item["plan_hash"] and by_row(fresh)[2]["state"] == "needs_review"


# ======================================================================= Apply
def test_apply_creates_updates_audits_and_backs_up(api):
    web = api.mk_client("Web Mgmt Ltd", bpid="104233", products="Terminal")
    zeta = api.mk_client("Zeta Inc", notes="keep me", next_contact="2026-12-31")
    item = insp(api, [["Web Ltd", "104233", "B-PIPE", None, "2026-11-05", "n"],
                      ["zeta inc", "Z1", "P", "Inactive", "2026-01-01", "other"],
                      ["Brand New Corp", "9", "NP", "Inactive", "2026-02-02", "note"],
                      ["Bad", None, None, "Dormant"],
                      ["brand new corp", None, None, None, None, "ignored: earlier row wins"]])
    r = by_row(item)
    dec = [same(2, r[2]["candidates"][0])]
    t0 = len(api.audit())
    j = apply_(api, item, dec)
    res = j["item"]
    assert res["counts"] == {"created": 1, "updated": 2, "unchanged": 0, "skipped": 1, "invalid": 1, "overwritten": 0}
    assert res["breakdown"] == {"created_new": 1, "created_marked_different": 0, "updated_exact": 1,
                                "updated_marked_same": 1, "skipped_invalid": 1, "skipped_by_user": 0}
    assert res["audit_events"] == 3 and res["fields_filled"] == 2 + 2 and res["backup"]["reason"] == "pre-import"
    assert j["side_effects"] == [{"type": "backup_created", "file": res["backup"]["name"], "reason": "pre-import"}]
    # fill blanks only: nothing non-blank was overwritten
    w = acct(api, "Web Mgmt Ltd")
    assert (w["products"], w["bpid"], w["next_contact"], w["notes"], w["row_version"]) == ("Terminal", "104233", "2026-11-05", "n", 2)
    z = acct(api, "Zeta Inc")
    assert (z["notes"], z["next_contact"], z["bpid"], z["products"], z["status"], z["row_version"]) == (
        "keep me", "2026-12-31", "Z1", "P", "Active", 2)
    n = acct(api, "Brand New Corp")
    assert (n["status"], n["bpid"], n["products"], n["next_contact"], n["notes"], n["row_version"]) == (
        "Inactive", "9", "NP", "2026-02-02", "note", 1)
    # audit: one row per created / updated account, shape per contract
    ev = api.audit(source="Bulk Import")
    assert len(ev) == 3 and len(api.audit()) == t0 + 3
    assert sorted((e["action"], e["object_label"]) for e in ev) == [
        ("created", "Brand New Corp"), ("updated", "Web Mgmt Ltd"), ("updated", "Zeta Inc")]
    for e in ev:
        assert (e["category"], e["object_type"], e["source"]) == ("crm", "client", "Bulk Import")
        assert e["client_id"] == e["object_id"] and e["reason"].startswith("Bulk import: accounts.xlsx, row")
    zu = next(e for e in ev if e["object_label"] == "Zeta Inc")
    assert json.loads(zu["old_value"]) == {"bpid": None, "products": None} and json.loads(zu["new_value"]) == {"bpid": "Z1", "products": "P"}
    assert zu["reason"] == "Bulk import: accounts.xlsx, row 3"
    assert next(e for e in ev if e["action"] == "created")["reason"] == "Bulk import: accounts.xlsx, rows 4, 6"
    # the safety backup exists, has integrity, and holds the PRE-import state
    bk = api.rt.config.backup_dir / res["backup"]["name"]
    c = sqlite3.connect(str(bk))
    assert c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert c.execute("SELECT count(*) FROM clients").fetchone()[0] == 2 and c.execute("SELECT count(*) FROM audit_events WHERE source='Bulk Import'").fetchone()[0] == 0
    c.close()
    assert res["skipped"] == [{"row": 5, "name": "Bad", "kind": "invalid", "reason": "Status: 'Dormant' is not in the list (Active, Inactive)"}]
    # activity sentence
    from server.sentences import system_sentence
    assert system_sentence("client", "created", "X", source="Bulk Import") == "Account created by bulk import"
    assert system_sentence("client", "updated", "X", source="Bulk Import") == "Account updated by bulk import"
    assert system_sentence("client", "created", "X", source="GUI") == "Account created: X"


def test_apply_different_account_and_skip_and_row_target(api):
    api.mk_client("Brightwater Pension Trust")
    item = insp(api, [["Pinecrest Advisory"], ["Pinecrest Advisory Group", None, "P"], ["Bright Water Holdings"], ["Orbit Co"],
                      ["Orbit Co Group", "5", None, None, None, "from later row"]])
    r = by_row(item)
    assert [r[i]["state"] for i in (2, 3, 4, 5, 6)] == ["new", "needs_review", "needs_review", "new", "needs_review"]
    dec = [same(3, r[3]["candidates"][0]), {"row": 4, "choice": "skip"}, {"row": 6, "choice": "different"}]
    res = apply_(api, item, dec)["item"]
    assert res["counts"]["created"] == 3 and res["counts"]["skipped"] == 1 and res["counts"]["updated"] == 0
    assert res["breakdown"]["created_marked_different"] == 1
    pa = acct(api, "Pinecrest Advisory")
    assert pa["products"] == "P"                                              # row 3 merged into earlier row 2 (new)
    assert api.sql("SELECT count(*) AS n FROM clients WHERE name = 'Pinecrest Advisory Group'")[0]["n"] == 0
    assert acct(api, "Orbit Co Group")["notes"] == "from later row"
    sk = res["skipped"][0]
    assert sk == {"row": 4, "name": "Bright Water Holdings", "kind": "by_user",
                  "reason": "Skipped by you (possible duplicate of Brightwater Pension Trust)"}
    assert res["skipped_url"].endswith("/skipped.xlsx")
    created = {c["name"]: c["rows"] for c in res["created"]}
    assert created["Pinecrest Advisory"] == [2, 3]


def test_inactive_stays_inactive_and_status_never_changes(api):
    api.mk_client("Dormant Co", status="Inactive")
    api.mk_client("Dormant Corp Intl", status="Inactive", products="Old")
    item = insp(api, [["Dormant Co", "7", "P", "Active"], ["Dormant Corp", None, "New", "Active"]])
    r = by_row(item)
    assert r[2]["state"] == "update" and r[3]["state"] == "needs_review"
    assert r[3]["candidates"][0]["notes"] == ["Existing account is Inactive and stays Inactive"]
    apply_(api, item, [same(3, r[3]["candidates"][0])])
    assert acct(api, "Dormant Co")["status"] == "Inactive" and acct(api, "Dormant Co")["bpid"] == "7"
    d = acct(api, "Dormant Corp Intl")
    assert d["status"] == "Inactive" and d["products"] == "Old"


FIELDS = ("bpid", "products", "next_contact", "notes")


@pytest.mark.parametrize("existing_filled", [True, False])
@pytest.mark.parametrize("file_filled", [True, False])
@pytest.mark.parametrize("status", ["Active", "Inactive"])
def test_fill_blanks_matrix_every_field(api, existing_filled, file_filled, status):
    ex = {"bpid": "E1", "products": "EP", "next_contact": "2030-01-01", "notes": "EN"} if existing_filled else {}
    api.mk_client("Matrix Co", status=status, **ex)
    fv = {"bpid": "F1", "products": "FP", "next_contact": "2031-02-02", "notes": "FN"} if file_filled else {}
    item = insp(api, [["Matrix Co", fv.get("bpid"), fv.get("products"), "Active" if file_filled else None,
                       fv.get("next_contact"), fv.get("notes")]])
    res = apply_(api, item)["item"]
    row = acct(api, "Matrix Co")
    for f in FIELDS:
        want = ex.get(f) if f in ex else fv.get(f)
        assert row[f] == want, f
    assert row["status"] == status
    assert res["counts"]["overwritten"] == 0
    changed = file_filled and not existing_filled
    assert res["counts"]["updated"] == (1 if changed else 0) and res["counts"]["unchanged"] == (0 if changed else 1)
    assert row["row_version"] == (2 if changed else 1) or not changed and row["row_version"] == 1
    if not changed:
        assert res["backup"] is None and res["audit_events"] == 0


def test_nothing_to_change_writes_nothing_and_is_idempotent(api):
    api.mk_client("Idem Co", bpid="1", products="P", next_contact="2030-05-05", notes="n")
    rows = [["Idem Co", "1", "P", "Active", "2030-05-05", "n"], ["New One", "2", "Q", None, None, None]]
    first = apply_(api, insp(api, rows))["item"]
    assert first["counts"]["created"] == 1 and first["counts"]["unchanged"] == 1
    snap = snapshot(api)
    second_item = insp(api, rows)
    assert second_item["counts"]["updates"] == 2 and second_item["counts"]["unchanged"] == 2
    second = apply_(api, second_item)["item"]
    assert second["counts"]["created"] == 0 and second["counts"]["updated"] == 0 and second["counts"]["unchanged"] == 2
    assert second["backup"] is None and second["audit_events"] == 0
    assert snapshot(api) == snap                          # no row_version bump, no audit, no backup file


def test_unresolved_confirm_validation_and_repeat(api):
    api.mk_client("Web Mgmt Ltd")
    item = insp(api, [["Web Ltd"], ["Fresh Co"]])
    snap = snapshot(api)
    e = err(api, item, [], 409, "unresolved_duplicates")
    assert e["details"] == {"rows": [2]}
    err(api, item, [{"row": 2, "choice": "skip"}], 409, "confirmation_required", confirm=False)
    assert snapshot(api) == snap                                       # no write, no backup
    cand = by_row(item)[2]["candidates"][0]
    bad = [
        [{"row": 3, "choice": "skip"}],                                  # not a flagged row
        [{"row": 9, "choice": "skip"}],                                  # unknown row
        [{"row": 2, "choice": "skip"}, {"row": 2, "choice": "different"}],
        [{"row": 2, "choice": "same"}],                                  # missing target
        [{"row": 2, "choice": "different", "target": cand["ref"]}],
        [{"row": 2, "choice": "same", "target": {"type": "account", "id": "nope"}}],
        [{"row": 2, "choice": "same", "target": {"type": "row", "row": 3}}],
    ]
    for d in bad:
        e = err(api, item, d, 422, "validation_error")
        assert "decisions" in e["fields"]
    api.post("/api/import/accounts/apply", {"import_id": item["import_id"], "plan_hash": item["plan_hash"], "confirm": True,
                                            "decisions": [], "extra": 1})
    assert api.post("/api/import/accounts/apply", {"import_id": "x", "plan_hash": "y", "extra": 1}).status_code == 422
    assert snapshot(api) == snap
    ok = apply_(api, item, [{"row": 2, "choice": "different"}])
    assert ok["item"]["counts"]["created"] == 2
    e = err(api, item, [{"row": 2, "choice": "different"}], 409, "already_applied")
    assert e["details"]["result"]["counts"]["created"] == 2
    assert api.sql("SELECT count(*) AS n FROM clients")[0]["n"] == 3     # second Apply wrote nothing
    assert api.c.get(f"/api/import/accounts/{item['import_id']}").status_code == 409


def test_same_onto_skipped_or_missing_row_is_422(api):
    item = insp(api, [["Pinecrest Advisory"], ["Pinecrest Advisory Group"], ["Pinecrest Advisory Ltd."]])
    r = by_row(item)
    assert r[3]["state"] == "needs_review"
    # row 2 is 'new' (not flagged); 4 exact-merges? 'Pinecrest Advisory Ltd.' core == row 2 core: flagged against row 2 and row 3
    cands4 = {json.dumps(c["ref"], sort_keys=True): c for c in r[4]["candidates"]}
    t3 = cands4[json.dumps({"type": "row", "row": 3}, sort_keys=True)]
    e = err(api, item, [{"row": 3, "choice": "skip"}, same(4, t3)], 422, "validation_error")
    assert "skipped" in e["fields"]["decisions"]
    ok = apply_(api, item, [{"row": 3, "choice": "different"}, same(4, t3)])["item"]
    assert ok["counts"]["created"] == 2


# ----------------------------------------------------------------------- staleness / concurrency
def stale_setup(api):
    api.mk_client("Web Mgmt Ltd", bpid="1")
    api.mk_client("Zeta Inc")
    item = insp(api, [["Web Ltd", "1", "P"], ["Zeta Inc", None, "Z"], ["Fresh Co"]])
    return item, [same(2, by_row(item)[2]["candidates"][0])]


def test_stale_when_a_matched_or_candidate_account_is_edited(api):
    item, dec = stale_setup(api)
    web = acct(api, "Web Mgmt Ltd")
    api.ok("/api/client/save", {"id": web["id"], "expected_row_version": 1, "notes": "edited elsewhere"})
    snap = snapshot(api)
    e = err(api, item, dec, 409, "preview_stale")
    assert e["details"]["changed"][0]["row"] == 2 and e["details"]["why"] == "plan_hash"
    assert snapshot(api) == snap
    fresh = api.c.get(f"/api/import/accounts/{item['import_id']}").json()["item"]       # refreshed plan keeps unchanged rows
    assert by_row(fresh)[3]["state"] == "update" and by_row(fresh)[2]["candidates"][0]["row_version"] == 2
    ok = apply_(api, fresh, [same(2, by_row(fresh)[2]["candidates"][0])], plan_hash=fresh["plan_hash"])
    assert ok["item"]["counts"]["updated"] == 2


def test_stale_when_account_deleted_or_new_similar_account_created(api):
    item, dec = stale_setup(api)
    z = acct(api, "Zeta Inc")
    api.ok("/api/client/delete", {"id": z["id"], "confirm": True, "confirm_name": "Zeta Inc"})
    err(api, item, dec, 409, "preview_stale")
    item2 = insp(api, [["Zeta Inc"], ["Fresh Co"]])
    api.ok("/api/client/save", {"name": "Fresh Co Ltd", "duplicates_confirmed": True})
    e = err(api, item2, [], 409, "preview_stale")
    assert {c["row"] for c in e["details"]["changed"]} == {3}


def test_stale_when_exact_twin_appears_after_preview(api):
    item = insp(api, [["Brand New"]])
    api.mk_client("Brand New")
    err(api, item, [], 409, "preview_stale")
    assert api.sql("SELECT count(*) AS n FROM clients WHERE name = 'Brand New'")[0]["n"] == 1


def test_edit_between_check_and_write_is_caught_inside_the_transaction(api):
    item, dec = stale_setup(api)
    web = acct(api, "Web Mgmt Ltd")

    def edit(rt, plan):                          # lands after the pre-check and the backup, before BEGIN IMMEDIATE
        with rt.db() as c:
            c.execute("UPDATE clients SET notes = 'raced', row_version = row_version + 1, updated_at = ? WHERE id = ?",
                      ("2030-01-01T00:00:00.000Z", web["id"]))
    svc.HOOKS["after_backup"] = edit
    snap_clients = api.sql("SELECT name FROM clients ORDER BY name")
    e = err(api, item, dec, 409, "preview_stale")
    assert e["details"]["why"] == "row_version"
    assert api.sql("SELECT name FROM clients ORDER BY name") == snap_clients
    assert api.audit(source="Bulk Import") == [] and acct(api, "Zeta Inc")["products"] is None
    assert acct(api, "Web Mgmt Ltd")["notes"] == "raced"


def test_atomic_rollback_when_the_last_write_fails(api):
    api.mk_client("Existing Co")
    item = insp(api, [["Existing Co", "1", "P"], ["N1"], ["N2"], ["N3"]])
    before = snapshot(api)
    calls = []

    def boom(n, dest):
        calls.append(n)
        if n == 4:
            raise RuntimeError("injected failure on the last write")
    svc.HOOKS["after_write"] = boom
    r = api.c.post("/api/import/accounts/apply", json={"import_id": item["import_id"], "plan_hash": item["plan_hash"],
                                                       "confirm": True, "decisions": []})
    assert r.status_code == 500 and r.json()["error"]["code"] == "internal"
    assert calls == [1, 2, 3, 4]
    after = snapshot(api)
    assert after["clients"] == before["clients"] and after["audit"] == before["audit"]      # content identical to before
    assert len(after["backups"]) == len(before["backups"]) + 1                              # only the backup file exists
    svc.HOOKS.clear()
    again = apply_(api, item)                                   # the session survives; a retry succeeds
    assert again["item"]["counts"]["created"] == 3


def test_backup_is_taken_before_the_first_write(api):
    api.mk_client("Existing Co")
    item = insp(api, [["Existing Co", "1"], ["Brand New"]])
    seen = {}

    def peek(rt, plan):
        seen["backups"] = [p.name for p in rt.config.backup_dir.iterdir() if "pre-import" in p.name]
        with rt.db() as c:
            seen["clients"] = c.execute("SELECT count(*) FROM clients").fetchone()[0]
            seen["audit"] = c.execute("SELECT count(*) FROM audit_events WHERE source='Bulk Import'").fetchone()[0]
    svc.HOOKS["after_backup"] = peek
    apply_(api, item)
    assert len(seen["backups"]) == 1 and seen["clients"] == 1 and seen["audit"] == 0


def test_backup_failure_writes_nothing(api, monkeypatch):
    api.mk_client("Existing Co")
    item = insp(api, [["Existing Co", "1"], ["Brand New"]])
    before = snapshot(api)
    from server import dbkit
    monkeypatch.setattr(dbkit.migrate, "online_backup", lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")))
    r = api.c.post("/api/import/accounts/apply", json={"import_id": item["import_id"], "plan_hash": item["plan_hash"],
                                                       "confirm": True, "decisions": []})
    assert r.status_code == 500 and r.json()["error"]["code"] == "backup_failed"
    assert snapshot(api) == before


def test_second_concurrent_apply_gets_409_conflict(api):
    item = insp(api, [["Brand New"]])
    st = svc.store(api.rt)
    assert st.apply_lock.acquire(blocking=False)
    try:
        err(api, item, [], 409, "conflict")
    finally:
        st.apply_lock.release()
    assert apply_(api, item)["item"]["counts"]["created"] == 1


def test_pre_import_backups_are_automatic_and_pruned_with_the_others(api):
    from server.backups import BackupService
    assert "pre-import" in C.AUTO_BACKUP_REASONS
    for i in range(C.AUTO_BACKUP_KEEP + 3):
        api.rt.backups.create(sqlite3.connect(str(api.rt.config.db_path)), "pre-import")
    assert len([b for b in api.rt.backups.list(100) if b["reason"] == "pre-import"]) == C.AUTO_BACKUP_KEEP
    assert BackupService  # noqa


def test_a_500_row_apply_is_one_short_transaction(api):
    import time
    item = insp(api, [[f"Unique Account {i:03d} {chr(97 + i % 26) * 5}"] for i in range(500)])
    assert item["counts"]["new"] + item["counts"]["needs_review"] == 500
    decisions = [{"row": r["row"], "choice": "different"} for r in item["rows"] if r["state"] == "needs_review"]
    t = time.perf_counter()
    res = apply_(api, item, decisions)["item"]
    print("500-row apply:", round(time.perf_counter() - t, 2), "s")
    assert res["counts"]["created"] == 500 and res["audit_events"] == 500
    assert len(api.audit(source="Bulk Import")) == 500


# ----------------------------------------------------------------------- skipped download
def test_skipped_xlsx_content_and_formula_safety_roundtrip(api):
    api.mk_client("Brightwater Pension Trust")
    rows = [["=SUM(A1)", "+1", "@x", "Dormant", None, "-5"],                       # invalid (bad status) with hostile text
            ["Bright Water Holdings", "9", "P", "Active", "2026-01-01", "=cmd|' /C calc'!A0"],
            ["Fine Co"]]
    item = insp(api, rows)
    assert by_row(item)[2]["state"] == "invalid"
    res = apply_(api, item, [{"row": 3, "choice": "skip"}])["item"]
    r = api.c.get(res["skipped_url"])
    assert r.status_code == 200 and r.headers["content-type"] == template_media()
    assert re.fullmatch(r'attachment; filename="tam_accounts_skipped_\d{8}_\d{4}\.xlsx"', r.headers["content-disposition"])
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Accounts"]
    assert [c.value for c in ws[1]] == K.HEAD + ["Source row", "Reason", "Possible duplicate of"]
    rows_out = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert rows_out[0][:6] == ["=SUM(A1)", "+1", "@x", "Dormant", None, "-5"] or rows_out[0][0] == "=SUM(A1)"
    assert rows_out[0][6] == "2" and rows_out[1][6] == "3" and rows_out[1][8] == "Brightwater Pension Trust"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            assert c.data_type != "f"
    assert rows_out[1][5] == "=cmd|' /C calc'!A0"
    # re-uploadable: extra columns ignored with a warning
    again = K.upload(api, r.content, "skipped.xlsx").json()["item"]
    assert {w["code"] for w in again["warnings"]} == {"unknown_column"}
    assert by_row(again)[2]["values"]["name"] == "=SUM(A1)"
    # wrong format / before apply
    assert api.c.get(res["skipped_url"][:-5] + ".csv").status_code == 404
    other = insp(api, [["Zed Zed"]])
    assert api.c.get(f"/api/import/accounts/{other['import_id']}/skipped.xlsx").status_code == 409


def test_nothing_skipped_means_no_url(api):
    res = apply_(api, insp(api, [["Only One"]]))["item"]
    assert res["skipped"] == [] and res["skipped_url"] is None


# ----------------------------------------------------------------------- csv end to end
def test_csv_import_end_to_end_and_skipped_csv(api):
    api.mk_client("Zeta Inc")
    text = ("Account name;BPID;Products;Status;Next contact date;Notes\r\n"
            "Zeta Inc;;P;;2026-05-05;n\r\n"
            "Caf\xe9 Corp;007;\"A;B\";Inactive;;\"two\r\nlines\"\r\n"
            "=HYPERLINK(\"x\");;;Nope;;\r\n").encode("cp1252")
    r = K.upload(api, text, "accounts.csv")
    item = r.json()["item"]
    assert r.status_code == 200 and item["file"]["format"] == "csv" and item["file"]["encoding"] == "Windows-1252"
    assert item["file"]["delimiter"] == "semicolon" and item["file"]["sheet"] == "CSV"
    assert item["file"]["notes"] == ["Read as CSV: encoding Windows-1252, delimiter semicolon."]
    assert item["counts"]["new"] == 1 and item["counts"]["updates"] == 1 and item["counts"]["invalid"] == 1
    res = apply_(api, item)["item"]
    assert res["skipped_url"].endswith("/skipped.csv")
    assert acct(api, "Café Corp")["products"] == "A;B" and acct(api, "Café Corp")["notes"] == "two\nlines"
    assert acct(api, "Café Corp")["bpid"] == "007"
    c = api.c.get(res["skipped_url"])
    assert c.headers["content-type"].startswith("text/csv") and c.content.startswith(b"\xef\xbb\xbf")
    body = c.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in body and "Source row,Reason,Possible duplicate of" in body
    assert api.c.get(res["skipped_url"][:-4] + ".xlsx").status_code == 404
    again = K.upload(api, c.content, "skipped.csv").json()["item"]
    assert by_row(again)[2]["values"]["name"] == '=HYPERLINK("x")'


def test_csv_and_xlsx_extension_content_mismatch_is_refused(api):
    r = K.upload(api, K.xlsx([["a"]]), "data.csv")
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "not_csv"
    r = K.upload(api, b"Account name\r\nx\r\n", "data.xlsx")
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "not_xlsx"


# ======================================================================= single add (T-86)
def test_check_similar_shapes_and_parity_with_import(api):
    api.mk_client("Web Mgmt Ltd", bpid="104233")
    api.mk_client("Castellan Peak")
    j = api.ok("/api/client/check-similar", {"name": "Web Ltd", "bpid": "104233"})
    assert j["side_effects"] == [] and j["item"]["exact"] is None
    c = j["item"]["candidates"][0]
    assert c["strength"] == "High" and [r["code"] for r in c["reasons"]] == ["contained_name", "same_bpid"]
    assert "fills" not in c and "kept" not in c and c["ref"]["type"] == "account" and c["row_version"] == 1
    ex = api.ok("/api/client/check-similar", {"name": "  castellan   PEAK ", "bpid": None})["item"]
    assert ex["exact"]["name"] == "Castellan Peak" and ex["candidates"] == []
    item = insp(api, [["Web Ltd", "104233"]])
    ic = by_row(item)[2]["candidates"][0]
    ic = {k: v for k, v in ic.items() if k not in ("fills", "kept")}
    assert ic == c                                              # identical candidates, reasons, strength, order
    assert api.err("/api/client/check-similar", {"name": ""}, 422, "validation_error")
    assert api.err("/api/client/check-similar", {"name": "x", "zzz": 1}, 422, "validation_error")
    assert api.err("/api/client/check-similar", {"name": "x" * 201}, 422, "validation_error")
    n = len(api.audit())
    api.ok("/api/client/check-similar", {"name": "Anything"})
    assert len(api.audit()) == n


def test_client_save_create_enforces_possible_duplicate(api):
    api.mk_client("Web Mgmt Ltd", bpid="104233")
    before = snapshot(api)
    e = api.err("/api/client/save", {"name": "Web Ltd"}, 409, "possible_duplicate")
    assert e["details"]["exact"] is None and e["details"]["candidates"][0]["name"] == "Web Mgmt Ltd"
    e2 = api.err("/api/client/save", {"name": "Web Ltd", "duplicates_confirmed": False}, 409, "possible_duplicate")
    assert snapshot(api) == before
    ok = api.ok("/api/client/save", {"name": "Web Ltd", "duplicates_confirmed": True})
    assert ok["item"]["name"] == "Web Ltd"
    # exact twin stays a conflict, even when confirmed
    api.err("/api/client/save", {"name": "web  ltd", "duplicates_confirmed": True}, 409, "conflict")
    api.err("/api/client/save", {"name": "WEB MGMT LTD"}, 409, "conflict")
    # no candidates: creates directly, as today
    api.ok("/api/client/save", {"name": "Totally Unrelated Name"})
    # updates are unaffected and ignore the flag
    a = acct(api, "Web Ltd")
    api.ok("/api/client/save", {"id": a["id"], "expected_row_version": 1, "notes": "x", "duplicates_confirmed": True})
    assert acct(api, "Web Ltd")["notes"] == "x"
    assert api.audit(object_id=a["id"], action="updated")


def test_import_never_surfaces_a_name_collision_as_a_database_error(api):
    api.mk_client("Société Générale")
    item = insp(api, [["SOCIÉTÉ  GÉNÉRALE"], ["société générale"], ["Societe Generale"]])
    r = by_row(item)
    assert r[2]["state"] == "update" and r[3]["state"] == "update" and r[3]["merged_into"] == {"row": 2}
    assert r[4]["state"] == "needs_review"                         # accents differ: flagged, never auto-merged
    res = apply_(api, item, [{"row": 4, "choice": "different"}])["item"]
    assert res["counts"]["created"] == 1
