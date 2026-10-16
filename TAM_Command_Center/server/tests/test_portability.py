"""Run 6 phase 1 (T-176..T-179): Excel export, database copy, restore from file (.xlsx and .csv)."""
from __future__ import annotations

import csv
import io
import sqlite3
import sys
import time
from pathlib import Path

import openpyxl
import pytest

from server.portability import restore as RS
from server.portability import tables as T
from server.portability.export import guard, unguard
from server.tests.importkit import rezip

FX = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "portability"


def snap(api):
    """Row sets of every exported table (all exported columns except row_version)."""
    out = {}
    for t in T.TABLES:
        cols = ", ".join(c.db for c in t.data_cols)
        out[t.key] = sorted(map(repr, api.sql(f"SELECT {cols} FROM {t.db}")))  # nosec B608 - fixed test identifiers
    return out


def counts(api):
    return {t.key: api.sql(f"SELECT count(*) n FROM {t.db}")[0]["n"] for t in T.TABLES}  # nosec B608


def seed(api):
    a = api.mk_client("Acme Capital", bpid="42", products="BPIPE")["id"]
    b = api.mk_client("=Evil Fund")["id"]
    api.ok("/api/contact/save", {"client_id": a, "name": "Pat Example", "email": "pat@example.test", "is_primary": True})
    api.ok("/api/contact/save", {"client_id": a, "name": "Sam Example"})
    api.mk_task(a, "Send report", priority="High", due="2026-10-30", notes="=cmd|' /C calc'!A0")
    api.ok("/api/note/save", {"client_id": a, "note_type": "Call Note", "subject": "Hi", "body": "line1\nline2"})
    obj = api.ok("/api/bams/object/save", {"client_id": a, "name": "BPOD-1", "region": "East", "environment": "PROD",
                                           "metrics": [{"metric_name": "CPU", "ticker": "T.CPU"}]})["item"]["object"]
    api.ok("/api/operations/quarter/save", {"year": 2026, "quarter": "Q4", "apics_version": "4.78.0",
                                            "east_prod_date": "2026-11-07"})
    api.ok("/api/operations/appliance/save", {"object_id": obj["id"], "year": 2026, "quarter": "Q4", "status": "Scheduled",
                                              "notes": "'quoted"})
    return a, b


def export_bytes(api):
    r = api.get("/api/export/excel")
    assert r.status_code == 200, r.text
    return r.content


def wipe(api):
    with api.rt.db() as conn:
        conn.execute("BEGIN")
        for t in ("server_upgrades", "quarterly_upgrade_windows", "bams_metrics", "bams_objects", "notes", "tasks",
                  "contacts", "clients"):
            conn.execute(f"DELETE FROM {t}")  # nosec B608 - fixed test identifiers
        conn.execute("COMMIT")


def up(api, data, name, session=None, table=None):
    params = {}
    if session:
        params["session"] = session
    if table:
        params["table"] = table
    return api.c.post("/api/restore/inspect", content=data, params=params, headers={"X-File-Name": name})


def ok_item(r):
    assert r.status_code == 200, r.text
    return r.json()["item"]


def apply(api, item, status=200, **kw):
    r = api.post("/api/restore/apply", {"restore_id": item["restore_id"], "plan_hash": item["plan_hash"], "confirm": True, **kw})
    assert r.status_code == status, r.text
    return r.json()


def csv_files_from_export(data):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    out = {}
    for ws in wb:
        if ws.title == "About":
            continue
        buf = io.StringIO(newline="")
        w = csv.writer(buf)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])
        out[ws.title] = buf.getvalue().encode("utf-8")
    return out


@pytest.fixture
def pa(build):
    return build()


# ------------------------------------------------------------------------------------------ export
def test_guard_roundtrip_is_exact():
    for s in ["=1+1", "+x", "-5", "@a", "'=x", "''=x", "plain", "'quote", "", "\t=x"]:
        assert unguard(guard(s)) == s, s
    assert guard("=1+1") == "'=1+1" and guard("'quote") == "'quote"


def test_excel_export_sheets_counts_guard_and_headers(pa):
    seed(pa)
    r = pa.get("/api/export/excel")
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    assert "tam_export_" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('_DEV.xlsx"')
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["About"] + [t.sheet for t in T.TABLES]
    about = {row[0].value: row[1].value for row in wb["About"].iter_rows(min_row=2)}
    assert about["app"] == "TAM Command Center" and about["channel"] == "DEV" and about["schema_version"] >= 1
    c = counts(pa)
    for t in T.TABLES:
        assert about[f"rows: {t.sheet}"] == c[t.key] == wb[t.sheet].max_row - 1
    assert [x.value for x in wb["Accounts"][1]][:3] == ["id", "name", "status"]
    names = [row[1].value for row in wb["Accounts"].iter_rows(min_row=2)]
    assert "'=Evil Fund" in names
    assert all(cell.data_type != "f" for ws in wb for row in ws.iter_rows() for cell in row)
    assert not any(x in wb.sheetnames for x in ("audit_events", "app_meta", "bams_metric_samples"))


def test_excel_export_of_empty_db_and_prod_filename(build):
    api = build()
    wb = openpyxl.load_workbook(io.BytesIO(export_bytes(api)))
    assert wb["Accounts"].max_row == 1 and len(wb.sheetnames) == 9
    prod = build(prod=True)
    r = prod.get("/api/export/excel")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('_PROD.xlsx"')


def test_export_writes_no_audit_event(pa):
    seed(pa)
    n = pa.sql("SELECT count(*) n FROM audit_events")[0]["n"]
    export_bytes(pa)
    pa.get("/api/export/database")
    assert pa.sql("SELECT count(*) n FROM audit_events")[0]["n"] == n


def test_database_copy_is_valid_sqlite_and_temp_removed(pa, tmp_path):
    seed(pa)
    r = pa.get("/api/export/database")
    assert r.status_code == 200 and r.content[:16] == b"SQLite format 3\x00"
    assert "tam_backup_" in r.headers["content-disposition"] and "_v" in r.headers["content-disposition"]
    p = tmp_path / "copy.db"
    p.write_bytes(r.content)
    c = sqlite3.connect(p)
    assert c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert c.execute("SELECT count(*) FROM clients").fetchone()[0] == 2
    c.close()
    assert not list(pa.rt.config.backup_dir.glob("*.tmp")) and not pa.rt.backups.list()


def test_database_copy_failure_is_plain_and_leaves_nothing(pa, monkeypatch):
    from server.portability import export
    def boom(*a, **k):
        raise RuntimeError("boom")
    monkeypatch.setattr(export.migrate, "online_backup", boom)
    r = pa.get("/api/export/database")
    assert r.status_code == 500 and r.json()["error"]["code"] == "backup_failed"
    assert not list(pa.rt.config.backup_dir.glob("*.tmp"))


# ------------------------------------------------------------------------------------------ restore, xlsx
def test_roundtrip_xlsx_into_emptied_db_equal_row_sets(pa):
    seed(pa)
    before = snap(pa)
    data = export_bytes(pa)
    wipe(pa)
    assert sum(counts(pa).values()) == 0
    item = ok_item(up(pa, data, "e.xlsx"))
    assert item["totals"]["new"] == sum(len(v) for v in before.values()) and item["can_apply"]
    assert item["totals"]["invalid"] == 0 and item["totals"]["differs"] == 0
    res = apply(pa, item)["item"]
    assert res["backup"]["reason"] == "pre-import" and list(pa.rt.config.backup_dir.glob("*pre-import*"))
    assert snap(pa) == before
    assert pa.sql("SELECT DISTINCT row_version v FROM clients") == [{"v": 1}]
    ev = pa.audit(object_type="import", action="restored")
    assert {e["object_id"] for e in ev} == {t.db for t in T.TABLES}
    assert all(e["source"] == "Bulk Import" and e["category"] == "data" for e in ev)
    pa.err("/api/restore/apply", {"restore_id": item["restore_id"], "plan_hash": item["plan_hash"], "confirm": True}, 409, "already_applied")
    again = ok_item(up(pa, data, "e.xlsx"))
    assert again["totals"]["new"] == 0 and again["totals"]["identical"] == res["counts"]["added"] and again["nothing_to_restore"]
    pa.err("/api/restore/apply", {"restore_id": again["restore_id"], "plan_hash": again["plan_hash"], "confirm": True}, 409, "conflict")


def test_inspect_writes_nothing(pa):
    seed(pa)
    data = export_bytes(pa)
    wipe(pa)
    before = (counts(pa), pa.sql("SELECT count(*) n FROM audit_events")[0]["n"], len(pa.rt.backups.list()))
    ok_item(up(pa, data, "e.xlsx"))
    assert (counts(pa), pa.sql("SELECT count(*) n FROM audit_events")[0]["n"], len(pa.rt.backups.list())) == before


def test_differs_never_overwritten_and_skipped_downloads(pa):
    a, _ = seed(pa)
    wb = openpyxl.load_workbook(io.BytesIO(export_bytes(pa)))
    wb["Accounts"].cell(row=2, column=2).value = "Acme Renamed"
    buf = io.BytesIO()
    wb.save(buf)
    item = ok_item(up(pa, buf.getvalue(), "edited.xlsx"))
    t = next(x for x in item["tables"] if x["key"] == "accounts")
    assert t["differs"] == 1 and t["identical"] == 1 and item["totals"]["new"] == 0
    assert item["issues"][0]["outcome"] == "differs" and "name" in item["issues"][0]["text"]
    assert pa.sql("SELECT name FROM clients WHERE id = ?", (a,))[0]["name"] == "Acme Capital"
    wb["Accounts"].append(["new-id-1", "Brand New Ltd", "Active"])
    buf2 = io.BytesIO()
    wb.save(buf2)
    item = ok_item(up(pa, buf2.getvalue(), "edited2.xlsx"))
    assert item["totals"]["new"] == 1 and item["totals"]["differs"] == 1
    apply(pa, item)
    assert pa.sql("SELECT name FROM clients WHERE id = ?", (a,))[0]["name"] == "Acme Capital"
    x = pa.get(f"/api/restore/{item['restore_id']}/skipped.xlsx")
    assert x.status_code == 200 and "Accounts" in openpyxl.load_workbook(io.BytesIO(x.content)).sheetnames
    c = pa.get(f"/api/restore/{item['restore_id']}/skipped.csv")
    assert c.status_code == 200 and b"differs" in c.content
    assert pa.get("/api/restore/nope/skipped.xlsx").status_code == 410


def test_xlsx_refusals(pa, build):
    seed(pa)
    data = export_bytes(pa)
    assert up(pa, (FX / "xlsx" / "not_an_export.xlsx").read_bytes(), "x.xlsx").json()["error"]["details"]["reason"] == "not_an_export"
    assert up(pa, (FX / "xlsx" / "newer_schema.xlsx").read_bytes(), "x.xlsx").json()["error"]["details"]["reason"] == "schema_newer"
    prod = build(prod=True)
    r = up(prod, data, "d.xlsx")
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "channel_mismatch"
    assert "production" in r.json()["error"]["message"]
    assert up(pa, export_bytes(prod), "p.xlsx").status_code == 200            # PROD export into DEV is allowed
    assert up(pa, b"MZ not a zip", "x.xlsx").status_code == 422
    assert up(pa, b"a,b\n1,2\n", "x.txt").status_code == 422
    assert up(pa, rezip(data, add={"xl/vbaProject.bin": b"x"}), "m.xlsx").json()["error"]["details"]["reason"] == "macros"
    assert up(pa, b"", "e.xlsx").status_code == 422


def test_formula_cell_row_is_invalid(pa):
    seed(pa)
    wb = openpyxl.load_workbook(io.BytesIO(export_bytes(pa)))
    wb["Accounts"].append(["f-1", "x", "Active"])
    last = wb["Accounts"].max_row
    wb["Accounts"].cell(row=last, column=2).value = "=1+1"
    buf = io.BytesIO()
    wb.save(buf)
    item = ok_item(up(pa, buf.getvalue(), "f.xlsx"))
    iss = [i for i in item["issues"] if i["code"] == "formula_not_allowed"]
    assert iss and iss[0]["table"] == "accounts" and iss[0]["row"] == last and iss[0]["column"] == "name"


def test_size_and_row_limits(pa):
    big = b"id,name\n" + b"x" * (25 * 1024 * 1024 + 10)
    assert pa.c.post("/api/restore/inspect", content=big, headers={"X-File-Name": "b.csv"}).status_code == 413
    rows = "\n".join(f"a{i},Name {i}" for i in range(10_005))
    r = up(pa, ("id,name\n" + rows).encode(), "Accounts.csv")
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "too_many_rows"


# ------------------------------------------------------------------------------------------ restore, csv
def test_roundtrip_csv_all_tables_into_emptied_db(pa):
    seed(pa)
    before = snap(pa)
    files = csv_files_from_export(export_bytes(pa))
    wipe(pa)
    sid = None
    for sheet, data in files.items():
        item = ok_item(up(pa, data, sheet.replace(" ", "_") + ".csv", session=sid))
        sid = item["restore_id"]
        f = item["files"][-1]
        assert f["table_label"] == sheet and f["table_source"] in ("headers", "filename"), (sheet, f)
    assert item["csv_notice"].startswith("CSV files cannot tell") and item["totals"]["invalid"] == 0
    apply(pa, item)
    assert snap(pa) == before


def test_csv_child_alone_is_invalid_with_hint_then_succeeds_with_parent(pa):
    item = ok_item(up(pa, (FX / "invalid" / "missing_parent_Contacts.csv").read_bytes(), "Contacts.csv"))
    assert item["totals"]["invalid"] == 1 and "Add the Accounts file" in item["issues"][0]["text"]
    assert not item["can_apply"]
    first = ok_item(up(pa, (FX / "valid" / "Accounts.csv").read_bytes(), "Accounts.csv"))
    item = ok_item(up(pa, (FX / "valid" / "Contacts.csv").read_bytes(), "Contacts.csv", session=first["restore_id"]))
    assert item["totals"]["invalid"] == 0 and item["totals"]["new"] == 5
    apply(pa, item)
    assert counts(pa)["accounts"] == 2 and counts(pa)["contacts"] == 3


def test_csv_all_valid_fixtures_together(pa):
    sid = None
    for f in sorted((FX / "valid").glob("*.csv")):
        item = ok_item(up(pa, f.read_bytes(), f.name, session=sid))
        sid = item["restore_id"]
    assert item["totals"]["invalid"] == 0, item["issues"]
    apply(pa, item)
    assert all(v >= 1 for v in counts(pa).values())
    assert pa.sql("SELECT count(*) n FROM audit_events WHERE action='restored'")[0]["n"] == 8


@pytest.mark.parametrize("name", ["variants_Accounts_semicolon.csv", "variants_Accounts_bom.csv", "variants_Accounts_shuffled.csv"])
def test_csv_variants_detected_as_accounts(pa, name):
    item = ok_item(up(pa, (FX / name).read_bytes(), name))
    assert item["files"][0]["table"] == "accounts" and item["totals"]["new"] == 2


def test_csv_encodings_and_delimiters(pa):
    text = "id,name,status\nw1,Café Fund,Active\n"
    assert ok_item(up(pa, text.encode("cp1252"), "Accounts.csv"))["totals"]["new"] == 1
    assert ok_item(up(pa, text.encode("utf-16"), "Accounts.csv"))["files"][0]["table"] == "accounts"
    assert ok_item(up(pa, text.replace(",", "\t").encode(), "Accounts.csv"))["totals"]["new"] == 1


def test_csv_unrecognized_choose_table_and_missing_columns(pa):
    item = ok_item(up(pa, (FX / "invalid" / "Unrecognized.csv").read_bytes(), "whatever.csv"))
    f = item["files"][0]
    assert f["state"] == "unrecognized" and f["table"] is None and not item["can_apply"]
    assert item["blocking"][0]["code"] == "unrecognized_file"
    r = pa.post(f"/api/restore/{item['restore_id']}/table", {"file_id": f["file_id"], "table": "accounts"})
    assert r.status_code == 200
    got = r.json()["item"]
    f2 = got["files"][0]
    assert f2["state"] == "missing_columns" and "id" in f2["missing_columns"] and "name" in f2["missing_columns"]
    assert pa.post("/api/restore/apply", {"restore_id": got["restore_id"], "plan_hash": got["plan_hash"], "confirm": True}).status_code == 409
    ok = ok_item(up(pa, b"id,name\nz1,Zed\n", "x.csv", table="accounts"))
    assert ok["files"][0]["table_source"] == "chosen"


def test_csv_duplicate_table_mixing_and_remove(pa):
    acc = (FX / "valid" / "Accounts.csv").read_bytes()
    item = ok_item(up(pa, acc, "Accounts.csv"))
    sid = item["restore_id"]
    r = up(pa, acc, "Accounts2.csv", session=sid)
    assert r.status_code == 422 and r.json()["error"]["details"]["reason"] == "table_already_added"
    assert "already added" in r.json()["error"]["message"]
    assert up(pa, export_bytes(pa), "e.xlsx", session=sid).json()["error"]["details"]["reason"] == "mixed_formats"
    gone = pa.post(f"/api/restore/{sid}/remove", {"file_id": item["files"][0]["file_id"]}).json()["item"]
    assert gone["files"] == [] and not gone["can_apply"]
    assert pa.post(f"/api/restore/{sid}/remove", {"file_id": "zz"}).status_code == 404
    assert ok_item(up(pa, acc, "Accounts.csv", session=sid))["totals"]["new"] == 2
    assert up(pa, acc, "A.csv", session="deadbeef").status_code == 410


# ------------------------------------------------------------------------------------------ invalid rows
def stage(pa, *names):
    sid = None
    for n in names:
        item = ok_item(up(pa, (FX / n).read_bytes(), Path(n).name, session=sid))
        sid = item["restore_id"]
    return item


def test_invalid_fixture_rows(pa):
    item = stage(pa, "valid/Accounts.csv", "invalid/bad_values_Tasks.csv")
    assert {i["code"] for i in item["issues"]} == {"value_not_allowed"} and item["totals"]["invalid"] == 2
    assert {i["column"] for i in item["issues"]} == {"status", "priority"}
    item = stage(pa, "valid/Accounts.csv", "invalid/two_primaries_Contacts.csv")
    assert [i["code"] for i in item["issues"]] == ["second_primary"] and item["totals"]["new"] == 3
    apply(pa, stage(pa, "valid/Accounts.csv"))
    item = stage(pa, "invalid/dup_name_Accounts.csv")
    assert item["totals"]["invalid"] == 1 and item["issues"][0]["code"] == "name_taken"
    item = stage(pa, "invalid/missing_parent_Tasks.csv")
    assert item["issues"][0]["code"] == "parent_missing"


def test_duplicate_ticker_against_database_and_within_file(pa):
    seed(pa)
    obj = pa.sql("SELECT id FROM bams_objects")[0]["id"]
    csvt = f"id,object_id,metric_name,ticker\nm-a,{obj},Other,T.CPU\nm-b,{obj},M2,T.NEW\nm-c,{obj},M3,T.NEW\n"
    item = ok_item(up(pa, csvt.encode(), "Metrics.csv"))
    assert item["totals"] == {"new": 1, "identical": 0, "differs": 0, "invalid": 2, "rows": 3}
    assert {i["code"] for i in item["issues"]} == {"ticker_taken"}


def test_duplicate_id_in_file(pa):
    item = ok_item(up(pa, b"id,name\na1,One\na1,Two\n", "Accounts.csv"))
    assert item["totals"]["new"] == 1 and item["issues"][0]["code"] == "repeats_row"


def test_value_rules_dates_enums_lengths(pa):
    rows = ("id,name,status,next_contact\n"
            "r1,Ok Co,Active,2026-02-30\n"
            "r2,Ok Two,Archived,\n"
            "r3,Ctl\x07Name,Active,\n"
            "r4," + "x" * 201 + ",Active,\n"
            "r5,Fine,Active,\n")
    item = ok_item(up(pa, rows.encode(), "Accounts.csv"))
    assert item["totals"]["invalid"] == 4 and item["totals"]["new"] == 1
    assert {i["code"] for i in item["issues"]} == {"date_invalid", "value_not_allowed", "control_characters", "too_long"}


def test_database_rule_rejection_uses_scratch_schema(pa):
    obj_csv = b"id,client_id,server_name,environment,polarity\nr1,c1,S1,PROD,DEV\n"
    sid = ok_item(up(pa, b"id,name\nc1,Acme\n", "Accounts.csv"))["restore_id"]
    item = ok_item(up(pa, obj_csv, "Operations_Records.csv", session=sid))
    assert item["totals"]["invalid"] == 1 and item["issues"][0]["code"] == "database_rule"


# ------------------------------------------------------------------------------------------ apply safety
def test_atomic_failure_rolls_back_and_reports_backup(pa):
    seed(pa)
    data = export_bytes(pa)
    wipe(pa)
    item = ok_item(up(pa, data, "e.xlsx"))

    def boom(n, key):
        if n == 5:
            raise RuntimeError("injected")
    RS.HOOKS["after_write"] = boom
    try:
        r = pa.post("/api/restore/apply", {"restore_id": item["restore_id"], "plan_hash": item["plan_hash"], "confirm": True})
    finally:
        RS.HOOKS.clear()
    assert r.status_code == 500 and r.json()["error"]["code"] == "import_failed"
    assert r.json()["error"]["details"]["rolled_back"] is True and r.json()["error"]["details"]["backup"]
    assert sum(counts(pa).values()) == 0 and not pa.audit(action="restored")
    assert apply(pa, item)["item"]["counts"]["added"] > 0


def test_preview_stale_confirm_and_unknown_session(pa):
    seed(pa)
    data = export_bytes(pa)
    wipe(pa)
    item = ok_item(up(pa, data, "e.xlsx"))
    pa.err("/api/restore/apply", {"restore_id": item["restore_id"], "plan_hash": item["plan_hash"]}, 409, "confirmation_required")
    pa.exec("INSERT INTO clients (id, name) VALUES ('late', 'Acme Capital')")      # now a name clash appears
    pa.err("/api/restore/apply", {"restore_id": item["restore_id"], "plan_hash": item["plan_hash"], "confirm": True}, 409, "preview_stale")
    assert counts(pa)["accounts"] == 1
    pa.err("/api/restore/apply", {"restore_id": "nope", "plan_hash": "x", "confirm": True}, 410, "import_expired")


def test_session_expiry_and_limit(pa):
    item = ok_item(up(pa, b"id,name\nq1,Q\n", "Accounts.csv"))
    st = RS.store(pa.rt)
    st.now = lambda: 10 ** 9
    assert pa.get(f"/api/restore/{item['restore_id']}").status_code == 410
    st.now = time.monotonic
    ids = [ok_item(up(pa, f"id,name\ns{i},S{i}\n".encode(), "Accounts.csv"))["restore_id"] for i in range(4)]
    assert pa.get(f"/api/restore/{ids[0]}").status_code == 410 and pa.get(f"/api/restore/{ids[3]}").status_code == 200


def test_get_plan_filters_paging_and_method_guard(pa):
    csvt = "id,name,status\n" + "\n".join(f"b{i},B{i},Nope" for i in range(5)) + "\nok1,Fine,Active\n"
    item = ok_item(up(pa, csvt.encode(), "Accounts.csv"))
    r = pa.get(f"/api/restore/{item['restore_id']}", outcome="invalid", limit=2, offset=1).json()["item"]
    assert len(r["issues"]) == 2 and r["issues_total"] == 5
    assert pa.get("/api/restore/inspect").status_code == 405 and pa.get("/api/restore/apply").status_code == 405
    assert pa.get(f"/api/restore/{item['restore_id']}", table="nope").status_code == 422


def test_reset_and_unknown_restore_paths_still_501(pa):
    assert pa.c.post("/api/restore/database/x/y/z").status_code == 501
    assert pa.c.post("/api/data/reset").status_code == 501


def test_injection_text_is_inert_in_database(pa):
    evil = "x'); DROP TABLE clients;--"
    item = ok_item(up(pa, f'id,name\ninj1,"{evil}"\n'.encode(), "Accounts.csv"))
    apply(pa, item)
    assert pa.sql("SELECT name FROM clients WHERE id='inj1'")[0]["name"] == evil
    assert pa.sql("SELECT count(*) n FROM clients")[0]["n"] == 1


def test_fixture_headers_match_table_definitions():
    sys.path.insert(0, str(FX))
    import build_fixtures as bf
    for sheet, (_, cols) in bf.SHEETS.items():
        t = T.BY_SHEET[sheet.casefold()]
        assert set(cols) <= {c.header for c in t.cols}, (sheet, set(cols) - {c.header for c in t.cols})


def test_zip_bomb_and_external_link_refused(pa):
    data = export_bytes(pa)
    bomb = rezip(data, add={"xl/bomb.xml": b"<a>" + b"0" * (30 * 1024 * 1024) + b"</a>"})
    assert up(pa, bomb, "b.xlsx").json()["error"]["details"]["reason"] == "zip_unsafe"
    ext = rezip(data, add={"xl/externalLinks/externalLink1.xml": b"<x/>"})
    assert up(pa, ext, "l.xlsx").json()["error"]["details"]["reason"] == "external_links"


def test_xlsx_files_list_shows_the_workbook_once_with_its_sheets(pa):
    """QA-7-2"""
    seed(pa)
    data = export_bytes(pa)
    item = ok_item(up(pa, data, "e.xlsx"))
    assert len(item["files"]) == 1 and item["files"][0]["name"] == "e.xlsx"
    f = item["files"][0]
    assert [s["table"] for s in f["sheets"]] == [t.key for t in T.TABLES] and f["rows"] == sum(s["rows"] for s in f["sheets"])
    wipe(pa)
    item = ok_item(up(pa, data, "e.xlsx"))
    res = apply(pa, item)["item"]
    assert res["files"] == ["e.xlsx"]
    sid = ok_item(up(pa, data, "e.xlsx"))
    gone = pa.post(f"/api/restore/{sid['restore_id']}/remove", {"file_id": sid["files"][0]["file_id"]}).json()["item"]
    assert gone["files"] == [] and gone["totals"]["rows"] == 0                  # one remove drops every sheet


def test_leading_tab_and_spacing_have_specific_reasons(pa):
    """QA-7-3"""
    rows = 'id,name\nt1,"\tTabbed"\nt2,"Line\nBreak"\nt3," Padded"\nt4,Fine\n'
    item = ok_item(up(pa, rows.encode(), "Accounts.csv"))
    by = {i["row"]: i for i in item["issues"]}
    assert by[2]["code"] == "tab_or_linebreak" and "starts with a tab" in by[2]["text"] and by[2]["column"] == "name"
    assert by[3]["code"] == "tab_or_linebreak" and by[4]["code"] == "leading_trailing_space"
    assert item["totals"]["new"] == 1 and not any(i["code"] == "database_rule" for i in item["issues"])
