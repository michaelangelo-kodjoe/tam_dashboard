"""BAMS import template (xlsx/csv), generated per request: no data, no formulas, no macros, no links (D-64)."""
from __future__ import annotations

import csv
import io

from .. import constants as C
from .bams_reader import HEADERS
from .reader import _load_openpyxl
from .template import CSV_MEDIA, LAST_ROW, XLSX_MEDIA  # noqa: F401  (re-exported for the service)

INSTRUCTIONS = [
    "TAM Command Center: BAMS appliance and metrics import",
    "",
    "Fill the Import sheet and upload it on the BAMS import page. A row may update appliance metadata only, or may "
    "contain a Metric + Ticker pair. Nothing is written until you confirm on the Preview page.",
    "",
    "Columns (row 1 is the header; keep the names, the order may change):",
    "Account: the exact account name already in the tracker. This is the only universally required column/value. "
    "BPID is optional and can be used as an additional match check. Accounts are never created by this import.",
    "Appliance: required only when the row is making a BAMS change. For an appliance-only metadata update, leave Metric and Ticker blank.",
    "Metric + Ticker: optional as a pair. If either is supplied, both are required. A ticker belongs to exactly one metric and is never changed by import.",
    "Region, Location and Environment (PROD or DEV): blank means leave the existing appliance value unchanged; an explicit non-blank value updates it.",
    "Polarity: optional East or West. When supplied, it is authoritative for the appliance's East/West Operations polarity and overrides Region for that row.",
    "Enabled (Yes or No, blank means Yes) and Stale after (s) (whole number above 0, blank means 120) apply to new metrics only.",
    "",
    "Values are percentages as returned by Bloomberg; there is no unit column. Limits: up to " + str(C.IMPORT_MAX_ROWS) +
    " rows and 5 MB per file. Formulas are not accepted: paste values. Alert rules, current values and history are not imported.",
    "You can also upload a CSV (.csv) file with the same column names.",
]
EXAMPLE = [
    ["Example Capital", "", "BA49686", "", "", "", "", "PROD", "", "", "West"],
    ["Example Capital", "", "BA51186", "", "", "", "", "PROD", "", "", "East"],
    ["Example Capital", "", "BPOD-NY-01", "CPU utilisation", "EXAMPLE.BPOD.NY01.CPU", "", "New York", "PROD", "Yes", "120", "East"],
]


def build_template_xlsx() -> bytes:
    openpyxl = _load_openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        c = ins.cell(row=i, column=1, value=line)
        c.data_type = "s"
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ins["A1"].font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 110
    ws = wb.create_sheet("Import")
    ex = wb.create_sheet("Example")
    lists = wb.create_sheet("Lists")
    for i, v in enumerate(("PROD", "DEV"), start=1):
        lists.cell(row=i, column=1, value=v)
    for i, v in enumerate(("Yes", "No"), start=1):
        lists.cell(row=i, column=2, value=v)
    for i, v in enumerate(("East", "West"), start=1):
        lists.cell(row=i, column=3, value=v)
    lists.sheet_state = "hidden"
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    req_fill = PatternFill("solid", fgColor="FDE9D9")
    for sheet in (ws, ex):
        for col, h in enumerate(HEADERS, start=1):
            c = sheet.cell(row=1, column=col, value=h)
            c.font = Font(bold=True)
            c.fill = req_fill if h == "Account" else head_fill
        for i, w in enumerate([32, 12, 22, 26, 30, 14, 18, 14, 10, 16, 12]):
            sheet.column_dimensions[chr(65 + i)].width = w
        sheet.freeze_panes = "A2"
    for r, vals in enumerate(EXAMPLE, start=2):
        for col, v in enumerate(vals, start=1):
            c = ex.cell(row=r, column=col, value=v)
            c.data_type = "s"
    for r in range(2, LAST_ROW + 1):
        ws.cell(row=r, column=2).number_format = "@"
    rng = lambda col: f"{col}2:{col}{LAST_ROW}"  # noqa: E731
    dv_env = DataValidation(type="list", formula1="=Lists!$A$1:$A$2", allow_blank=True, showErrorMessage=True,
                            errorTitle="Environment", error="Choose PROD or DEV.")
    dv_en = DataValidation(type="list", formula1="=Lists!$B$1:$B$2", allow_blank=True, showErrorMessage=True,
                           errorTitle="Enabled", error="Choose Yes or No.")
    dv_stale = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True,
                              showErrorMessage=True, errorTitle="Stale after (s)", error="Use a whole number above 0.")
    dv_pol = DataValidation(type="list", formula1="=Lists!$C$1:$C$2", allow_blank=True, showErrorMessage=True,
                            errorStyle="stop", errorTitle="Polarity",
                            error="Choose East or West.")
    for dv, col in ((dv_env, "H"), (dv_en, "I"), (dv_stale, "J"), (dv_pol, "K")):
        ws.add_data_validation(dv)
        dv.add(rng(col))
    wb.properties.creator = "TAM Command Center"
    wb.active = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template_csv() -> bytes:
    out = io.StringIO(newline="")
    csv.writer(out, lineterminator="\r\n").writerow(HEADERS)
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")


def build_template(fmt: str) -> bytes:
    return build_template_csv() if fmt == "csv" else build_template_xlsx()
