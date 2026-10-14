"""Import template and skipped-rows files (xlsx via openpyxl, csv via the csv module). Generated per request, in
memory, from constants; never stored; no account data, no formula, no macro, no external link (D-37, D-46)."""
from __future__ import annotations

import csv
import io
from typing import Any, Dict, List, Optional, Sequence

from .. import constants as C
from .reader import FIELDS, LABELS, _load_openpyxl

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA = "text/csv; charset=utf-8"
HEADERS = [LABELS[f] for f in FIELDS]
EXTRA_HEADERS = ["Source row", "Reason", "Possible duplicate of"]
LAST_ROW = 501
FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")

INSTRUCTIONS = [
    "TAM Command Center: account import",
    "",
    "Fill the Accounts sheet, save, and upload the file on the Accounts import page. Nothing is written until you "
    "confirm on the Preview page.",
    "",
    "Columns (row 1 is the header; keep the names, the order may change):",
    "Account name (required, up to 200 characters). An account with the same name already in the tracker is "
    "matched automatically: only blank fields are filled, existing values are never overwritten.",
    "BPID (optional, up to 32 characters, stored as text so leading zeros survive).",
    "Products (optional free text, up to 500 characters).",
    "Status (optional): " + " or ".join(C.CLIENT_STATUSES) + ". Blank means Active for a new account. "
    "An Inactive account stays Inactive.",
    "Next contact date (optional): type it as YYYY-MM-DD, for example 2026-11-05.",
    "Notes (optional, up to 20,000 characters).",
    "",
    "Limits: up to " + str(C.IMPORT_MAX_ROWS) + " account rows and 5 MB per file. Formulas are not accepted: "
    "paste values. Contacts, servers, tasks and notes records are not imported.",
    "You can also upload a CSV (.csv) file with the same column names (comma, semicolon or tab separated).",
]


def build_template_xlsx() -> bytes:
    openpyxl = _load_openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        c = ins.cell(row=i, column=1, value=line)
        c.data_type = "s"
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ins["A1"].font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 110
    ws = wb.create_sheet("Accounts")
    lists = wb.create_sheet("Lists")
    for i, st in enumerate(C.CLIENT_STATUSES, start=1):
        lists.cell(row=i, column=1, value=st)
    lists.sheet_state = "hidden"
    req_fill = PatternFill("solid", fgColor="FDE9D9")
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = req_fill if col == 1 else head_fill
    widths = [34, 14, 30, 12, 20, 50]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    for r in range(2, LAST_ROW + 1):
        ws.cell(row=r, column=2).number_format = "@"
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
    rng = lambda col: f"{col}2:{col}{LAST_ROW}"  # noqa: E731
    dv_name = DataValidation(type="textLength", operator="between", formula1="1", formula2="200", allow_blank=True,
                             showInputMessage=True, promptTitle="Account name (required)",
                             prompt="Required. Up to 200 characters.", showErrorMessage=True,
                             errorTitle="Account name", error="Use 1 to 200 characters.")
    dv_status = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(C.CLIENT_STATUSES)}", allow_blank=True,
                               showInputMessage=True, promptTitle="Status",
                               prompt="Choose " + " or ".join(C.CLIENT_STATUSES) + ". Blank means Active.",
                               showErrorMessage=True, errorTitle="Status",
                               error="Choose " + " or ".join(C.CLIENT_STATUSES) + ".")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                             showInputMessage=True, promptTitle="Next contact date",
                             prompt="Type the date as YYYY-MM-DD, for example 2026-11-05.", showErrorMessage=True,
                             errorTitle="Next contact date", error="Use a date in YYYY-MM-DD format.")
    for dv, col in ((dv_name, "A"), (dv_status, "D"), (dv_date, "E")):
        ws.add_data_validation(dv)
        dv.add(rng(col))
    wb.properties.creator = "TAM Command Center"
    wb.active = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template_csv() -> bytes:
    out = io.StringIO(newline="")
    csv.writer(out, lineterminator="\r\n").writerow(HEADERS)
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")        # BOM: Excel opens UTF-8 correctly


def build_template(fmt: str) -> bytes:
    return build_template_csv() if fmt == "csv" else build_template_xlsx()


def _as_text(v: Optional[str]) -> str:
    return "" if v is None else str(v)


def build_skipped_xlsx(rows: Sequence[Sequence[str]]) -> bytes:
    """`rows`: lists of the 9 string columns (six import columns, Source row, Reason, Possible duplicate of).
    Every value is stored as a string (never a formula)."""
    openpyxl = _load_openpyxl()
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounts"
    for col, h in enumerate(HEADERS + EXTRA_HEADERS, start=1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for r, vals in enumerate(rows, start=2):
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col)
            c.value = _as_text(v)[:32767]
            c.data_type = "s"                       # a value like =SUM(A1) stays text
            if c.value.startswith(FORMULA_LEAD):
                c.quotePrefix = True
    for i, w in enumerate([34, 14, 30, 12, 20, 50, 11, 60, 34]):
        ws.column_dimensions[chr(65 + i)].width = w
    wb.properties.creator = "TAM Command Center"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_skipped_csv(rows: Sequence[Sequence[str]]) -> bytes:
    out = io.StringIO(newline="")
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow(HEADERS + EXTRA_HEADERS)
    for vals in rows:
        w.writerow([("'" + v if v.startswith(FORMULA_LEAD) else v) for v in (_as_text(x) for x in vals)])
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")
