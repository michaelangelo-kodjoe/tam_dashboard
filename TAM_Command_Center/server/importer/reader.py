"""Hardened reader/validator for account import files (PM D-37, D-38, D-47; user amendment: .xlsx AND .csv).

The upload is untrusted. Order of defence for .xlsx: size cap -> ZIP signature -> ZIP structure (entry count, sizes,
ratios, names, encryption) -> forbidden parts (macros, external links, embedded objects) -> XML parts scanned for
DOCTYPE/ENTITY -> only then `openpyxl` (read-only, links not kept, formulas never evaluated, cached values never
read). Nothing is extracted to disk. A refusal raises InvalidFile (or PayloadTooLarge) with `details.reason`.
.csv files are decoded (BOM / UTF-8 / Windows-1252), the delimiter is sniffed, and every cell is plain text.
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from .. import constants as C
from ..dupes import clean_name
from ..errors import InvalidFile, PayloadTooLarge

FIELDS = ("name", "bpid", "products", "status", "next_contact", "notes")
LABELS = {"name": "Account name", "bpid": "BPID", "products": "Products", "status": "Status",
          "next_contact": "Next contact date", "notes": "Notes"}
HEADER_ALIASES = {"account name": "name", "bpid": "bpid", "products": "products", "status": "status",
                  "next contact date": "next_contact", "notes": "notes"}
LIMITS = {"name": 200, "bpid": 32, "products": 500, "notes": 20000}

MAX_ENTRIES = 200
MAX_TOTAL_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ENTRY = 20 * 1024 * 1024
MAX_RATIO = 100
MAX_HEADER_COLS = 100
RAW_KEEP = 32767                      # Excel's own cell limit; raw text kept for the skipped file

_CFB_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ENC_MARK = "EncryptedPackage".encode("utf-16-le")
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\u200b-\u200f\u2028\u2029\u202a-\u202e\u2066-\u2069\ufeff]")
_WS = re.compile(r"\s+")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_XML_EXT = (".xml", ".rels", ".vml")
_DANGEROUS = (b"<!doctype", b"<!entity")
_DANGEROUS_ENC = tuple(p.decode().encode(enc) for p in _DANGEROUS for enc in ("utf-16-le", "utf-16-be"))


class Formula(str):
    """A formula cell: the text is never evaluated and never used as a value."""


class CellError(str):
    """An Excel error cell such as #REF!."""


@dataclass
class ParsedRow:
    row: int
    values: Dict[str, Optional[str]]
    raw: Dict[str, str]
    errors: List[Dict[str, str]] = field(default_factory=list)

    @property
    def valid(self) -> bool:
        return not self.errors


@dataclass
class ParsedFile:
    format: str                       # "xlsx" | "csv"
    sheet: str
    rows: List[ParsedRow]
    rows_found: int
    warnings: List[Dict[str, str]]
    size_bytes: int
    encoding: Optional[str] = None
    delimiter: Optional[str] = None
    notes: List[str] = field(default_factory=list)


@dataclass(frozen=True)
class FileSpec:
    """What differs between import files: the sheet, the columns and the row parser (accounts vs BAMS)."""
    sheet: str
    aliases: Dict[str, str]
    labels: Dict[str, str]
    parse_row: Any
    check_headers: Any                 # (mapping, header_cells) -> None, raises InvalidFile
    noun: str                          # "account" / "metric"
    no_rows_hint: str


# ======================================================================= helpers


def display_file_name(raw: Optional[str]) -> str:
    """Display only: strip path parts and control characters, <= 200 chars."""
    name = (raw or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = _WS.sub(" ", _CTRL.sub("", name)).strip()
    return name[:200] or "import"


def _bad(reason: str, message: str, **details: Any) -> InvalidFile:
    return InvalidFile(message, details={"reason": reason, **details})


def _clip(s: str, n: int = 60) -> str:
    s = _WS.sub(" ", _CTRL.sub("", s)).strip()
    return s if len(s) <= n else s[:n - 1] + "…"


def detect_format(data: bytes, filename: str) -> str:
    """Decide by content AND extension; either alone is not trusted."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in ("xlsx", "csv"):
        raise _bad("unsupported_type", "Only Excel (.xlsx) or CSV (.csv) files can be imported.")
    if data[:8] == _CFB_MAGIC:
        if _ENC_MARK in data:
            raise _bad("encrypted", "This workbook is password protected. Remove the password and upload it again.")
        raise _bad("not_xlsx", "This is an older Excel (.xls) or other non-.xlsx file. "
                               "Open it in Excel and save it as Excel Workbook (.xlsx) or CSV (.csv).")
    is_zip = data[:4] in (b"PK\x03\x04", b"PK\x05\x06")
    if ext == "xlsx":
        if not is_zip:
            raise _bad("not_xlsx", "This file is not a real .xlsx workbook. Open it in Excel and save it as .xlsx.")
        return "xlsx"
    if is_zip:
        raise _bad("not_csv", "This file is named .csv but is a compressed/Excel file. Save it as CSV (.csv) first.")
    return "csv"


# ======================================================================= ZIP / XML safety (.xlsx)


def _read_bounded(zf: zipfile.ZipFile, info: zipfile.ZipInfo, budget: List[int]) -> bytes:
    out = bytearray()
    try:
        with zf.open(info) as fh:
            while True:
                chunk = fh.read(65536)
                if not chunk:
                    break
                out += chunk
                budget[0] -= len(chunk)
                if len(out) > MAX_ENTRY or budget[0] < 0:
                    raise _bad("zip_unsafe", "The workbook expands to an unsafe size and was refused.")
    except InvalidFile:
        raise
    except (zipfile.BadZipFile, RuntimeError, EOFError, zlib_error()) as exc:
        raise _bad("corrupt", "The workbook is damaged and cannot be read.") from exc
    return bytes(out)


def zlib_error() -> type:
    import zlib
    return zlib.error


def check_zip(data: bytes) -> zipfile.ZipFile:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except (zipfile.BadZipFile, ValueError, OSError, EOFError) as exc:
        raise _bad("corrupt", "The workbook is damaged and cannot be read.") from exc
    infos = zf.infolist()
    if not infos:
        raise _bad("corrupt", "The workbook is empty.")
    if len(infos) > MAX_ENTRIES:
        raise _bad("zip_unsafe", f"The workbook has too many internal parts ({len(infos)}; limit {MAX_ENTRIES}).")
    total = 0
    seen = set()
    for i in infos:
        n = i.filename
        low = n.lower()
        if i.flag_bits & 0x1:
            raise _bad("encrypted", "This workbook is password protected. Remove the password and upload it again.")
        if (n.startswith(("/", "\\")) or "\\" in n or re.match(r"^[A-Za-z]:", n)
                or any(seg == ".." for seg in n.split("/")) or "\x00" in n):
            raise _bad("zip_unsafe", "The workbook contains an unsafe internal path and was refused.")
        if low in seen:
            raise _bad("zip_unsafe", "The workbook contains duplicate internal parts and was refused.")
        seen.add(low)
        if i.file_size > MAX_ENTRY:
            raise _bad("zip_unsafe", "The workbook contains an oversized internal part and was refused.")
        if i.file_size > 0 and (i.compress_size == 0 or i.file_size / i.compress_size > MAX_RATIO):
            raise _bad("zip_unsafe", "The workbook has an abnormal compression ratio and was refused.")
        total += i.file_size
        if total > MAX_TOTAL_UNCOMPRESSED:
            raise _bad("zip_unsafe", "The workbook expands to an unsafe size and was refused.")
        # forbidden content
        if "vbaproject" in low or low.endswith(".vba"):
            raise _bad("macros", "Workbooks with macros are not accepted. Save a macro-free .xlsx and upload it.")
        if low.startswith("xl/externallinks/"):
            raise _bad("external_links", "Workbooks with links to other files are not accepted. Paste the values instead.")
        if (low.startswith(("xl/embeddings/", "xl/activex/", "xl/ole")) or "oleobject" in low
                or (low.endswith(".bin") and not low.startswith("xl/printersettings/"))):
            raise _bad("embedded_objects", "Workbooks with embedded objects are not accepted.")
    names = {i.filename for i in infos}
    if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
        raise _bad("not_xlsx", "This file is not a real .xlsx workbook. Open it in Excel and save it as .xlsx.")
    budget = [MAX_TOTAL_UNCOMPRESSED]
    for i in infos:
        low = i.filename.lower()
        if not low.endswith(_XML_EXT):
            continue
        body = _read_bounded(zf, i, budget)
        lowered = body.lower()
        if any(p in lowered for p in _DANGEROUS) or any(p in lowered for p in _DANGEROUS_ENC):
            raise _bad("xml_unsafe", "The workbook contains unsafe XML (DOCTYPE/ENTITY) and was refused.")
        if low == "[content_types].xml":
            if b"macroenabled" in lowered or b"vbaproject" in lowered:
                raise _bad("macros", "Workbooks with macros are not accepted. Save a macro-free .xlsx and upload it.")
            if b"spreadsheetml" not in lowered:
                raise _bad("not_xlsx", "This file is not a real .xlsx workbook. Open it in Excel and save it as .xlsx.")
        if low.endswith(".rels") and re.search(rb'targetmode\s*=\s*["\']external["\']', lowered):
            raise _bad("external_links", "Workbooks with links to other files are not accepted. Paste the values instead.")
    return zf


# ======================================================================= rows (shared by .xlsx and .csv)


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not isinstance(v, (Formula, CellError)) and not v.strip())


def _raw_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, _dt.datetime):
        return v.date().isoformat() if v.time() == _dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, (_dt.date, _dt.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)[:RAW_KEEP]


def _err(row: int, fld: str, code: str, problem: str, labels: Optional[Dict[str, str]] = None) -> Dict[str, str]:
    return {"field": fld, "code": code,
            "text": f"Row {row}: {(labels or LABELS)[fld]}: {problem}. This row will be skipped; fix it in the file and "
                    "re-upload, or leave it out."}


def _text(v: Any, fld: str, row: int, errors: List[Dict[str, str]], labels: Optional[Dict[str, str]] = None) -> Optional[str]:
    if _is_blank(v):
        return None
    if isinstance(v, Formula):
        errors.append(_err(row, fld, "formula_not_allowed", "formulas are not accepted; paste the value", labels))
        return None
    if isinstance(v, CellError):
        errors.append(_err(row, fld, "cell_error", f"the cell holds an error value ({_clip(str(v), 20)}); fix or clear it", labels))
        return None
    if isinstance(v, bool):
        errors.append(_err(row, fld, "value_type", "TRUE/FALSE is not text; type the value as text", labels))
        return None
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        errors.append(_err(row, fld, "value_type", "a date or time is not accepted here; type the value as text", labels))
        return None
    if isinstance(v, float):
        if fld == "bpid" and not v.is_integer():
            errors.append(_err(row, fld, "bpid_invalid", f"'{_raw_text(v)}' is not a whole number; type the BPID as text", labels))
            return None
        v = str(int(v)) if v.is_integer() else repr(v)
    elif isinstance(v, int):
        v = str(v)
    s = _CTRL.sub("", str(v))
    if fld == "notes":
        s = s.replace("\r\n", "\n").replace("\r", "\n").strip()
    else:
        s = _WS.sub(" ", s).strip()
    s = unicodedata.normalize("NFC", s)
    return s or None


def parse_row(rownum: int, cells: Dict[str, Any]) -> ParsedRow:
    errors: List[Dict[str, str]] = []
    values: Dict[str, Optional[str]] = {f: None for f in FIELDS}
    raw = {f: _raw_text(cells.get(f)) for f in FIELDS}
    for fld in FIELDS:
        if fld == "next_contact":
            continue
        values[fld] = _text(cells.get(fld), fld, rownum, errors)
    # date
    v = cells.get("next_contact")
    if not _is_blank(v):
        hint = "use YYYY-MM-DD, for example 2026-11-05"
        if isinstance(v, Formula):
            errors.append(_err(rownum, "next_contact", "formula_not_allowed", "formulas are not accepted; paste the value"))
        elif isinstance(v, CellError):
            errors.append(_err(rownum, "next_contact", "cell_error", "the cell holds an error value; fix or clear it"))
        elif isinstance(v, _dt.datetime):
            values["next_contact"] = v.date().isoformat()
        elif isinstance(v, _dt.date):
            values["next_contact"] = v.isoformat()
        elif isinstance(v, str) and _DATE.match(v.strip()):
            try:
                values["next_contact"] = _dt.date.fromisoformat(v.strip()).isoformat()
            except ValueError:
                errors.append(_err(rownum, "next_contact", "date_invalid", f"'{_clip(v)}' is not a real calendar date; {hint}"))
        else:
            errors.append(_err(rownum, "next_contact", "date_invalid", f"'{_clip(_raw_text(v))}' is not a valid date; {hint}"))
        if values["next_contact"] and not 1900 <= int(values["next_contact"][:4]) <= 2200:
            values["next_contact"] = None
            errors.append(_err(rownum, "next_contact", "date_invalid", f"the year is out of range; {hint}"))
    # name / lengths / status
    if values["name"] is not None:
        values["name"] = clean_name(values["name"])
    if values["name"] is None and not any(e["field"] == "name" for e in errors):
        errors.append(_err(rownum, "name", "name_required", "is required"))
    for fld, limit in LIMITS.items():
        val = values[fld]
        if val is not None and len(val) > limit:
            errors.append(_err(rownum, fld, f"{fld}_too_long",
                               f"is {len(val)} characters; the limit is {limit}. Shorten it (nothing is cut automatically)"))
            values[fld] = None
    st = values["status"]
    if st is not None:
        canon = {s.casefold(): s for s in C.CLIENT_STATUSES}.get(st.casefold())
        if canon is None:
            errors.append(_err(rownum, "status", "status_not_allowed",
                               f"'{_clip(st, 30)}' is not in the list ({', '.join(C.CLIENT_STATUSES)})"))
            values["status"] = None
        else:
            values["status"] = canon
    return ParsedRow(row=rownum, values=values, raw=raw, errors=errors)


def map_headers(header_cells: List[Any], spec: Optional["FileSpec"] = None) -> Tuple[Dict[str, int], List[Dict[str, str]]]:
    """-> ({field: column index (0-based)}, warnings). Raises InvalidFile for missing/duplicate headers."""
    mapping: Dict[str, int] = {}
    warnings: List[Dict[str, str]] = []
    for idx, v in enumerate(header_cells):
        if _is_blank(v):
            continue
        key = _WS.sub(" ", _CTRL.sub("", str(v))).strip().rstrip("*").strip().casefold() if not isinstance(v, Formula) else ""
        fld = (spec.aliases if spec else HEADER_ALIASES).get(key)
        col = _col_letter(idx)
        if fld is None:
            warnings.append({"code": "unknown_column", "text": f'Column {col} "{_clip(str(v), 40)}" is ignored.'})
        elif fld in mapping:
            lab = (spec.labels if spec else LABELS)[fld]
            raise _bad("header_duplicate", f'The column "{lab}" appears twice in row 1. Keep one.', column=lab)
        else:
            mapping[fld] = idx
    if spec is not None:
        spec.check_headers(mapping, header_cells)
    elif "name" not in mapping:
        raise _bad("header_missing", 'Row 1 must contain the column "Account name". Download the template to see the layout.')
    return mapping, warnings


def _col_letter(idx: int) -> str:
    s, n = "", idx + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _finish(rows_cells: List[Tuple[int, Dict[str, Any]]], found: int, spec: Optional["FileSpec"] = None) -> List[ParsedRow]:
    fn = spec.parse_row if spec else parse_row
    return [fn(n, c) for n, c in rows_cells]


def _too_many(found: int, noun: str = "account") -> InvalidFile:
    return _bad("too_many_rows", f"This file has {found} {noun} rows. The limit is {C.IMPORT_MAX_ROWS} per import. "
                                 "Split the file and import it in parts.", found=found, limit=C.IMPORT_MAX_ROWS)


# ======================================================================= .xlsx


def _load_openpyxl() -> Any:
    try:
        import openpyxl  # noqa: PLC0415  (lazy: the app must start without it)
        return openpyxl
    except ImportError as exc:                                  # pragma: no cover - exercised by a test
        from ..errors import AppError
        import logging
        logging.getLogger("tam.import").error("openpyxl is not installed; account import and the template are "
                                              "unavailable (pip install -r server/requirements.txt)")
        raise AppError("Account import is unavailable on this installation (a required component is missing). "
                       "See the log.") from exc


def read_xlsx(data: bytes, spec: Optional[FileSpec] = None) -> ParsedFile:
    zf = check_zip(data)
    zf.close()
    openpyxl = _load_openpyxl()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # noqa: BLE001 - any parser failure on an untrusted file is "corrupt"
        raise _bad("corrupt", "The workbook is damaged and cannot be read.") from exc
    try:
        names = list(wb.sheetnames)
        sheet_name = spec.sheet if spec else "Accounts"
        noun = spec.noun if spec else "account"
        match = [n for n in names if n.strip().casefold() == sheet_name.casefold()]
        if not match:
            raise _bad("sheet_missing", f'The workbook has no sheet named "{sheet_name}". '
                                        'Sheets found: ' + (", ".join(_clip(n, 30) for n in names[:10]) or "none") + ".",
                       sheets=[_clip(n, 30) for n in names[:10]])
        ws = wb[match[0]]
        declared = None
        try:
            declared = ws.max_row
        except Exception:  # noqa: BLE001
            declared = None
        if declared and declared > C.IMPORT_SCAN_ROWS:
            raise _bad("too_many_rows", f"The {sheet_name} sheet is formatted down to row {declared}, beyond the "
                                        f"{C.IMPORT_SCAN_ROWS} rows that are read. Delete the empty formatted rows "
                                        f"(or split the file; the limit is {C.IMPORT_MAX_ROWS} {noun} rows).",
                       found=declared, limit=C.IMPORT_MAX_ROWS, physical=True)
        try:
            header = next(ws.iter_rows(min_row=1, max_row=1, max_col=MAX_HEADER_COLS), ())
        except StopIteration:
            header = ()
        hvals = [_cell_value(c) for c in header]
        mapping, warnings = map_headers(hvals, spec)
        maxcol = max(mapping.values()) + 1
        rows_cells: List[Tuple[int, Dict[str, Any]]] = []
        found = 0
        for rownum, row in enumerate(ws.iter_rows(min_row=2, max_row=C.IMPORT_SCAN_ROWS, max_col=maxcol), start=2):
            cells = {f: (_cell_value(row[i]) if i < len(row) else None) for f, i in mapping.items()}
            if all(_is_blank(v) for v in cells.values()):
                continue
            found += 1
            if found <= C.IMPORT_MAX_ROWS:
                rows_cells.append((rownum, cells))
        probe = next(ws.iter_rows(min_row=C.IMPORT_SCAN_ROWS + 1, max_row=C.IMPORT_SCAN_ROWS + 1, max_col=maxcol), ())
        if any(not _is_blank(_cell_value(c)) for c in probe):
            raise _too_many(found + 1, noun)
        if found > C.IMPORT_MAX_ROWS:
            raise _too_many(found, noun)
        if found == 0:
            raise _bad("no_rows", f"The file has no {noun} rows under the header. " + (spec.no_rows_hint if spec else "Add accounts and upload it again."))
        return ParsedFile(format="xlsx", sheet=ws.title, rows=_finish(rows_cells, found, spec), rows_found=found,
                          warnings=warnings, size_bytes=len(data))
    finally:
        wb.close()


def _cell_value(cell: Any) -> Any:
    v = getattr(cell, "value", None)
    t = getattr(cell, "data_type", "n")
    if t == "f":
        return Formula(str(v) if v is not None else "=")
    if t == "e":
        return CellError(str(v))
    return v


# ======================================================================= .csv


def _decode(data: bytes) -> Tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", "strict"), "UTF-8 (with BOM)"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "UTF-16"
    try:
        return data.decode("utf-8", "strict"), "UTF-8"
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1252", "strict"), "Windows-1252"
    except UnicodeDecodeError:
        return data.decode("latin-1"), "ISO-8859-1"


def _sniff_delimiter(text: str) -> str:
    """Count candidate delimiters outside quotes on the header record; ties prefer comma, semicolon, tab, pipe."""
    counts = {d: 0 for d in ",;\t|"}
    inq = False
    for ch in text[:20000]:
        if ch == '"':
            inq = not inq
        elif not inq:
            if ch in "\r\n":
                break
            if ch in counts:
                counts[ch] += 1
    best = max(",;\t|", key=lambda d: counts[d])
    return best if counts[best] > 0 else ","


def _csv_cell(s: str) -> Any:
    s = s if len(s) <= RAW_KEEP * 4 else s[:RAW_KEEP * 4]
    if len(s) > 1 and s[0] == "'" and s[1] in "=+-@\t\r":      # our own export guard (see skipped.py); data otherwise verbatim
        s = s[1:]
    return s


def read_csv(data: bytes, spec: Optional[FileSpec] = None) -> ParsedFile:
    if b"\x00" in data and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise _bad("not_csv", "This does not look like a text CSV file.")
    try:
        text, enc = _decode(data)
    except UnicodeError as exc:
        raise _bad("not_csv", "The CSV text could not be decoded. Save it as UTF-8 and upload it again.") from exc
    delim = _sniff_delimiter(text)
    csv.field_size_limit(C.MAX_IMPORT_BYTES + 1)
    rows: List[List[str]] = []
    try:
        for rec in csv.reader(io.StringIO(text, newline=""), delimiter=delim, strict=False):
            rows.append(rec)
            if len(rows) > C.IMPORT_SCAN_ROWS + 1:
                break
    except csv.Error as exc:
        raise _bad("corrupt", "The CSV file is malformed and cannot be read.") from exc
    noun = spec.noun if spec else "account"
    if not rows:
        if spec is not None:
            spec.check_headers({}, [])
        raise _bad("header_missing", 'Row 1 must contain the column "Account name". Download the template to see the layout.')
    mapping, warnings = map_headers([_csv_cell(c) for c in rows[0][:MAX_HEADER_COLS]], spec)
    rows_cells: List[Tuple[int, Dict[str, Any]]] = []
    found = 0
    if len(rows) > C.IMPORT_SCAN_ROWS:
        extra = rows[C.IMPORT_SCAN_ROWS:]
        if any(any(c.strip() for c in r) for r in extra):
            raise _too_many(sum(1 for r in rows[1:] if any(c.strip() for c in r)), noun)
    for rownum, rec in enumerate(rows[1:C.IMPORT_SCAN_ROWS], start=2):
        cells = {f: (_csv_cell(rec[i]) if i < len(rec) else None) for f, i in mapping.items()}
        if all(_is_blank(v) for v in cells.values()):
            continue
        found += 1
        if found <= C.IMPORT_MAX_ROWS:
            rows_cells.append((rownum, cells))
    if found > C.IMPORT_MAX_ROWS:
        raise _too_many(found, noun)
    if found == 0:
        raise _bad("no_rows", f"The file has no {noun} rows under the header. " + (spec.no_rows_hint if spec else "Add accounts and upload it again."))
    names = {"," : "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}
    return ParsedFile(format="csv", sheet="CSV", rows=_finish(rows_cells, found, spec), rows_found=found, warnings=warnings,
                      size_bytes=len(data), encoding=enc, delimiter=names[delim],
                      notes=[f"Read as CSV: encoding {enc}, delimiter {names[delim]}."])


# ======================================================================= entry point


def read_upload(data: bytes, filename: str, spec: Optional[FileSpec] = None) -> ParsedFile:
    if len(data) > C.MAX_IMPORT_BYTES:
        raise PayloadTooLarge(f"The file is larger than {C.MAX_IMPORT_BYTES // (1024 * 1024)} MB.",
                              details={"max_bytes": C.MAX_IMPORT_BYTES, "reason": "too_large"})
    if not data:
        raise _bad("corrupt", "The file is empty.")
    fmt = detect_format(data, filename)
    return read_xlsx(data, spec) if fmt == "xlsx" else read_csv(data, spec)
