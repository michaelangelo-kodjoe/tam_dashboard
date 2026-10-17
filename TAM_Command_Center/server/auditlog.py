"""Audit log reads (read-only; audit_events is append-only). Filters are fixed SQL fragments with bound parameters."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import re
import sqlite3
from typing import Any, Dict, List, Optional, Tuple

from . import auditprune
from .dbkit import like_pattern
from .errors import PayloadTooLarge, ValidationFailed
from .sentences import system_sentence

CATEGORIES = ("crm", "operations", "bams", "settings", "data", "system")
SOURCES = ("GUI", "System", "Migration", "Bulk Import", "ChatGPT Bridge")
MAX_LIMIT = 200
MAX_OFFSET = 100_000
EXPORT_MAX = 50_000
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_COLS = ("a.id, a.timestamp, a.category, a.client_id, c.name AS client_name, a.object_type, a.object_id, a.object_label,"
         " a.action, a.old_value, a.new_value, a.reason, a.source")
_FROM = " FROM audit_events a LEFT JOIN clients c ON c.id = a.client_id"


def _bound(v: str, field: str, end: bool) -> str:
    try:
        if _DATE.match(v):
            d = _dt.date.fromisoformat(v)
            return (d + _dt.timedelta(days=1 if end else 0)).isoformat() + "T00:00:00.000Z"
        t = _dt.datetime.fromisoformat(v.replace("Z", "+00:00"))
        t = (t if t.tzinfo else t.replace(tzinfo=_dt.timezone.utc)).astimezone(_dt.timezone.utc)
        return t.strftime("%Y-%m-%dT%H:%M:%S.") + f"{t.microsecond // 1000:03d}Z"
    except ValueError:
        raise ValidationFailed("Request validation failed.", fields={field: "must be YYYY-MM-DD or an ISO date-time"}) from None


def _where(f: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
    w: List[str] = []
    p: Dict[str, Any] = {}
    if f.get("from"):
        w.append("a.timestamp >= :from_ts")
        p["from_ts"] = _bound(f["from"], "from", False)
    if f.get("to"):
        is_date = bool(_DATE.match(f["to"]))
        w.append("a.timestamp < :to_ts" if is_date else "a.timestamp <= :to_ts")
        p["to_ts"] = _bound(f["to"], "to", True)
    for key, col in (("category", "a.category"), ("object_type", "a.object_type"), ("action", "a.action"),
                     ("source", "a.source"), ("client_id", "a.client_id")):
        if f.get(key):
            w.append(f"{col} = :{key}")
            p[key] = f[key]
    if f.get("q"):
        w.append("(a.object_label LIKE :q ESCAPE '\\' OR a.reason LIKE :q ESCAPE '\\' OR a.action LIKE :q ESCAPE '\\'"
                 " OR a.object_type LIKE :q ESCAPE '\\' OR c.name LIKE :q ESCAPE '\\')")
        p["q"] = like_pattern(f["q"])
    return (" WHERE " + " AND ".join(w)) if w else "", p


def _json(v: Any) -> Any:
    if not v:
        return None
    try:
        return json.loads(v)
    except (TypeError, ValueError):
        return None


def summary(r: Dict[str, Any]) -> str:
    t, a, label = r["object_type"], r["action"], r["object_label"]
    if t in ("import", "backup", "schema", "setting"):
        base = label or t
        return f"{base} ({a.replace('_', ' ')})"
    if t == "quarterly_upgrade_window":
        return f"Operations quarter {a}: {label or ''}".strip()
    if t.startswith("bams_") and t not in ("bams_alert", "bams_alert_event"):
        return f"BAMS {t[5:].replace('_', ' ')} {a.replace('_', ' ')}: {label or '(unnamed)'}"
    return system_sentence(t, a, label, r["old_value"], r["new_value"], None, r["source"])


def _item(r: sqlite3.Row) -> Dict[str, Any]:
    d = dict(r)
    return {"id": d["id"], "timestamp": d["timestamp"], "category": d["category"], "client_id": d["client_id"],
            "client_name": d["client_name"], "object_type": d["object_type"], "object_id": d["object_id"],
            "object_label": d["object_label"], "action": d["action"], "source": d["source"], "reason": d["reason"],
            "summary": summary(d), "old_value": _json(d["old_value"]), "new_value": _json(d["new_value"])}


def list_page(conn: sqlite3.Connection, f: Dict[str, Any], limit: int, offset: int) -> Dict[str, Any]:
    where, p = _where(f)
    total = conn.execute("SELECT count(*)" + _FROM + where, p).fetchone()[0]  # nosec B608 - fixed fragments, values bound
    rows = conn.execute("SELECT " + _COLS + _FROM + where + " ORDER BY a.timestamp DESC, a.id DESC LIMIT :lim OFFSET :off",  # nosec B608
                        {**p, "lim": limit, "off": offset}).fetchall()
    return {"ok": True, "total": total, "limit": limit, "offset": offset, "has_more": offset + len(rows) < total,
            "retention": auditprune.status(conn), "items": [_item(r) for r in rows]}


def facets(conn: sqlite3.Connection) -> Dict[str, Any]:
    def distinct(col: str) -> List[str]:
        return [r[0] for r in conn.execute(f"SELECT DISTINCT {col} FROM audit_events WHERE {col} IS NOT NULL ORDER BY {col}")]  # nosec B608 - fixed column names
    return {"ok": True, "categories": distinct("category"), "object_types": distinct("object_type"),
            "actions": distinct("action"), "sources": distinct("source"),
            "retention": auditprune.status(conn)}


EXPORT_HEADERS = ["id", "timestamp", "category", "account", "object_type", "object_label", "action", "source", "reason", "summary"]


def export_rows(conn: sqlite3.Connection, f: Dict[str, Any]) -> List[List[Any]]:
    from .portability.export import guard
    where, p = _where(f)
    n = conn.execute("SELECT count(*)" + _FROM + where, p).fetchone()[0]  # nosec B608 - fixed fragments, values bound
    if n > EXPORT_MAX:
        raise PayloadTooLarge(f"{n} audit rows match; the export is limited to {EXPORT_MAX}. Narrow the filters.",
                              details={"reason": "too_many_rows", "found": n, "limit": EXPORT_MAX})
    out = []
    for r in conn.execute("SELECT " + _COLS + _FROM + where + " ORDER BY a.timestamp DESC, a.id DESC", p):  # nosec B608
        i = _item(r)
        out.append([i["id"], i["timestamp"], i["category"], guard(i["client_name"] or ""), i["object_type"],
                    guard(i["object_label"] or ""), i["action"], i["source"], guard(i["reason"] or ""), guard(i["summary"])])
    return out


def to_csv(rows: List[List[Any]]) -> bytes:
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(EXPORT_HEADERS)
    w.writerows(rows)
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def to_xlsx(rows: List[List[Any]]) -> bytes:
    from .importer.reader import _load_openpyxl
    from openpyxl.styles import Font
    wb = _load_openpyxl().Workbook()
    ws = wb.active
    ws.title = "Audit log"
    ws.append(EXPORT_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
